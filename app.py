
from __future__ import annotations
import re
import threading
import csv, io, os, sqlite3, time, hashlib, hmac, secrets, json
from datetime import datetime, timedelta, timezone
from functools import lru_cache
from pathlib import Path
from typing import Optional
from fastapi import FastAPI, HTTPException, Query, Request, Response, Depends
from fastapi.responses import FileResponse, StreamingResponse, RedirectResponse, HTMLResponse
from fastapi.staticfiles import StaticFiles
from dotenv import load_dotenv
from services.product_catalog_router_v1 import router as product_catalog_router
from market_amount_filter import bounds as _market_amount_bounds, execute as _market_amount_execute, completed_dimension as _amount_completed_dimension, xt_national_dimension as _amount_xt_national_dimension

ROOT = Path(__file__).resolve().parent
load_dotenv(ROOT / ".env", override=False)
STATIC = ROOT / "static"
UZBI_ROOT = Path(os.getenv("UZBI_ROOT", r"C:\Qudrat\xarid_uz\UZBI"))
DATA = UZBI_ROOT / "data"
PRODUCT_SEARCH_DB = DATA / "uzbi_product_search.sqlite3"
ANALYTICS_DB = DATA / "uzbi_analytics.sqlite3"
PRODUCT_CODES_DB = DATA / "product_codes.sqlite3"
ORGANIZATIONS_DB = DATA / "bidata_organizations.sqlite3"

def _org_norm_name(v):
    s=str(v or "").strip().lower()
    s=(s.replace("`","'")
         .replace("?","'")
         .replace("?","'")
         .replace("?","'"))
    s=s.replace("?","?")
    s=re.sub(r"\s+"," ",s)
    s=re.sub(r"[^\w?-??????' -]+"," ",s,flags=re.I)
    s=re.sub(r"\s+"," ",s).strip()
    return s

@lru_cache(maxsize=8192)
def _organization_stir_lookup(org_name):
    """
    BIData SAFE organization -> STIR lookup.
    Read only. Source databases are not modified.
    """
    raw=str(org_name or "").strip()
    if not raw or not ORGANIZATIONS_DB.exists():
        return None

    # Match the normalization used by bidata_organizations.sqlite3.
    n=raw.upper().strip()

    for ch in ('"', '?', '?', '?', '?'):
        n=n.replace(ch, '')

    for ch in ('`', '?', '?', '?', '?', '?'):
        n=n.replace(ch, "'")

    n=n.replace("?","?")
    n=re.sub(r"\s+"," ",n).strip()

    try:
        uri=f"file:{ORGANIZATIONS_DB.as_posix()}?mode=ro"
        c=sqlite3.connect(uri,uri=True,timeout=5)
        try:
            # First: exact normalized_name.
            r=c.execute(
                """
                SELECT stir
                FROM organizations
                WHERE normalized_name=?
                  AND status='SAFE'
                LIMIT 1
                """,
                (n,)
            ).fetchone()

            if r and r[0]:
                return str(r[0]).strip()

            # Second: exact canonical source name.
            r=c.execute(
                """
                SELECT stir
                FROM organizations
                WHERE canonical_name=?
                  AND status='SAFE'
                LIMIT 1
                """,
                (raw,)
            ).fetchone()

            if r and r[0]:
                return str(r[0]).strip()

            return None
        finally:
            c.close()

    except Exception:
        return None

def envdb(key, default):
    return Path(os.getenv(f"UZBI_DB_{key.upper()}", str(DATA/default)))

PLATFORMS = {
 "auction":{"title":"xarid.uzex.uz — Аукцион савдоси","kind":"auction","db":envdb("auction","auction_completed_deals.sqlite3"),"icon":"A","logo":"/static/platform-icons/auction.png","source":"xarid.uzex.uz","trade_type":"Аукцион савдоси"},
 "shop":{"title":"xarid.uzex.uz — Электрон дўкон","kind":"completed","db":envdb("shop","shop_completed_deals.sqlite3"),"icon":"S","logo":"/static/platform-icons/shop.png","source":"xarid.uzex.uz","trade_type":"Электрон дўкон"},
 "national":{"title":"xarid.uzex.uz — Миллий дўкон","kind":"completed","db":envdb("national","national_completed_deals.sqlite3"),"icon":"UZ","logo":"/static/platform-icons/national.png","source":"xarid.uzex.uz","trade_type":"Миллий дўкон"},

 "xt_auction":{"title":"xt-xarid.uz — Аукцион савдоси","kind":"xt_auction_completed","db":envdb("xt_auction","xt_auction_completed.sqlite3"),"icon":"A","logo":"/static/platform-icons/xt.png","source":"xt-xarid.uz","trade_type":"Аукцион савдоси"},
 "xt_national":{"title":"xt-xarid.uz — Миллий дўкон","kind":"xt_national_completed","db":envdb("xt_national","xt_national_completed.sqlite3"),"icon":"UZ","logo":"/static/platform-icons/xt.png","source":"xt-xarid.uz","trade_type":"Миллий дўкон"},

 "etender":{"title":"etender.uzex.uz — Тендер савдолари","kind":"etender","db":envdb("etender","etender_completed_deals.sqlite3"),"icon":"T","logo":"/static/platform-icons/etender.png","source":"etender.uzex.uz","trade_type":"Тендер савдолари"},
 "xt":{"title":"xt-xarid.uz — Тендер савдолари","kind":"xt","db":envdb("xt","xt_xarid_tenders.sqlite3"),"icon":"T","logo":"/static/platform-icons/xt.png","source":"xt-xarid.uz","trade_type":"Тендер савдолари"},
 "shaffof":{"title":"Shaffof Xarid — Тендер савдолари","kind":"shaffof","db":envdb("shaffof","shaffof_tenders.sqlite3"),"icon":"T","logo":"/static/platform-icons/shaffof.png","source":"Shaffof Xarid","trade_type":"Тендер савдолари"},
 "mc":{"title":"MC.UZ — Тендер савдолари","kind":"mc","db":envdb("mc","mc_tenders.sqlite3"),"icon":"T","logo":"/static/platform-icons/mc.png","source":"MC.UZ","trade_type":"Тендер савдолари"},
}

PENDING = {
 "shop":{"title":"xarid.uzex.uz ? ???????? ????? ???? ??????","db":envdb("shop_pending","shop_pending_lots.sqlite3"),"kind":"shop_pending","logo":"/static/platform-icons/shop.png","source":"xarid.uzex.uz","trade_type":"???????? ?????"},
 "national":{"title":"xarid.uzex.uz ? ?????? ????? ???? ??????","db":envdb("national_pending","national_pending_lots.sqlite3"),"kind":"national_pending","logo":"/static/platform-icons/shop.png","source":"xarid.uzex.uz","trade_type":"?????? ?????"},
 "auction":{"title":"xarid.uzex.uz ? ??????? ??????? ???? ??????","db":envdb("auction_pending","auction_pending_lots.sqlite3"),"kind":"auction_pending","logo":"/static/platform-icons/auction.png","source":"xarid.uzex.uz","trade_type":"??????? ???????"},
 "xt_auction":{"title":"xt-xarid.uz ? ??????? ??????? ???? ??????","db":envdb("xt_auction_pending","xt_auction_pending_lots.sqlite3"),"kind":"xt_auction_pending","logo":"/static/platform-icons/xt.png","source":"xt-xarid.uz","trade_type":"??????? ???????"},
 "xt_national":{"title":"xt-xarid.uz ? ?????? ????? ???? ??????","db":envdb("xt_national_pending","xt_national_pending.sqlite3"),"kind":"xt_national_pending","logo":"/static/platform-icons/xt.png","source":"xt-xarid.uz","trade_type":"?????? ?????"},

 "etender":{"title":"etender.uzex.uz ? ???? ?????????","db":envdb("etender_pending","etender_pending_lots.sqlite3"),"kind":"etender_pending","logo":"/static/platform-icons/etender.png","source":"etender.uzex.uz","trade_type":"?????? ?????????"},
 "xt":{"title":"xt-xarid.uz ? ???? ?????????","db":envdb("xt","xt_xarid_tenders.sqlite3"),"kind":"xt","logo":"/static/platform-icons/xt.png","source":"xt-xarid.uz","trade_type":"?????? ?????????"},
 "shaffof":{"title":"Shaffof Xarid ? ???? ?????????","db":envdb("shaffof","shaffof_tenders.sqlite3"),"kind":"shaffof_pending","logo":"/static/platform-icons/shaffof.png","source":"Shaffof Xarid","trade_type":"?????? ?????????"},
 "mc":{"title":"MC.UZ ? ???? ?????????","db":envdb("mc","mc_tenders.sqlite3"),"kind":"mc","logo":"/static/platform-icons/mc.png","source":"MC.UZ","trade_type":"?????? ?????????"},
}

APP_DB = Path(os.getenv("UZBI_APP_DB", str(ROOT / "data" / "uzbi_web_app.sqlite3")))
SESSION_HOURS = int(os.getenv("UZBI_SESSION_HOURS","24"))
COOKIE_SECURE = os.getenv("UZBI_COOKIE_SECURE","0").strip().lower() in {"1","true","yes"}

def appdb():
    APP_DB.parent.mkdir(parents=True, exist_ok=True)
    c=sqlite3.connect(APP_DB,timeout=30,check_same_thread=False)
    c.row_factory=sqlite3.Row
    c.execute("PRAGMA journal_mode=WAL")
    c.execute("PRAGMA busy_timeout=30000")
    return c

def hash_password(password:str, salt:bytes|None=None):
    salt=salt or secrets.token_bytes(16)
    dk=hashlib.pbkdf2_hmac("sha256",password.encode("utf-8"),salt,260000)
    return salt.hex()+"$"+dk.hex()

def verify_password(password:str, stored:str):
    try:
        salt_hex,digest=stored.split("$",1)
        test=hash_password(password,bytes.fromhex(salt_hex)).split("$",1)[1]
        return hmac.compare_digest(test,digest)
    except Exception:
        return False

def init_app_db():
    with appdb() as c:
        c.executescript("""
        CREATE TABLE IF NOT EXISTS users(
          id INTEGER PRIMARY KEY AUTOINCREMENT,
          username TEXT UNIQUE NOT NULL,
          password_hash TEXT NOT NULL,
          full_name TEXT,
          role TEXT NOT NULL CHECK(role IN ('admin','analyst','viewer')),
          is_active INTEGER NOT NULL DEFAULT 1,
          created_at TEXT NOT NULL,
          last_login_at TEXT
        );
        CREATE TABLE IF NOT EXISTS sessions(
          token TEXT PRIMARY KEY,
          user_id INTEGER NOT NULL,
          created_at TEXT NOT NULL,
          expires_at TEXT NOT NULL,
          ip TEXT,
          user_agent TEXT,
          FOREIGN KEY(user_id) REFERENCES users(id) ON DELETE CASCADE
        );
        CREATE INDEX IF NOT EXISTS ix_sessions_expires ON sessions(expires_at);
        """)
        admin=c.execute("SELECT id FROM users WHERE username='admin'").fetchone()
        if not admin:
            pwd=os.getenv("UZBI_ADMIN_PASSWORD", "").strip()
            if not pwd:
                raise RuntimeError("Set UZBI_ADMIN_PASSWORD before the first start to create the admin account.")
            c.execute("INSERT INTO users(username,password_hash,full_name,role,is_active,created_at) VALUES(?,?,?,?,1,?)",
                      ("admin",hash_password(pwd),"BIData Administrator","admin",datetime.now(timezone.utc).isoformat()))
        c.commit()

init_app_db()

def current_user(request:Request):
    token=request.cookies.get("uzbi_session")
    if not token:
        raise HTTPException(401,"LOGIN_REQUIRED")
    now=datetime.now(timezone.utc).isoformat()
    with appdb() as c:
        r=c.execute("""SELECT u.id,u.username,u.full_name,u.role,u.is_active,s.expires_at
                     FROM sessions s JOIN users u ON u.id=s.user_id
                     WHERE s.token=? AND s.expires_at>?""",(token,now)).fetchone()
    if not r or not r["is_active"]:
        raise HTTPException(401,"LOGIN_REQUIRED")
    return dict(r)

def require_role(*roles):
    def dep(u=Depends(current_user)):
        if u["role"] not in roles:
            raise HTTPException(403,"PERMISSION_DENIED")
        return u
    return dep

app=FastAPI(title="BIData WEB",version="10.7-product-5plus3")
app.mount("/static",StaticFiles(directory=str(STATIC)),name="static")

@app.middleware("http")
async def _utf8_headers(request: Request, call_next):
    response = await call_next(request)
    ct = response.headers.get("content-type", "")
    if ct.startswith("application/json") and "charset=" not in ct.lower():
        response.headers["content-type"] = "application/json; charset=utf-8"
    response.headers["X-UZBI-Version"] = "10.7-product-5plus3"
    return response



def ro(path:Path):
    if not path.exists(): raise HTTPException(503,f"DB Р В Р Р‹Р Р†Р вЂљРЎв„ўР В Р’В Р РЋРІР‚СћР В Р’В Р РЋРІР‚вЂќР В Р’В Р РЋРІР‚ВР В Р’В Р вЂ™Р’В»Р В Р’В Р РЋР’ВР В Р’В Р вЂ™Р’В°Р В Р’В Р СћРІР‚ВР В Р’В Р РЋРІР‚В: {path}")
    uri=path.resolve().as_uri()+"?mode=ro"
    c=sqlite3.connect(uri,uri=True,timeout=60,check_same_thread=False)
    c.create_function("CASEFOLD", 1, lambda s: s.casefold() if s is not None else "")
    c.row_factory=sqlite3.Row
    c.execute("PRAGMA query_only=ON")
    c.execute("PRAGMA busy_timeout=60000")
    c.execute("PRAGMA temp_store=MEMORY")
    c.execute("PRAGMA cache_size=-65536")
    return c

def ident(s): return '"'+str(s).replace('"','""')+'"'
def tables(c): return [r[0] for r in c.execute("select name from sqlite_master where type='table' and name not like 'sqlite_%' order by name")]
def cols(c,t): return [r["name"] for r in c.execute(f"pragma table_info({ident(t)})")]
ALIASES={
 "id":["deal_id","proc_id","id","tender_id","lot_id"],
 "date":["deal_date","confirmed_at","created_at","publicated_at","published_at","placement_date","start_date"],
 "close":["close_at","end_date","deadline","placement_end_date"],
 "product":["product_name","name","title","category_name","subject"],
 "category":["category_name","category","type_name"],
 "amount":["deal_cost","total_amount","winner_amount","contract_price","price","sum","start_amount","start_price"],
 "quantity":["amount","quantity","qty","count"],
 "customer":["customer_name","customer","buyer_name","organization_name","org_name"],
 "customer_inn":["customer_inn","customer_tin","buyer_inn","customer_stir"],
 "supplier":["provider_name","supplier_name","winner_name","provider"],
 "supplier_inn":["provider_inn","provider_tin","supplier_inn","winner_inn"],
 "region":["customer_region_name","region_name_uz","region_name","customer_region","region"],
 "status":["deal_status_name","status_name","status","state"],
 "display":["lot_display_no","lot_number","number","proc_number"],
 "unit":["unit_name","unit","measure_name"],
}
def detect(cs,key):
    low={x.lower():x for x in cs}
    for a in ALIASES[key]:
        if a.lower() in low:return low[a.lower()]
    return None
def best_table(c):
    ts=tables(c)
    preferred=["completed_deals","tenders","pending_lots","lots","deals"]
    for p in preferred:
        if p in ts:return p
    scored=[]
    for t in ts:
        try:
            cc=cols(c,t); score=sum(bool(detect(cc,k)) for k in ("id","date","amount","status","customer","product"))
            n=c.execute(f"select count(*) from {ident(t)}").fetchone()[0]
            scored.append((score,n,t))
        except: pass
    return max(scored,default=(0,0,None))[2]

def platform(key):
    if key not in PLATFORMS: raise HTTPException(404,"Р В Р’В Р РЋРЎС™Р В Р’В Р РЋРІР‚СћР В Р’В Р РЋР’ВР В Р’В Р вЂ™Р’В°Р В Р Р‹Р В РІР‚В°Р В Р’В Р вЂ™Р’В»Р В Р Р‹Р РЋРІР‚СљР В Р’В Р РЋР’В Р В Р’В Р РЋРІР‚вЂќР В Р’В Р вЂ™Р’В»Р В Р’В Р вЂ™Р’В°Р В Р Р‹Р Р†Р вЂљРЎв„ўР В Р Р‹Р Р†Р вЂљРЎвЂєР В Р’В Р РЋРІР‚СћР В Р Р‹Р В РІР‚С™Р В Р’В Р РЋР’ВР В Р’В Р вЂ™Р’В°")
    return PLATFORMS[key]

def completed_filter(date_from=None,date_to=None,region=None,q=None,key=None):
    w=["1=1"];p=[]
    if date_from:w+=["deal_date>=?"];p+=[date_from]
    if date_to:w+=["deal_date<date(?,'+1 day')"];p+=[date_to]
    if region:w+=["customer_region_name=?"];p+=[region]

    if q:
        q_clean=q.strip()

        if key=="shop":
            q_is_tin=(
                q_clean.isdigit()
                and len(q_clean)>=8
            )

            if q_is_tin:
                w+=["(customer_inn=? or provider_inn=?)"]
                p += [q_clean,q_clean]
            else:
                product_names=_resolve_shop_product_names(
                    q_clean,
                    limit=50
                )

                if product_names:
                    placeholders=",".join(
                        "?" for _ in product_names
                    )
                    w += [
                        f"product_name in ({placeholders})"
                    ]
                    p += product_names
                else:
                    # Valid text search with no FTS match:
                    # return no rows without scanning the 18+ GB DB.
                    w += ["1=0"]
        else:
            like=f"%{q_clean}%"
            w+=["(product_name like ? or category_name like ? or lot_display_no like ? or customer_name like ? or customer_inn like ? or provider_name like ? or provider_inn like ?)"]
            p += [like]*7

    return " and ".join(w),p


@lru_cache(maxsize=32)
def cached_meta(key: str):
    """Metadata cache: source DB remains read-only; result lives only in process memory."""
    v=platform(key)
    with ro(v["db"]) as c:
        if v["kind"]=="completed":
            # Avoid expensive DISTINCT region on every page open.
            # total_count exists in SHOP; fallback to COUNT only when needed.
            cs=cols(c,"completed_deals")
            if "total_count" in cs:
                rr=c.execute("select max(total_count), min(deal_date), max(deal_date) from completed_deals").fetchone()
                row_count=int(rr[0] or 0)
                min_date=(rr[1] or "")[:10]; max_date=(rr[2] or "")[:10]
            else:
                rr=c.execute("select count(*), min(deal_date), max(deal_date) from completed_deals").fetchone()
                row_count=int(rr[0] or 0)
                min_date=(rr[1] or "")[:10]; max_date=(rr[2] or "")[:10]
            return {"mode":"REAL_READ_ONLY","min_date":min_date,"max_date":max_date,"rows":row_count,"regions":[],"regions_lazy":True}
        if v["kind"]=="auction":
            rr=c.execute("select count(*),min(deal_date),max(deal_date) from completed_deals").fetchone()
            return {"mode":"AUCTION_SPECIAL","min_date":(rr[1] or "")[:10],"max_date":(rr[2] or "")[:10],"rows":rr[0],"regions":[],"regions_lazy":True}
        if v["kind"]=="etender":
            rr=c.execute("select count(*),min(deal_date),max(deal_date) from completed_deals").fetchone()
            return {"mode":"ETENDER_SPECIAL","min_date":(rr[1] or "")[:10],"max_date":(rr[2] or "")[:10],"rows":rr[0],"regions":[]}
        if v["kind"]=="shaffof":
            rr=c.execute("select count(*),min(created_at),max(created_at) from tenders").fetchone()
            return {"mode":"SHAFFOF_SPECIAL","min_date":(rr[1] or "")[:10],"max_date":(rr[2] or "")[:10],"rows":rr[0],"regions":[],"regions_lazy":True}
        if v["kind"]=="xt_auction_completed":
            rr=c.execute("""
                select
                    count(distinct proc_id),
                    min(contract_close_date),
                    max(contract_close_date)
                from contracts
                where status in ('execution_published','format_payment_success')
                  and proc_id is not null
            """).fetchone()
            return {
                "mode":"XT_AUCTION_SPECIAL",
                "min_date":(rr[1] or "")[:10],
                "max_date":(rr[2] or "")[:10],
                "rows":int(rr[0] or 0),
                "regions":[],
                "regions_lazy":True
            }
        if v["kind"]=="xt":
            rr=c.execute("select count(*),min(publicated_at),max(close_at) from tenders").fetchone()
            return {"mode":"XT_SPECIAL","min_date":(rr[1] or "")[:10],"max_date":(rr[2] or "")[:10],"rows":rr[0],"regions":[]}
        if v["kind"]=="mc":
            rr=c.execute("select count(*),min(confirmed_date),max(confirmed_date) from tenders").fetchone()
            return {"mode":"MC_SPECIAL","min_date":(rr[1] or "")[:10],"max_date":(rr[2] or "")[:10],"rows":rr[0],"regions":[],"regions_lazy":True}
        t=best_table(c)
        return {"mode":"AUTO_READ_ONLY","min_date":"","max_date":"","rows":c.execute(f"select count(*) from {ident(t)}").fetchone()[0],"regions":[]}

@app.get("/api/{key}/regions")
def regions_fast(key:str):
    v=platform(key)
    with ro(v["db"]) as c:
        if v["kind"]=="completed":
            # only one scan, and only when the user opens region dropdown
            return [x[0] for x in c.execute("select distinct customer_region_name from completed_deals where coalesce(customer_region_name,'')<>'' order by 1")]
        if v["kind"]=="xt_auction_completed":
            return [x[0] for x in c.execute("select distinct trim(region_name) from auction_procedures where coalesce(trim(region_name),'')<>'' order by 1")]
        if v["kind"]=="auction":
            return [x[0] for x in c.execute("select distinct region_name from completed_deals where coalesce(region_name,'')<>'' order by 1")]
        if v["kind"]=="shaffof":
            return [x[0] for x in c.execute("select distinct region from tenders where coalesce(region,'')<>'' order by 1")]
        if v["kind"]=="mc":
            return [x[0] for x in c.execute("select distinct coalesce(region_name_uz,region_name) from tenders where coalesce(region_name_uz,region_name,'')<>'' order by 1")]
        return []


@app.get("/")
def index(): return FileResponse(STATIC/"index.html")



# BIDATA PLATFORM CACHE BACKGROUND REFRESH V1
_PLATFORM_CACHE_TTL = 3600
_PLATFORM_CACHE_LOCK = threading.Lock()
_PLATFORM_CACHE_REFRESHING = False

def _calculate_platform_summary():
    out=[]
    for k,v in PLATFORMS.items():
        x={"key":k,"title":v["title"],"kind":v["kind"],"icon":v["icon"],"logo":v.get("logo"),"source":v.get("source"),"trade_type":v.get("trade_type"),"path":str(v["db"]),"exists":v["db"].exists()}
        if v["db"].exists():
            try:
                with ro(v["db"]) as c:
                    t=best_table(c)
                    x["table"]=t
    
                    if k=="auction":
                        x["table"]="completed_deals"
                        x["rows"]=c.execute("""
                            select count(*)
                            from completed_deals
                            where status in (
                                ?,
                                ?,
                                ?
                            )
                        """, (
                            "\u041f\u043e\u0441\u0442\u0430\u0432\u043a\u0430 \u0441\u043e\u0432\u0435\u0440\u0448\u0435\u043d\u0430",
                            "\u041e\u043f\u043b\u0430\u0447\u0435\u043d\u0430",
                            "\u041f\u043e\u0441\u0442\u0430\u0432\u043a\u0430 \u0441\u043e\u0432\u0435\u0440\u0448\u0435\u043d\u0430 (\u0412\u043e\u0441\u0441\u0442\u0430\u043d\u043e\u0432\u043b\u0435\u043d\u0430)"
                        )).fetchone()[0]
    
                    elif k in ("shop","national"):
                        x["table"]="completed_deals"
                        x["rows"]=c.execute("""
                            select count(*)
                            from completed_deals
                            where deal_status_name in (
                                ?,
                                ?,
                                ?
                            )
                        """, (
                            "\u041f\u043e\u0441\u0442\u0430\u0432\u043b\u0435\u043d\u0430",
                            "\u041f\u043e\u0441\u0442\u0430\u0432\u043b\u0435\u043d\u0430 (\u0412\u043e\u0441\u0441\u0442\u0430\u043d\u043e\u0432\u043b\u0435\u043d\u0430)",
                            "\u041e\u043f\u043b\u0430\u0447\u0435\u043d\u0430"
                        )).fetchone()[0]
    
                    elif k=="xt_auction":
                        x["table"]="auction_bi_procedures"
                        x["rows"]=c.execute("""
                            select count(*)
                            from auction_bi_procedures
                            where status_group='success'
                        """).fetchone()[0]
    
                    elif k=="xt_national":
                        x["table"]="completed_lots"
                        x["rows"]=c.execute(
                            "select count(*) from completed_lots"
                        ).fetchone()[0]
    
                    elif k=="etender":
                        x["table"]="completed_deals"
                        x["rows"]=c.execute("""
                            select count(*)
                            from completed_deals
                            where deal_status_name in (?,?)
                        """, (
                            "G\u2018olib tomonidan qabul qilingan (Amalga oshgan)",
                            "Bajarilgan"
                        )).fetchone()[0]
    
                    elif k=="xt":
                        x["table"]="tenders"
                        x["rows"]=c.execute("""
                            select count(*)
                            from tenders
                            where status='close'
                        """).fetchone()[0]
    
                    elif k=="shaffof":
                        x["table"]="tenders"
                        x["rows"]=c.execute("""
                            select count(*)
                            from tenders
                            where status_name='Yakunlandi'
                              and contract_status_name='Shartnoma tasdiqlandi'
                        """).fetchone()[0]
    
                    elif k=="mc":
                        x["table"]="tenders"
                        x["rows"]=c.execute("""
                            select count(*)
                            from tenders
                            where winner_id is not null
                              and winner_determined_date is not null
                              and trim(winner_determined_date)<>''
                        """).fetchone()[0]
    
                    else:
                        x["rows"]=c.execute(
                            f"select count(*) from {ident(t)}"
                        ).fetchone()[0] if t else 0
    
                    x["tables"]=len(tables(c))
            except Exception as e:x["error"]=str(e)
        out.append(x)
    return out

def _write_platform_summary_cache(rows):
    cache_dir = ROOT / "cache"
    cache_dir.mkdir(parents=True, exist_ok=True)
    cache_file = cache_dir / "platforms_summary.json"

    payload = {
        "generated_at": time.time(),
        "platforms": {}
    }

    for x in rows:
        k = x.get("key")
        if not k:
            continue

        payload["platforms"][k] = {
            "rows": int(x.get("rows", 0) or 0),
            "table": x.get("table")
        }

    tmp = cache_file.with_suffix(".json.tmp")
    tmp.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2),
        encoding="utf-8"
    )
    tmp.replace(cache_file)


def _platform_cache_refresh_worker():
    global _PLATFORM_CACHE_REFRESHING

    try:
        rows = _calculate_platform_summary()
        if len(rows) == len(PLATFORMS):
            _write_platform_summary_cache(rows)
    except Exception as e:
        print("[PLATFORM CACHE] refresh error:", e)
    finally:
        with _PLATFORM_CACHE_LOCK:
            _PLATFORM_CACHE_REFRESHING = False


def _start_platform_cache_refresh():
    global _PLATFORM_CACHE_REFRESHING

    with _PLATFORM_CACHE_LOCK:
        if _PLATFORM_CACHE_REFRESHING:
            return False

        _PLATFORM_CACHE_REFRESHING = True

    t = threading.Thread(
        target=_platform_cache_refresh_worker,
        name="bidata-platform-cache-refresh",
        daemon=True
    )
    t.start()
    return True


def _platform_cache_is_stale(cache_file):
    try:
        age = time.time() - cache_file.stat().st_mtime
        return age >= _PLATFORM_CACHE_TTL
    except Exception:
        return True


@app.get("/api/platforms")
def api_platforms():
    # Fast persistent summary cache.
    # The original SQL rules below remain as fallback.
    cache_file = ROOT / "cache" / "platforms_summary.json"

    try:
        if cache_file.exists():
            cached = json.loads(cache_file.read_text(encoding="utf-8"))
            cached_platforms = cached.get("platforms", {})

            if all(k in cached_platforms for k in PLATFORMS):
                if _platform_cache_is_stale(cache_file):
                    _start_platform_cache_refresh()

                out = []

                for k, v in PLATFORMS.items():
                    item = cached_platforms[k]

                    x = {
                        "key": k,
                        "title": v["title"],
                        "kind": v["kind"],
                        "icon": v["icon"],
                        "logo": v.get("logo"),
                        "source": v.get("source"),
                        "trade_type": v.get("trade_type"),
                        "path": str(v["db"]),
                        "exists": v["db"].exists(),
                        "table": item.get("table"),
                        "rows": int(item.get("rows", 0)),
                        "cached": True
                    }

                    if v["db"].exists():
                        try:
                            with ro(v["db"]) as c:
                                x["tables"] = len(tables(c))
                        except Exception:
                            x["tables"] = 0
                    else:
                        x["tables"] = 0

                    out.append(x)

                return out
    except Exception:
        pass

    # No usable cache: calculate synchronously once as the safe fallback.
    # Future requests will use the persistent cache.
    out=[]
    for k,v in PLATFORMS.items():
        x={"key":k,"title":v["title"],"kind":v["kind"],"icon":v["icon"],"logo":v.get("logo"),"source":v.get("source"),"trade_type":v.get("trade_type"),"path":str(v["db"]),"exists":v["db"].exists()}
        if v["db"].exists():
            try:
                with ro(v["db"]) as c:
                    t=best_table(c)
                    x["table"]=t

                    if k=="auction":
                        x["table"]="completed_deals"
                        x["rows"]=c.execute("""
                            select count(*)
                            from completed_deals
                            where status in (
                                ?,
                                ?,
                                ?
                            )
                        """, (
                            "\u041f\u043e\u0441\u0442\u0430\u0432\u043a\u0430 \u0441\u043e\u0432\u0435\u0440\u0448\u0435\u043d\u0430",
                            "\u041e\u043f\u043b\u0430\u0447\u0435\u043d\u0430",
                            "\u041f\u043e\u0441\u0442\u0430\u0432\u043a\u0430 \u0441\u043e\u0432\u0435\u0440\u0448\u0435\u043d\u0430 (\u0412\u043e\u0441\u0441\u0442\u0430\u043d\u043e\u0432\u043b\u0435\u043d\u0430)"
                        )).fetchone()[0]

                    elif k in ("shop","national"):
                        x["table"]="completed_deals"
                        x["rows"]=c.execute("""
                            select count(*)
                            from completed_deals
                            where deal_status_name in (
                                ?,
                                ?,
                                ?
                            )
                        """, (
                            "\u041f\u043e\u0441\u0442\u0430\u0432\u043b\u0435\u043d\u0430",
                            "\u041f\u043e\u0441\u0442\u0430\u0432\u043b\u0435\u043d\u0430 (\u0412\u043e\u0441\u0441\u0442\u0430\u043d\u043e\u0432\u043b\u0435\u043d\u0430)",
                            "\u041e\u043f\u043b\u0430\u0447\u0435\u043d\u0430"
                        )).fetchone()[0]

                    elif k=="xt_auction":
                        x["table"]="auction_bi_procedures"
                        x["rows"]=c.execute("""
                            select count(*)
                            from auction_bi_procedures
                            where status_group='success'
                        """).fetchone()[0]

                    elif k=="xt_national":
                        x["table"]="completed_lots"
                        x["rows"]=c.execute(
                            "select count(*) from completed_lots"
                        ).fetchone()[0]

                    elif k=="etender":
                        x["table"]="completed_deals"
                        x["rows"]=c.execute("""
                            select count(*)
                            from completed_deals
                            where deal_status_name in (?,?)
                        """, (
                            "G\u2018olib tomonidan qabul qilingan (Amalga oshgan)",
                            "Bajarilgan"
                        )).fetchone()[0]

                    elif k=="xt":
                        x["table"]="tenders"
                        x["rows"]=c.execute("""
                            select count(*)
                            from tenders
                            where status='close'
                        """).fetchone()[0]

                    elif k=="shaffof":
                        x["table"]="tenders"
                        x["rows"]=c.execute("""
                            select count(*)
                            from tenders
                            where status_name='Yakunlandi'
                              and contract_status_name='Shartnoma tasdiqlandi'
                        """).fetchone()[0]

                    elif k=="mc":
                        x["table"]="tenders"
                        x["rows"]=c.execute("""
                            select count(*)
                            from tenders
                            where winner_id is not null
                              and winner_determined_date is not null
                              and trim(winner_determined_date)<>''
                        """).fetchone()[0]

                    else:
                        x["rows"]=c.execute(
                            f"select count(*) from {ident(t)}"
                        ).fetchone()[0] if t else 0

                    x["tables"]=len(tables(c))
            except Exception as e:x["error"]=str(e)
        out.append(x)
    return out

@app.get("/api/system/health")
def system_health():
    start=time.time(); arr=api_platforms()
    return {"status":"ok","version":"10.3","uzbi_root":str(UZBI_ROOT),"online":sum(1 for x in arr if x["exists"]),"total":len(arr),"elapsed_ms":round((time.time()-start)*1000),"platforms":arr}

@app.get("/api/{key}/meta")
def meta(key:str):
    platform(key)
    return cached_meta(key)

def end_exclusive(end_date):
    return (datetime.strptime(end_date, "%Y-%m-%d") + timedelta(days=1)).strftime("%Y-%m-%d")
@app.get("/api/{key}/summary")
def summary(key:str,date_from:Optional[str]=None,date_to:Optional[str]=None,region:Optional[str]=None,q:Optional[str]=None,amount_min:Optional[str]=None,amount_max:Optional[str]=None):
    amount_min,amount_max=_market_amount_bounds(key,amount_min,amount_max)
    v=platform(key)
    with ro(v["db"]) as c:
        if v["kind"]=="completed":
            w,p=completed_filter(date_from,date_to,region,q,key)

            if key in ("shop","national"):
                success=(
                    "\u041f\u043e\u0441\u0442\u0430\u0432\u043b\u0435\u043d\u0430",
                    "\u041f\u043e\u0441\u0442\u0430\u0432\u043b\u0435\u043d\u0430 (\u0412\u043e\u0441\u0441\u0442\u0430\u043d\u043e\u0432\u043b\u0435\u043d\u0430)",
                    "\u041e\u043f\u043b\u0430\u0447\u0435\u043d\u0430",
                )
                w=f"({w}) and deal_status_name in (?,?,?)"
                p=list(p)+list(success)

            r=_market_amount_execute(c,key,amount_min,amount_max,date_from,date_to,region,f"""select count(*) deals,coalesce(sum(deal_cost),0) amount,coalesce(avg(deal_cost),0) avg_deal,
            count(distinct product_name) products,count(distinct customer_inn) customers,count(distinct provider_inn) providers
            from completed_deals where {w}""",p,search=q).fetchone(); return dict(r)
        if v["kind"]=="auction":
            success=(
                "\u041f\u043e\u0441\u0442\u0430\u0432\u043a\u0430 \u0441\u043e\u0432\u0435\u0440\u0448\u0435\u043d\u0430",
                "\u041e\u043f\u043b\u0430\u0447\u0435\u043d\u0430",
                "\u041f\u043e\u0441\u0442\u0430\u0432\u043a\u0430 \u0441\u043e\u0432\u0435\u0440\u0448\u0435\u043d\u0430 (\u0412\u043e\u0441\u0441\u0442\u0430\u043d\u043e\u0432\u043b\u0435\u043d\u0430)",
            )
            where=["d.status in (?,?,?)"]
            params=list(success)
            if date_from:
                where.append("substr(d.deal_date,1,10)>=?")
                params.append(date_from[:10])
            if date_to:
                where.append("substr(d.deal_date,1,10)<=?")
                params.append(date_to[:10])
            if region:
                where.append("d.region_name=?")
                params.append(region)
            if q:
                q_clean=q.strip()
                q_digits="".join(ch for ch in q_clean if ch.isdigit())
                q_is_tin=q_clean.isdigit() and len(q_digits)>=8
                qq=f"%{q_clean}%"

                if q_is_tin:
                    where.append("""(
                        coalesce(d.customer_tin,'') like ? or
                        coalesce(d.provider_tin,'') like ?
                    )""")
                    params.extend([qq,qq])
                else:
                    where.append("""exists (
                        select 1
                        from auction_products aq
                        where aq.completed_deal_id=d.id
                          and coalesce(aq.product_name,'') like ?
                    )""")
                    params.append(qq)

            w=" and ".join(where)
            r=_market_amount_execute(c,key,amount_min,amount_max,date_from,date_to,region,f"""select count(*) deals,
            coalesce(sum(d.contract_price),0) amount,
            coalesce(avg(d.contract_price),0) avg_deal,
            count(distinct d.customer_tin) customers,
            count(distinct d.provider_tin) providers
            from completed_deals d where {w}""",params,search=q).fetchone()
            d=dict(r)
            if "auction_products" in tables(c):
                d["products"]=_market_amount_execute(c,key,amount_min,amount_max,date_from,date_to,region,f"""select count(distinct p.product_name)
                    from auction_products p
                    join completed_deals d on d.id=p.completed_deal_id
                    where {w} and coalesce(p.product_name,'')<>''""",params,search=q).fetchone()[0]
            else:
                d["products"]=0
            return d
        if v["kind"]=="xt_auction_completed":
            # XT Auction completed analytics:
            # successful contracts define completed trade procedures.
            success_statuses=("execution_published","format_payment_success")
            where=[
                "ct.status in (?,?)",
                "ct.proc_id is not null"
            ]
            params=[*success_statuses]

            if date_from:
                where.append("ct.contract_close_date>=?")
                params.append(date_from)
            if date_to:
                where.append("ct.contract_close_date<?")
                params.append(end_exclusive(date_to))

            if region:
                where.append("trim(ap.region_name)=?")
                params.append(region)

            if q:
                q_clean=q.strip()
                q_digits="".join(ch for ch in q_clean if ch.isdigit())
                q_is_tin=q_clean.isdigit() and len(q_digits)>=8
                qq=f"%{q_clean}%"

                if q_is_tin:
                    where.append("""(
                        coalesce(ct.org_inn,'') like ? or
                        coalesce(ct.winner_inn,'') like ?
                    )""")
                    params.extend([qq,qq])
                else:
                    where.append("""exists (
                        select 1
                        from auction_items qi
                        where qi.proc_id=ct.proc_id
                          and CASEFOLD(coalesce(qi.product_name,'')) like CASEFOLD(?)
                    )""")
                    params.append(qq)

            w=" and ".join(where)

            r=_market_amount_execute(c,key,amount_min,amount_max,date_from,date_to,region,f"""
                select
                    count(distinct ct.proc_id) deals,
                    coalesce(sum(ct.contract_totalcost),0) amount,
                    count(distinct nullif(trim(ct.org_inn),'')) customers,
                    count(distinct nullif(trim(ct.winner_inn),'')) providers
                from contracts ct
                left join auction_procedures ap
                  on ap.proc_id=ct.proc_id
                where {w}
            """,params,search=q).fetchone()

            d=dict(r)
            d["avg_deal"]=(float(d["amount"] or 0)/int(d["deals"])) if d["deals"] else 0

            d["products"]=_market_amount_execute(c,key,amount_min,amount_max,date_from,date_to,region,f"""
                select count(distinct nullif(trim(i.product_name),''))
                from auction_items i
                where i.proc_id in (
                    select distinct ct.proc_id
                    from contracts ct
                    left join auction_procedures ap
                      on ap.proc_id=ct.proc_id
                    where {w}
                )
            """,params,search=q).fetchone()[0]

            return d
        if v["kind"]=="xt":
            r=_market_amount_execute(c,key,amount_min,amount_max,date_from,date_to,region,"""select count(*) deals,coalesce(sum(total_amount),0) amount,coalesce(avg(total_amount),0) avg_deal,
            count(distinct customer_company_id) customers,0 providers from tenders""",search=q).fetchone()
            d=dict(r); d["products"]=_market_amount_execute(c,key,amount_min,amount_max,date_from,date_to,region,"select count(distinct product_name) from items",search=q).fetchone()[0] if "items" in tables(c) else 0; return d
        if v["kind"]=="etender":
            r=_market_amount_execute(c,key,amount_min,amount_max,date_from,date_to,region,"""select count(*) deals,coalesce(sum(deal_cost),0) amount,coalesce(avg(deal_cost),0) avg_deal,
            count(distinct category_name) products,count(distinct customer_inn) customers,count(distinct provider_inn) providers
            from completed_deals""",search=q).fetchone(); return dict(r)
        if v["kind"]=="shaffof":
            r=_market_amount_execute(c,key,amount_min,amount_max,date_from,date_to,region,"""select count(*) deals,coalesce(sum(coalesce(winner_offer,start_price)),0) amount,
            coalesce(avg(coalesce(winner_offer,start_price)),0) avg_deal,
            count(distinct coalesce(service_type,object_type,lot_type_name)) products,
            count(distinct customer_inn) customers,count(distinct winner_inn) providers from tenders""",search=q).fetchone(); return dict(r)
        if v["kind"]=="mc":
            r=_market_amount_execute(c,key,amount_min,amount_max,date_from,date_to,region,"""select count(*) deals,coalesce(sum(start_price),0) amount,coalesce(avg(start_price),0) avg_deal,
            count(distinct name) products,count(distinct customer_inn) customers,0 providers from tenders""",search=q).fetchone()
            d=dict(r)
            if "tender_items" in tables(c): d["products"]=_market_amount_execute(c,key,amount_min,amount_max,date_from,date_to,region,"select count(distinct product_name) from tender_items",search=q).fetchone()[0]
            return d
        t=best_table(c); cs=cols(c,t)
        ac,pc,cc,sc=detect(cs,"amount"),detect(cs,"product"),detect(cs,"customer"),detect(cs,"supplier")
        sel=["count(*) deals",f"coalesce(sum({ident(ac)}),0) amount" if ac else "0 amount",f"coalesce(avg({ident(ac)}),0) avg_deal" if ac else "0 avg_deal",
             f"count(distinct {ident(pc)}) products" if pc else "0 products",f"count(distinct {ident(cc)}) customers" if cc else "0 customers",f"count(distinct {ident(sc)}) providers" if sc else "0 providers"]
        return dict(_market_amount_execute(c,key,amount_min,amount_max,date_from,date_to,region,f"select {','.join(sel)} from {ident(t)}",search=q).fetchone())

@app.get("/api/{key}/monthly")
def monthly(key:str,date_from:Optional[str]=None,date_to:Optional[str]=None,region:Optional[str]=None,q:Optional[str]=None,amount_min:Optional[str]=None,amount_max:Optional[str]=None):
    amount_min,amount_max=_market_amount_bounds(key,amount_min,amount_max)
    v=platform(key)
    with ro(v["db"]) as c:
        if v["kind"]=="completed":
            w,p=completed_filter(date_from,date_to,region,q,key)

            if key in ("shop","national"):
                success=(
                    "\u041f\u043e\u0441\u0442\u0430\u0432\u043b\u0435\u043d\u0430",
                    "\u041f\u043e\u0441\u0442\u0430\u0432\u043b\u0435\u043d\u0430 (\u0412\u043e\u0441\u0441\u0442\u0430\u043d\u043e\u0432\u043b\u0435\u043d\u0430)",
                    "\u041e\u043f\u043b\u0430\u0447\u0435\u043d\u0430",
                )
                w=f"({w}) and deal_status_name in (?,?,?)"
                p=list(p)+list(success)

            return [dict(x) for x in _market_amount_execute(c,key,amount_min,amount_max,date_from,date_to,region,f"select substr(deal_date,1,7) period,count(*) deals,coalesce(sum(deal_cost),0) amount from completed_deals where {w} group by 1 order by 1",p,search=q)]
        if v["kind"]=="etender":
            return [dict(x) for x in _market_amount_execute(c,key,amount_min,amount_max,date_from,date_to,region,"select substr(deal_date,1,7) period,count(*) deals,coalesce(sum(deal_cost),0) amount from completed_deals where deal_date is not null group by 1 order by 1",search=q)]
        if v["kind"]=="shaffof":
            return [dict(x) for x in _market_amount_execute(c,key,amount_min,amount_max,date_from,date_to,region,"select substr(created_at,1,7) period,count(*) deals,coalesce(sum(coalesce(winner_offer,start_price)),0) amount from tenders where created_at is not null group by 1 order by 1",search=q)]
        if v["kind"]=="mc":
            return [dict(x) for x in _market_amount_execute(c,key,amount_min,amount_max,date_from,date_to,region,"select substr(confirmed_date,1,7) period,count(*) deals,coalesce(sum(start_price),0) amount from tenders where confirmed_date is not null group by 1 order by 1",search=q)]
        if v["kind"]=="auction":
            success=(
                "\u041f\u043e\u0441\u0442\u0430\u0432\u043a\u0430 \u0441\u043e\u0432\u0435\u0440\u0448\u0435\u043d\u0430",
                "\u041e\u043f\u043b\u0430\u0447\u0435\u043d\u0430",
                "\u041f\u043e\u0441\u0442\u0430\u0432\u043a\u0430 \u0441\u043e\u0432\u0435\u0440\u0448\u0435\u043d\u0430 (\u0412\u043e\u0441\u0441\u0442\u0430\u043d\u043e\u0432\u043b\u0435\u043d\u0430)",
            )
            where=["status in (?,?,?)","deal_date is not null"]
            params=list(success)
            if date_from:
                where.append("substr(deal_date,1,10)>=?")
                params.append(date_from[:10])
            if date_to:
                where.append("substr(deal_date,1,10)<=?")
                params.append(date_to[:10])
            if region:
                where.append("region_name=?")
                params.append(region)
            if q:
                q_clean=q.strip()
                q_digits="".join(ch for ch in q_clean if ch.isdigit())
                q_is_tin=q_clean.isdigit() and len(q_digits)>=8
                qq=f"%{q_clean}%"

                if q_is_tin:
                    where.append("""(
                        coalesce(customer_tin,'') like ? or
                        coalesce(provider_tin,'') like ?
                    )""")
                    params.extend([qq,qq])
                else:
                    where.append("""exists (
                        select 1
                        from auction_products aq
                        where aq.completed_deal_id=completed_deals.id
                          and coalesce(aq.product_name,'') like ?
                    )""")
                    params.append(qq)

            w=" and ".join(where)
            return [dict(x) for x in _market_amount_execute(c,key,amount_min,amount_max,date_from,date_to,region,f"""
                select substr(deal_date,1,7) period,
                       count(*) deals,
                       coalesce(sum(contract_price),0) amount
                from completed_deals
                where {w}
                group by 1
                order by 1
            """,params,search=q)]

        if v["kind"]=="xt_auction_completed":
            where=[
                "ct.status in ('execution_published','format_payment_success')",
                "ct.proc_id is not null",
                "ct.contract_close_date is not null"
            ]
            params=[]

            if date_from:
                where.append("ct.contract_close_date>=?")
                params.append(date_from)

            if date_to:
                where.append("ct.contract_close_date<?")
                params.append(end_exclusive(date_to))

            if region:
                where.append("trim(ap.region_name)=?")
                params.append(region)

            if q:
                q_clean=q.strip()
                q_digits="".join(ch for ch in q_clean if ch.isdigit())
                q_is_tin=q_clean.isdigit() and len(q_digits)>=8
                qq=f"%{q_clean}%"

                if q_is_tin:
                    where.append("""(
                        coalesce(ct.org_inn,'') like ? or
                        coalesce(ct.winner_inn,'') like ?
                    )""")
                    params.extend([qq,qq])
                else:
                    where.append("""exists (
                        select 1
                        from auction_items qi
                        where qi.proc_id=ct.proc_id
                          and CASEFOLD(coalesce(qi.product_name,'')) like CASEFOLD(?)
                    )""")
                    params.append(qq)

            w=" and ".join(where)

            return [dict(x) for x in _market_amount_execute(c,key,amount_min,amount_max,date_from,date_to,region,f"""
                with success_contracts as (
                    select
                        ct.proc_id,
                        min(ct.contract_close_date) proc_date,
                        sum(coalesce(ct.contract_totalcost,0)) proc_amount
                    from contracts ct
                    left join auction_procedures ap
                      on ap.proc_id=ct.proc_id
                    where {w}
                    group by ct.proc_id
                )
                select
                    substr(proc_date,1,7) period,
                    count(*) deals,
                    coalesce(sum(proc_amount),0) amount
                from success_contracts
                group by substr(proc_date,1,7)
                order by period
            """,params,search=q)]
        if key=='xt_national' and (amount_min is not None or amount_max is not None):
            sql="SELECT substr(close_date,1,7) period,count(*) deals,sum(total_amount) amount FROM completed_lots GROUP BY 1 ORDER BY 1"
            return [dict(r) for r in _market_amount_execute(c,key,amount_min,amount_max,date_from,date_to,region,sql,search=q)]
        t="tenders" if v["kind"]=="xt" else best_table(c)
        cs=cols(c,t); dc=detect(cs,"date"); ac=detect(cs,"amount")
        if not dc:return []
        return [dict(x) for x in _market_amount_execute(c,key,amount_min,amount_max,date_from,date_to,region,f"select substr({ident(dc)},1,7) period,count(*) deals,{('coalesce(sum('+ident(ac)+'),0)' if ac else '0')} amount from {ident(t)} where {ident(dc)} is not null group by 1 order by 1",search=q)]

def top_auto(c,t,col,amount_col,limit):
    if not col:return []
    val=f"coalesce(sum({ident(amount_col)}),0)" if amount_col else "count(*)"
    return [dict(x) for x in c.execute(f"select {ident(col)} name,count(*) deals,{val} amount from {ident(t)} where coalesce({ident(col)},'')<>'' group by 1 order by amount desc limit ?",(limit,))]

@app.get("/api/{key}/top/{dim}")
def top(key:str,dim:str,limit:int=Query(5,ge=1,le=20),date_from:Optional[str]=None,date_to:Optional[str]=None,region:Optional[str]=None,q:Optional[str]=None,amount_min:Optional[str]=None,amount_max:Optional[str]=None):
    amount_min,amount_max=_market_amount_bounds(key,amount_min,amount_max)
    if dim not in {"products","categories","regions","customers","providers","statuses"}:raise HTTPException(400,"dimension")
    v=platform(key)
    with ro(v["db"]) as c:
        if v["kind"]=="completed":
            mp={"products":"product_name","categories":"category_name","regions":"customer_region_name","customers":"customer_name","providers":"provider_name","statuses":"deal_status_name"}
            col=mp[dim]

            if key=="shop":
                # SHOP BI CACHE V1
                cache_tables={
                    "products":"daily_products",
                    "categories":"daily_categories",
                    "regions":"daily_regions",
                    "customers":"daily_customers",
                    "providers":"daily_providers",
                }
                cache_table=cache_tables.get(dim)

                cache_period_ok=(
                    date_from
                    and date_to
                    and date_from[:4]=="2026"
                    and date_to[:4]=="2026"
                )

                if (
                    cache_table
                    and cache_period_ok
                    and not region
                    and not q
                    and amount_min is None and amount_max is None
                ):
                    cache_db=Path(
                        r"C:\Qudrat\xarid_uz\UZBI\data"
                        r"\shop_bi_cache.sqlite3"
                    )

                    if cache_db.exists():
                        try:
                            with ro(cache_db) as cc:
                                cache_meta=dict(
                                    cc.execute(
                                        "select key,value from cache_meta"
                                    ).fetchall()
                                )

                                data_min=cache_meta.get(
                                    "data_date_min"
                                )
                                data_max=cache_meta.get(
                                    "data_date_max"
                                )

                                full_cache_period=(
                                    data_min
                                    and data_max
                                    and date_from[:10]<=data_min
                                    and date_to[:10]>=data_max
                                )

                                if full_cache_period:
                                    summary_table=(
                                        f"summary_{dim}"
                                    )

                                    if dim=="products":
                                        cache_sql=f"""
                                            select name,
                                                   unit_name,
                                                   deals,
                                                   quantity,
                                                   case
                                                       when quantity>0
                                                       then amount/quantity
                                                       else 0
                                                   end unit_price,
                                                   amount
                                            from {summary_table}
                                            order by amount desc
                                            limit ?
                                        """
                                    else:
                                        cache_sql=f"""
                                            select name,
                                                   deals,
                                                   amount
                                            from {summary_table}
                                            order by amount desc
                                            limit ?
                                        """

                                    cache_params=[limit]
                                else:
                                    if dim=="products":
                                        cache_sql=f"""
                                            select name,
                                                   unit_name,
                                                   coalesce(sum(deals),0) deals,
                                                   coalesce(sum(quantity),0) quantity,
                                                   case
                                                       when coalesce(sum(quantity),0)>0
                                                       then coalesce(sum(amount),0)/sum(quantity)
                                                       else 0
                                                   end unit_price,
                                                   coalesce(sum(amount),0) amount
                                            from {cache_table}
                                            where deal_date>=?
                                              and deal_date<=?
                                            group by name,unit_name
                                            order by amount desc
                                            limit ?
                                        """
                                    else:
                                        cache_sql=f"""
                                            select name,
                                                   coalesce(sum(deals),0) deals,
                                                   coalesce(sum(amount),0) amount
                                            from {cache_table}
                                            where deal_date>=?
                                              and deal_date<=?
                                            group by name
                                            order by amount desc
                                            limit ?
                                        """

                                    cache_params=[
                                        date_from[:10],
                                        date_to[:10],
                                        limit,
                                    ]

                                return [
                                    dict(x)
                                    for x in cc.execute(
                                        cache_sql,
                                        cache_params,
                                    )
                                ]
                        except sqlite3.Error:
                            # Cache failure: use original database.
                            pass

                success=(
                    "\u041f\u043e\u0441\u0442\u0430\u0432\u043b\u0435\u043d\u0430",
                    "\u041f\u043e\u0441\u0442\u0430\u0432\u043b\u0435\u043d\u0430 (\u0412\u043e\u0441\u0441\u0442\u0430\u043d\u043e\u0432\u043b\u0435\u043d\u0430)",
                    "\u041e\u043f\u043b\u0430\u0447\u0435\u043d\u0430",
                )
                where=["deal_status_name in (?,?,?)",f"coalesce({col},'')<>''"]
                params=list(success)

                if date_from:
                    where.append("substr(deal_date,1,10)>=?")
                    params.append(date_from[:10])
                if date_to:
                    where.append("substr(deal_date,1,10)<=?")
                    params.append(date_to[:10])
                if region:
                    where.append("customer_region_name=?")
                    params.append(region)
                if q:
                    q_clean=q.strip()
                    q_is_tin=(
                        q_clean.isdigit()
                        and len(q_clean)>=8
                    )

                    if q_is_tin:
                        where.append("""(
                            customer_inn=? or
                            provider_inn=?
                        )""")
                        params.extend([
                            q_clean,
                            q_clean,
                        ])
                    else:
                        where.append(
                            "coalesce(product_name,'') like ?"
                        )
                        params.append(
                            f"%{q_clean}%"
                        )

                w=" and ".join(where)
                sql=f"""select {col} name,
                               count(*) deals,
                               coalesce(sum(deal_cost),0) amount
                        from completed_deals
                        where {w}
                        group by 1
                        order by amount desc
                        limit ?"""
                return [dict(x) for x in _market_amount_execute(c,key,amount_min,amount_max,date_from,date_to,region,sql,params+[limit])]

            if key=="national":
                success=(
                    "\u041f\u043e\u0441\u0442\u0430\u0432\u043b\u0435\u043d\u0430",
                    "\u041f\u043e\u0441\u0442\u0430\u0432\u043b\u0435\u043d\u0430 (\u0412\u043e\u0441\u0441\u0442\u0430\u043d\u043e\u0432\u043b\u0435\u043d\u0430)",
                    "\u041e\u043f\u043b\u0430\u0447\u0435\u043d\u0430",
                )
                where=[
                    "deal_status_name in (?,?,?)",
                    f"coalesce({col},'')<>''"
                ]
                params=list(success)

                if date_from:
                    where.append("substr(deal_date,1,10)>=?")
                    params.append(date_from[:10])
                if date_to:
                    where.append("substr(deal_date,1,10)<=?")
                    params.append(date_to[:10])
                if region:
                    where.append("customer_region_name=?")
                    params.append(region)
                if q:
                    qq=f"%{q.strip()}%"
                    where.append("""(
                        coalesce(product_name,'') like ? or
                        coalesce(category_name,'') like ? or
                        coalesce(lot_display_no,'') like ? or
                        coalesce(customer_name,'') like ? or
                        coalesce(customer_inn,'') like ? or
                        coalesce(provider_name,'') like ? or
                        coalesce(provider_inn,'') like ?
                    )""")
                    params.extend([qq]*7)

                w=" and ".join(where)
                sql=f"""select {col} name,
                               count(*) deals,
                               coalesce(sum(deal_cost),0) amount
                        from completed_deals
                        where {w}
                        group by 1
                        order by amount desc
                        limit ?"""
                return [
                    dict(x)
                    for x in _market_amount_execute(c,key,amount_min,amount_max,date_from,date_to,region,sql,params+[limit])
                ]

            return [dict(x) for x in _market_amount_execute(c,key,amount_min,amount_max,date_from,date_to,region,f"select {col} name,count(*) deals,coalesce(sum(deal_cost),0) amount from completed_deals where coalesce({col},'')<>'' group by 1 order by amount desc limit ?",(limit,))]
        if v["kind"]=="auction":
            success=(
                "\u041f\u043e\u0441\u0442\u0430\u0432\u043a\u0430 \u0441\u043e\u0432\u0435\u0440\u0448\u0435\u043d\u0430",
                "\u041e\u043f\u043b\u0430\u0447\u0435\u043d\u0430",
                "\u041f\u043e\u0441\u0442\u0430\u0432\u043a\u0430 \u0441\u043e\u0432\u0435\u0440\u0448\u0435\u043d\u0430 (\u0412\u043e\u0441\u0441\u0442\u0430\u043d\u043e\u0432\u043b\u0435\u043d\u0430)",
            )
            where=["d.status in (?,?,?)"]
            params=list(success)

            if date_from:
                where.append("substr(d.deal_date,1,10)>=?")
                params.append(date_from[:10])
            if date_to:
                where.append("substr(d.deal_date,1,10)<=?")
                params.append(date_to[:10])
            if region:
                where.append("d.region_name=?")
                params.append(region)
            if q:
                qq=f"%{q.strip()}%"
                where.append("""(
                    coalesce(d.lot_number,'') like ? or
                    coalesce(d.category_name,'') like ? or
                    coalesce(d.customer_name,'') like ? or
                    coalesce(d.customer_tin,'') like ? or
                    coalesce(d.provider_name,'') like ? or
                    coalesce(d.provider_tin,'') like ?
                )""")
                params.extend([qq,qq,qq,qq,qq,qq])

            w=" and ".join(where)

            if dim=="products" and "auction_products" in tables(c):
                sql=f"""select p.product_name name,
                               coalesce(nullif(trim(p.unit_name),''),'') unit_name,
                               count(distinct d.id) deals,
                               coalesce(sum(p.quantity),0) quantity,
                               coalesce(sum(p.total_amount),0) amount
                        from auction_products p
                        join completed_deals d on d.id=p.completed_deal_id
                        where {w}
                          and coalesce(p.product_name,'')<>''
                        group by p.product_name,
                                 coalesce(nullif(trim(p.unit_name),''),'')
                        order by amount desc
                        limit ?"""
                return [dict(x) for x in _market_amount_execute(c,key,amount_min,amount_max,date_from,date_to,region,sql,params+[limit])]

            if dim=="categories":
                sql=f"""select d.category_name name,
                               count(*) deals,
                               coalesce(sum(d.contract_price),0) amount
                        from completed_deals d
                        where {w}
                          and coalesce(d.category_name,'')<>''
                        group by d.category_name
                        order by amount desc
                        limit ?"""
                return [dict(x) for x in _market_amount_execute(c,key,amount_min,amount_max,date_from,date_to,region,sql,params+[limit])]

            mp={
                "regions":"region_name",
                "customers":"customer_name",
                "providers":"provider_name",
                "statuses":"status",
            }
            col=mp.get(dim)
            if not col:
                return []
            sql=f"""select d.{col} name,
                           count(*) deals,
                           coalesce(sum(d.contract_price),0) amount
                    from completed_deals d
                    where {w}
                      and coalesce(d.{col},'')<>''
                    group by d.{col}
                    order by amount desc
                    limit ?"""
            return [dict(x) for x in _market_amount_execute(c,key,amount_min,amount_max,date_from,date_to,region,sql,params+[limit])]
        if v["kind"]=="xt_auction_completed":
            success_statuses=("execution_published","format_payment_success")

            where=[
                "ct.status in (?,?)",
                "ct.proc_id is not null"
            ]
            params=[*success_statuses]

            if date_from:
                where.append("ct.contract_close_date>=?")
                params.append(date_from)

            if date_to:
                where.append("ct.contract_close_date<?")
                params.append(end_exclusive(date_to))

            if dim=="customers":
                where.append("nullif(trim(ct.org_company_title),'') is not null")
                if region:
                    where.append("trim(ap.region_name)=?")
                    params.append(region)
                if q:
                    q_clean=q.strip()
                    q_digits="".join(ch for ch in q_clean if ch.isdigit())
                    q_is_tin=q_clean.isdigit() and len(q_digits)>=8
                    qq=f"%{q_clean}%"

                    if q_is_tin:
                        where.append("coalesce(ct.org_inn,'') like ?")
                        params.append(qq)
                    else:
                        where.append("CASEFOLD(coalesce(ct.org_company_title,'')) like CASEFOLD(?)")
                        params.append(qq)

                w=" and ".join(where)

                sql=f"""
                    select
                        ct.org_company_title name,
                        ct.org_inn customer_tin,
                        count(distinct ct.proc_id) deals,
                        coalesce(sum(ct.contract_totalcost),0) amount,
                        coalesce(sum(ct.contract_totalcost),0) / nullif(count(distinct ct.proc_id),0) avg_amount
                    from contracts ct
                    left join auction_procedures ap on ap.proc_id=ct.proc_id
                    where {w}
                    group by ct.org_inn, ct.org_company_title
                    order by amount desc
                    limit ?
                """
                return [dict(x) for x in _market_amount_execute(c,key,amount_min,amount_max,date_from,date_to,region,sql,params+[limit])]

            if dim=="providers":
                where.append("nullif(trim(ct.winner_company_title),'') is not null")
                if region:
                    where.append("trim(ap.region_name)=?")
                    params.append(region)
                if q:
                    q_clean=q.strip()
                    q_digits="".join(ch for ch in q_clean if ch.isdigit())
                    q_is_tin=q_clean.isdigit() and len(q_digits)>=8
                    qq=f"%{q_clean}%"

                    if q_is_tin:
                        where.append("coalesce(ct.winner_inn,'') like ?")
                        params.append(qq)
                    else:
                        where.append("CASEFOLD(coalesce(ct.winner_company_title,'')) like CASEFOLD(?)")
                        params.append(qq)

                w=" and ".join(where)

                sql=f"""
                    select
                        ct.winner_company_title name,
                        ct.winner_inn provider_tin,
                        count(distinct ct.proc_id) deals,
                        coalesce(sum(ct.contract_totalcost),0) amount,
                        coalesce(sum(ct.contract_totalcost),0) / nullif(count(distinct ct.proc_id),0) avg_amount
                    from contracts ct
                    left join auction_procedures ap on ap.proc_id=ct.proc_id
                    where {w}
                    group by ct.winner_inn, ct.winner_company_title
                    order by amount desc
                    limit ?
                """
                return [dict(x) for x in _market_amount_execute(c,key,amount_min,amount_max,date_from,date_to,region,sql,params+[limit])]


            if dim=="products":
                product_where=list(where)
                if region:
                    product_where.append("trim(ap.region_name)=?")
                    params.append(region)
                contract_w=" and ".join(product_where)

                item_where=[
                    "i.proc_id in (" +
                    "select distinct ct.proc_id from contracts ct left join auction_procedures ap on ap.proc_id=ct.proc_id where " +
                    contract_w + ")",
                    "nullif(trim(i.product_name),'') is not null",
                    "nullif(trim(i.unit_name),'') is not null"
                ]
                if q:
                    item_where.append("CASEFOLD(i.product_name) like CASEFOLD(?)")
                    params.append(f"%{q.strip()}%")

                iw=" and ".join(item_where)

                sql=f"""
                    select
                        i.product_name name,
                        i.unit_name unit_name,
                        count(distinct i.proc_id) deals,
                        coalesce(sum(i.quantity),0) quantity,
                        coalesce(sum(i.quantity * i.last_price),0) amount
                    from auction_items i
                    where {iw}
                    group by i.product_name, i.unit_name
                    order by amount desc
                    limit ?
                """

                return [dict(x) for x in _market_amount_execute(c,key,amount_min,amount_max,date_from,date_to,region,sql,params+[limit])]

            if dim=="categories":
                category_where=list(where)
                if region:
                    category_where.append("trim(ap.region_name)=?")
                    params.append(region)
                contract_w=" and ".join(category_where)

                sql=f"""
                    select
                        coalesce(
                            nullif(trim(i.category_name),''),
                            nullif(trim(p.category_name),'')
                        ) name,
                        count(*) items,
                        count(distinct i.proc_id) deals,
                        coalesce(sum(i.quantity * i.last_price),0) amount
                    from auction_items i
                    left join auction_product_canonical_v3 p
                      on p.product_code=i.product_code
                    where i.proc_id in (
                        select distinct ct.proc_id
                        from contracts ct left join auction_procedures ap on ap.proc_id=ct.proc_id
                        where {contract_w}
                    )
                      and coalesce(
                            nullif(trim(i.category_name),''),
                            nullif(trim(p.category_name),'')
                          ) is not null
                      and (? = '' or CASEFOLD(coalesce(
                            nullif(trim(i.category_name),''),
                            nullif(trim(p.category_name),'')
                          )) like CASEFOLD(?))
                    group by coalesce(
                        nullif(trim(i.category_name),''),
                        nullif(trim(p.category_name),'')
                    )
                    order by amount desc
                    limit ?
                """

                return [dict(x) for x in _market_amount_execute(c,key,amount_min,amount_max,date_from,date_to,region,sql,params+[(q or "").strip(),f"%{(q or "").strip()}%",limit])]

            if dim=="regions":
                region_where=[
                    "ct.status in (?,?)",
                    "ct.proc_id is not null",
                    "nullif(trim(p.region_name),'') is not null"
                ]
                region_params=[*success_statuses]

                if date_from:
                    region_where.append("ct.contract_close_date>=?")
                    region_params.append(date_from)

                if date_to:
                    region_where.append("ct.contract_close_date<?")
                    region_params.append(end_exclusive(date_to))


                if region:
                    region_where.append("trim(p.region_name)=?")
                    region_params.append(region)
                if q:
                    region_where.append("CASEFOLD(trim(p.region_name)) like CASEFOLD(?)")
                    region_params.append(f"%{q.strip()}%")
                rw=" and ".join(region_where)

                sql=f"""
                    select
                        trim(p.region_name) name,
                        count(distinct ct.proc_id) deals,
                        count(distinct ct.org_inn) customers,
                        count(distinct ct.winner_inn) providers,
                        coalesce(sum(ct.contract_totalcost),0) amount
                    from contracts ct
                    join auction_procedures p
                      on p.proc_id=ct.proc_id
                    where {rw}
                    group by trim(p.region_name)
                    order by amount desc
                    limit ?
                """

                return [dict(x) for x in _market_amount_execute(c,key,amount_min,amount_max,date_from,date_to,region,
                    sql,region_params+[limit]
                )]

            if dim=="statuses":
                status_where=list(where)
                status_params=list(params)

                if region:
                    status_where.append("trim(ap.region_name)=?")
                    status_params.append(region)

                if q:
                    status_where.append("ct.status like ?")
                    status_params.append(f"%{q.strip()}%")

                sw=" and ".join(status_where)

                sql=f"""
                    select
                        ct.status name,
                        count(*) contracts,
                        count(distinct ct.proc_id) deals,
                        coalesce(sum(ct.contract_totalcost),0) amount
                    from contracts ct
                    left join auction_procedures ap on ap.proc_id=ct.proc_id
                    where {sw}
                    group by ct.status
                    order by amount desc
                    limit ?
                """

                return [dict(x) for x in _market_amount_execute(c,key,amount_min,amount_max,date_from,date_to,region,sql,status_params+[limit])]

            return []
        if v["kind"]=="xt":
            if dim in {"products","categories"} and "items" in tables(c):
                col="product_name" if dim=="products" else "category_name"; return top_auto(c,"items",col,"total_price",limit)
            mp={"customers":"customer_name","statuses":"status"}; return top_auto(c,"tenders",mp.get(dim),"total_amount",limit)
        if v["kind"]=="etender":
            mp={"products":"category_name","categories":"category_name","customers":"customer_name","providers":"provider_name","statuses":"deal_status_name"}
            return top_auto(c,"completed_deals",mp.get(dim),"deal_cost",limit)
        if v["kind"]=="shaffof":
            mp={"products":"service_type","categories":"object_type","regions":"region","customers":"customer_name","providers":"winner_name","statuses":"status_name"}
            return top_auto(c,"tenders",mp.get(dim),"winner_offer",limit)
        if v["kind"]=="mc":
            if dim=="products" and "tender_items" in tables(c): return top_auto(c,"tender_items","product_name","total_price",limit)
            mp={"categories":"service_type_name","regions":"region_name_uz","customers":"customer_name","statuses":"status"}
            return top_auto(c,"tenders",mp.get(dim),"start_price",limit)
        if key=='xt_national' and (amount_min is not None or amount_max is not None):
            return _amount_xt_national_dimension(c,dim,limit,q,date_from,date_to,region,amount_min,amount_max)
        t=best_table(c); cs=cols(c,t); amap={"products":"product","categories":"category","regions":"region","customers":"customer","providers":"supplier","statuses":"status"}
        return top_auto(c,t,detect(cs,amap[dim]),detect(cs,"amount"),limit)

@app.get("/api/{key}/records")
def records(key:str,limit:int=Query(50,ge=1,le=200),offset:int=Query(0,ge=0),q:Optional[str]=None,
            date_from:Optional[str]=None,date_to:Optional[str]=None,region:Optional[str]=None,amount_min:Optional[str]=None,amount_max:Optional[str]=None):
    amount_min,amount_max=_market_amount_bounds(key,amount_min,amount_max)
    v=platform(key)
    with ro(v["db"]) as c:
        if v["kind"]=="completed":
            where,p=completed_filter(date_from,date_to,region,q,key)
            rows=_market_amount_execute(c,key,amount_min,amount_max,date_from,date_to,region,f"""select deal_date date,lot_display_no display_no,product_name product,amount quantity,
            coalesce(nullif(unit_name_enriched,''),unit_name,'') unit_name,
            case when amount<>0 then deal_cost/amount end unit_price,deal_cost amount,
            customer_region_name region,customer_name customer,provider_name supplier,deal_status_name status
            from completed_deals where {where} order by deal_date desc limit ? offset ?""",p+[limit,offset],search=q).fetchall()
            return {"columns":["date","display_no","product","quantity","unit_name","unit_price","amount","region","customer","supplier","status"],"rows":[dict(x) for x in rows]}
        if v["kind"]=="auction":
            success=(
                "\u041f\u043e\u0441\u0442\u0430\u0432\u043a\u0430 \u0441\u043e\u0432\u0435\u0440\u0448\u0435\u043d\u0430",
                "\u041e\u043f\u043b\u0430\u0447\u0435\u043d\u0430",
                "\u041f\u043e\u0441\u0442\u0430\u0432\u043a\u0430 \u0441\u043e\u0432\u0435\u0440\u0448\u0435\u043d\u0430 (\u0412\u043e\u0441\u0441\u0442\u0430\u043d\u043e\u0432\u043b\u0435\u043d\u0430)",
            )

            where=["status in (?,?,?)"]
            params=list(success)

            if date_from:
                where.append("substr(deal_date,1,10)>=?")
                params.append(date_from[:10])

            if date_to:
                where.append("substr(deal_date,1,10)<=?")
                params.append(date_to[:10])

            if region:
                where.append("region_name=?")
                params.append(region)

            if q:
                q_clean=q.strip()
                q_digits="".join(ch for ch in q_clean if ch.isdigit())
                q_is_tin=q_clean.isdigit() and len(q_digits)>=8
                qq=f"%{q_clean}%"

                if q_is_tin:
                    where.append("""(
                        coalesce(customer_tin,'') like ? or
                        coalesce(provider_tin,'') like ?
                    )""")
                    params.extend([qq,qq])
                else:
                    where.append("""exists (
                        select 1
                        from auction_products aq
                        where aq.completed_deal_id=completed_deals.id
                          and CASEFOLD(coalesce(aq.product_name,'')) like CASEFOLD(?)
                    )""")
                    params.append(qq)

            w=" and ".join(where)

            sql=f"""select deal_date date,
                           lot_number display_no,
                           category_name product,
                           contract_price amount,
                           region_name region,
                           customer_name customer,
                           provider_name supplier,
                           status
                    from completed_deals
                    where {w}
                    order by deal_date desc
                    limit ? offset ?"""

            rows=_market_amount_execute(c,key,amount_min,amount_max,date_from,date_to,region,
                sql,
                params+[limit,offset]
            ,search=q).fetchall()

            return {
                "columns":[
                    "date","display_no","product","amount",
                    "region","customer","supplier","status"
                ],
                "rows":[dict(x) for x in rows]
            }
        if v["kind"]=="xt_auction_completed":
            success_statuses=("execution_published","format_payment_success")

            where=[
                "ct.status in (?,?)",
                "ct.proc_id is not null"
            ]
            params=[*success_statuses]

            if date_from:
                where.append("ct.contract_close_date>=?")
                params.append(date_from)

            if date_to:
                where.append("ct.contract_close_date<?")
                params.append(end_exclusive(date_to))

            if region:
                where.append("trim(p.region_name)=?")
                params.append(region)

            if q:
                q_clean=q.strip()
                qq=f"%{q_clean}%"

                if q_clean.isdigit() and len(q_clean)>=8:
                    where.append("""(
                        coalesce(ct.org_inn,'') like ? or
                        coalesce(ct.winner_inn,'') like ?
                    )""")
                    params.extend([qq,qq])
                else:
                    where.append("""exists (
                        select 1
                        from auction_items qi
                        where qi.proc_id=ct.proc_id
                          and CASEFOLD(coalesce(qi.product_name,'')) like CASEFOLD(?)
                    )""")
                    params.append(qq)

            w=" and ".join(where)

            rows=_market_amount_execute(c,key,amount_min,amount_max,date_from,date_to,region,f"""
                select
                    ct.contract_close_date date,
                    ct.contract_number display_no,
                    ct.proc_id proc_id,
                    ct.contract_totalcost amount,
                    p.region_name region,
                    ct.org_company_title customer,
                    ct.org_inn customer_id,
                    ct.winner_company_title supplier,
                    ct.winner_inn supplier_id,
                    ct.status status
                from contracts ct
                left join auction_procedures p
                  on p.proc_id=ct.proc_id
                where {w}
                order by ct.contract_close_date desc, ct.contract_id desc
                limit ? offset ?
            """,params+[limit,offset],search=q).fetchall()

            return {
                "columns":[
                    "date","display_no","proc_id","amount","region",
                    "customer","customer_id","supplier","supplier_id","status"
                ],
                "rows":[dict(x) for x in rows]
            }
        if v["kind"]=="xt":
            rows=_market_amount_execute(c,key,amount_min,amount_max,date_from,date_to,region,"""select publicated_at date,coalesce(lot_number,cast(proc_id as text)) display_no,name product,total_amount amount,
            customer_name customer,status,close_at deadline,participant_count participants from tenders order by publicated_at desc limit ? offset ?""",(limit,offset),search=q).fetchall()
            return {"columns":list(rows[0].keys()) if rows else [],"rows":[dict(x) for x in rows]}
        if v["kind"]=="etender":
            rows=_market_amount_execute(c,key,amount_min,amount_max,date_from,date_to,region,"""select deal_date date,display_no,category_name product,start_cost,deal_cost amount,saving_amount,
            participants_count participants,customer_name customer,provider_name supplier,deal_status_name status
            from completed_deals order by deal_date desc limit ? offset ?""",(limit,offset),search=q).fetchall()
            return {"columns":list(rows[0].keys()) if rows else [],"rows":[dict(x) for x in rows]}
        if v["kind"]=="shaffof":
            rows=_market_amount_execute(c,key,amount_min,amount_max,date_from,date_to,region,"""select created_at date,unique_name display_no,name product,start_price,winner_offer amount,saving_amount,
            region,customer_name customer,winner_name supplier,status_name status,bidders_count participants
            from tenders order by created_at desc limit ? offset ?""",(limit,offset),search=q).fetchall()
            return {"columns":list(rows[0].keys()) if rows else [],"rows":[dict(x) for x in rows]}
        if v["kind"]=="mc":
            rows=_market_amount_execute(c,key,amount_min,amount_max,date_from,date_to,region,"""select confirmed_date date,coalesce(unique_name,lotid,cast(id as text)) display_no,name product,start_price amount,
            coalesce(region_name_uz,region_name) region,customer_name customer,status,placement_term deadline
            from tenders order by confirmed_date desc limit ? offset ?""",(limit,offset),search=q).fetchall()
            return {"columns":list(rows[0].keys()) if rows else [],"rows":[dict(x) for x in rows]}
        t=best_table(c); cs=cols(c,t)
        # Safe generic browser: choose useful detected columns, then first columns.
        wanted=[]
        for k in ["date","display","product","amount","region","customer","supplier","status"]:
            x=detect(cs,k)
            if x and x not in wanted:wanted.append(x)
        for x in cs:
            if x not in wanted and len(wanted)<12 and "json" not in x.lower():wanted.append(x)
        order=detect(cs,"date")
        sql=f"select {','.join(ident(x) for x in wanted)} from {ident(t)}"+(f" order by {ident(order)} desc" if order else "")+" limit ? offset ?"
        rr=_market_amount_execute(c,key,amount_min,amount_max,date_from,date_to,region,sql,(limit,offset),search=q).fetchall()
        return {"table":t,"columns":wanted,"rows":[dict(x) for x in rr]}

@app.get("/api/{key}/export.csv")
def export_csv(key:str,limit:int=Query(5000,ge=1,le=20000),date_from:Optional[str]=None,date_to:Optional[str]=None,region:Optional[str]=None,q:Optional[str]=None,amount_min:Optional[str]=None,amount_max:Optional[str]=None):
    amount_min,amount_max=_market_amount_bounds(key,amount_min,amount_max)
    data=records(key,limit=limit,offset=0,q=q,date_from=date_from,date_to=date_to,region=region,amount_min=amount_min,amount_max=amount_max)
    sio=io.StringIO(); w=csv.DictWriter(sio,fieldnames=data["columns"],extrasaction="ignore");w.writeheader()
    for r in data["rows"]:w.writerow(r)
    b=io.BytesIO(("\ufeff"+sio.getvalue()).encode("utf-8-sig"))
    return StreamingResponse(b,media_type="text/csv",headers={"Content-Disposition":f'attachment; filename="uzbi_{key}_export.csv"'})

@app.post("/api/system/cache/clear")
def clear_cache():
    cached_meta.cache_clear()
    return {"status":"ok"}



def _auction_dimension_query(c, dim:str, limit:int=100, q:str|None=None):
    if dim=="products":
        table="auction_products"; col="product_name"; amount_col="total_amount"
    elif dim=="categories":
        table="completed_deals"; col="category_name"; amount_col="contract_price"
    elif dim=="customers":
        table="completed_deals"; col="customer_name"; amount_col="contract_price"
    elif dim=="providers":
        table="completed_deals"; col="provider_name"; amount_col="contract_price"
    elif dim=="regions":
        table="completed_deals"; col="region_name"; amount_col="contract_price"
    elif dim=="statuses":
        table="completed_deals"; col="status"; amount_col="contract_price"
    else:
        raise HTTPException(400,"dimension")
    where=f"coalesce(trim({col}),'')<>''"
    p=[]
    if q:
        where+=f" and {col} like ?"
        p.append(f"%{q}%")
    p.append(limit)
    sql=f"""
        select trim({col}) name,
               count(*) deals,
               coalesce(sum({amount_col}),0) amount,
               coalesce(avg({amount_col}),0) avg_amount
        from {table}
        where {where}
        group by trim({col})
        order by amount desc
        limit ?
    """
    try:
        return [dict(x) for x in c.execute(sql,p).fetchall()]
    except sqlite3.Error as e:
        raise HTTPException(500,f"AUCTION_SQL_ERROR [{table}.{col}]: {e}")


# ==================== v10.5 FAST PRODUCT SEARCH / BI ====================
FAST_MARKET_LABELS={"shop":"xarid.uzex.uz Р В Р вЂ Р В РІР‚С™Р Р†Р вЂљРЎСљ Р В Р’В Р вЂ™Р’В­Р В Р’В Р вЂ™Р’В»Р В Р’В Р вЂ™Р’ВµР В Р’В Р РЋРІР‚СњР В Р Р‹Р Р†Р вЂљРЎв„ўР В Р Р‹Р В РІР‚С™Р В Р’В Р РЋРІР‚СћР В Р’В Р В РІР‚В¦ Р В Р’В Р СћРІР‚ВР В Р Р‹Р РЋРІР‚С”Р В Р’В Р РЋРІР‚СњР В Р’В Р РЋРІР‚СћР В Р’В Р В РІР‚В¦","national":"xarid.uzex.uz Р В Р вЂ Р В РІР‚С™Р Р†Р вЂљРЎСљ Р В Р’В Р РЋРЎв„ўР В Р’В Р РЋРІР‚ВР В Р’В Р вЂ™Р’В»Р В Р’В Р вЂ™Р’В»Р В Р’В Р РЋРІР‚ВР В Р’В Р Р†РІР‚С›РІР‚вЂњ Р В Р’В Р СћРІР‚ВР В Р Р‹Р РЋРІР‚С”Р В Р’В Р РЋРІР‚СњР В Р’В Р РЋРІР‚СћР В Р’В Р В РІР‚В¦","auction":"xarid.uzex.uz Р В Р вЂ Р В РІР‚С™Р Р†Р вЂљРЎСљ Р В Р’В Р РЋРІР‚в„ўР В Р Р‹Р РЋРІР‚СљР В Р’В Р РЋРІР‚СњР В Р Р‹Р Р†Р вЂљР’В Р В Р’В Р РЋРІР‚ВР В Р’В Р РЋРІР‚СћР В Р’В Р В РІР‚В¦ Р В Р Р‹Р В РЎвЂњР В Р’В Р вЂ™Р’В°Р В Р’В Р В РІР‚В Р В Р’В Р СћРІР‚ВР В Р’В Р РЋРІР‚Сћ","xt_auction":"xt-xarid.uz Р В Р вЂ Р В РІР‚С™Р Р†Р вЂљРЎСљ Р В Р’В Р РЋРІР‚в„ўР В Р Р‹Р РЋРІР‚СљР В Р’В Р РЋРІР‚СњР В Р Р‹Р Р†Р вЂљР’В Р В Р’В Р РЋРІР‚ВР В Р’В Р РЋРІР‚СћР В Р’В Р В РІР‚В¦ Р В Р Р‹Р В РЎвЂњР В Р’В Р вЂ™Р’В°Р В Р’В Р В РІР‚В Р В Р’В Р СћРІР‚ВР В Р’В Р РЋРІР‚Сћ","xt_national":"xt-xarid.uz Р В Р вЂ Р В РІР‚С™Р Р†Р вЂљРЎСљ Р В Р’В Р РЋРЎв„ўР В Р’В Р РЋРІР‚ВР В Р’В Р вЂ™Р’В»Р В Р’В Р вЂ™Р’В»Р В Р’В Р РЋРІР‚ВР В Р’В Р Р†РІР‚С›РІР‚вЂњ Р В Р’В Р СћРІР‚ВР В Р Р‹Р РЋРІР‚С”Р В Р’В Р РЋРІР‚СњР В Р’В Р РЋРІР‚СћР В Р’В Р В РІР‚В¦"}
FAST_MARKET_ICONS={"shop":"Р РЋР вЂљР РЋРЎСџР Р†Р вЂљРЎвЂќР Р†Р вЂљРІвЂћСћ","national":"UZ","auction":"Р В Р вЂ Р Р†Р вЂљРІР‚СљР СћРЎвЂ™","xt_auction":"Р В Р вЂ Р Р†Р вЂљРІР‚СњР Р†Р вЂљР’В°","xt_national":"Р В Р вЂ Р Р†Р вЂљРІР‚СњР Р†Р вЂљР’В°"}
FAST_MARKET_LOGOS={"shop":"/static/platform-icons/shop.png","national":"/static/platform-icons/national.png","auction":"/static/platform-icons/auction.png","xt_auction":"/static/platform-icons/xt.png","xt_national":"/static/platform-icons/xt.png"}
FAST_MARKET_SOURCES={"shop":"xarid.uzex.uz","national":"xarid.uzex.uz","auction":"xarid.uzex.uz","xt_auction":"xt-xarid.uz","xt_national":"xt-xarid.uz"}

def _normalize_product_query(s:str|None):
    s=(s or "").strip().lower().replace("\u0451","\u0435")
    s=re.sub(r"[\u201c\u201d\u201e\u00ab\u00bb\"'`\u2019]+"," ",s)
    s=re.sub(r"[^\w\s\-\+\/\.]"," ",s,flags=re.UNICODE)
    s=s.replace("_"," ")
    s=re.sub(r"[\s\-\+\/\.]+"," ",s)
    return s.strip()

def _product_search_ro():
    if not PRODUCT_SEARCH_DB.exists(): raise HTTPException(503,f"Product search DB Р В Р Р‹Р Р†Р вЂљРЎв„ўР В Р’В Р РЋРІР‚СћР В Р’В Р РЋРІР‚вЂќР В Р’В Р РЋРІР‚ВР В Р’В Р вЂ™Р’В»Р В Р’В Р РЋР’ВР В Р’В Р вЂ™Р’В°Р В Р’В Р СћРІР‚ВР В Р’В Р РЋРІР‚В: {PRODUCT_SEARCH_DB}")
    return ro(PRODUCT_SEARCH_DB)

def _fts_query(q:str):
    return " ".join(f'"{p}"*' for p in _normalize_product_query(q).split() if p)


def _resolve_shop_product_names(q, limit=50):
    """
    Resolve free-text Shop product search through the existing
    product_search_fts cache.

    Returns exact display_name values that can be used with:
        completed_deals.product_name IN (...)
    """
    qn = _normalize_product_query(q)

    if len(qn) < 2:
        return []

    try:
        with _product_search_ro() as c:
            rows = c.execute(
                """
                select p.display_name
                from product_search_fts f
                join product_search p
                  on p.product_id=f.rowid
                join product_market_stats m
                  on m.product_id=p.product_id
                 and m.market='shop'
                where product_search_fts match ?
                order by
                    case
                        when p.normalized_name=? then 0
                        when p.normalized_name like ? then 1
                        else 2
                    end,
                    m.deals_count desc,
                    p.display_name
                limit ?
                """,
                (
                    _fts_query(qn),
                    qn,
                    qn + "%",
                    limit,
                ),
            ).fetchall()

            return [
                str(r[0]).strip()
                for r in rows
                if r[0] and str(r[0]).strip()
            ]

    except sqlite3.Error:
        return []


def _manual_product_code_override(name):
    """Return a manually verified canonical product code, if one exists."""
    name = str(name or "").strip()
    if not name or not PRODUCT_CODES_DB.exists():
        return None

    key = name.casefold()

    try:
        with ro(PRODUCT_CODES_DB) as c:
            if "product_code_manual_overrides" not in set(tables(c)):
                return None

            row = c.execute("""
                select canonical_code
                from product_code_manual_overrides
                where normalized_name = ?
                  and status = 'manual_verified'
                limit 1
            """, (key,)).fetchone()

            if not row:
                return None

            code = str(row["canonical_code"] or "").strip()
            return code if code else None

    except Exception:
        return None


def _product_codes_ro():
    if not PRODUCT_CODES_DB.exists():
        raise HTTPException(404,"Р В Р’В Р РЋРЎвЂєР В Р’В Р РЋРІР‚СћР В Р’В Р В РІР‚В Р В Р’В Р вЂ™Р’В°Р В Р Р‹Р В РІР‚С™ Р В Р’В Р РЋРІР‚СњР В Р’В Р РЋРІР‚СћР В Р’В Р СћРІР‚ВР В Р’В Р вЂ™Р’В»Р В Р’В Р вЂ™Р’В°Р В Р Р‹Р В РІР‚С™Р В Р’В Р РЋРІР‚В Р В Р’В Р вЂ™Р’В±Р В Р’В Р вЂ™Р’В°Р В Р’В Р вЂ™Р’В·Р В Р’В Р вЂ™Р’В°Р В Р Р‹Р В РЎвЂњР В Р’В Р РЋРІР‚В Р В Р Р‹Р Р†Р вЂљРЎв„ўР В Р’В Р РЋРІР‚СћР В Р’В Р РЋРІР‚вЂќР В Р’В Р РЋРІР‚ВР В Р’В Р вЂ™Р’В»Р В Р’В Р РЋР’ВР В Р’В Р вЂ™Р’В°Р В Р’В Р СћРІР‚ВР В Р’В Р РЋРІР‚В")
    return ro(PRODUCT_CODES_DB)

def _code_norm(v):
    return re.sub(r"[^0-9a-zР В Р’В Р вЂ™Р’В°-Р В Р Р‹Р В Р РЏР В Р Р‹Р Р†Р вЂљР’В]+","",str(v or "").lower())

def _product_codes_for_name(name:str, limit:int=12):
    if not PRODUCT_CODES_DB.exists() or not str(name or "").strip(): return []
    with ro(PRODUCT_CODES_DB) as c:
        ts=set(tables(c))
        rows=[]
        if "product_codes" in ts:
            rows += [dict(r) for r in c.execute("""
                select product_code,product_name,category_name,last_seen_at,updated_at
                from product_codes
                where lower(trim(product_name))=lower(trim(?))
                order by coalesce(last_seen_at,updated_at,'') desc, product_code
                limit ?
            """,(name,limit)).fetchall()]
        # Fallback/extra aliases from historical lot mappings, deduplicated later.
        if len(rows)<limit and "product_code_lots" in ts:
            rows += [dict(r) for r in c.execute("""
                select product_code,product_name,category_name,seen_at as last_seen_at,null as updated_at
                from product_code_lots
                where lower(trim(product_name))=lower(trim(?))
                order by coalesce(seen_at,'') desc, product_code
                limit ?
            """,(name,limit*2)).fetchall()]
        out=[];seen=set()
        for r in rows:
            code=str(r.get("product_code") or "").strip()
            if not code or code in seen: continue
            seen.add(code);out.append(r)
            if len(out)>=limit: break
        return out

@app.get("/api/product-codes/search")
def product_codes_search(q:str=Query(...,min_length=1),limit:int=Query(30,ge=1,le=100)):
    q=str(q or "").strip()
    if not q:return {"query":q,"count":0,"rows":[]}
    with _product_codes_ro() as c:
        ts=set(tables(c))
        if "product_codes" not in ts:
            raise HTTPException(500,"product_codes Р В Р’В Р вЂ™Р’В¶Р В Р’В Р вЂ™Р’В°Р В Р’В Р СћРІР‚ВР В Р’В Р В РІР‚В Р В Р’В Р вЂ™Р’В°Р В Р’В Р вЂ™Р’В»Р В Р’В Р РЋРІР‚В Р В Р Р‹Р Р†Р вЂљРЎв„ўР В Р’В Р РЋРІР‚СћР В Р’В Р РЋРІР‚вЂќР В Р’В Р РЋРІР‚ВР В Р’В Р вЂ™Р’В»Р В Р’В Р РЋР’ВР В Р’В Р вЂ™Р’В°Р В Р’В Р СћРІР‚ВР В Р’В Р РЋРІР‚В")
        like=f"%{q}%"; prefix=f"{q}%"
        rows=[dict(r) for r in c.execute("""
            select product_code,product_name,category_name,first_seen_at,last_seen_at,updated_at
            from product_codes
            where product_code like ?
               or lower(product_name) like lower(?)
               or lower(coalesce(category_name,'')) like lower(?)
            order by case
              when product_code=? then 0
              when product_code like ? then 1
              when lower(trim(product_name))=lower(trim(?)) then 2
              when lower(product_name) like lower(?) then 3
              else 4 end,
              product_name, product_code
            limit ?
        """,(like,like,like,q,prefix,q,prefix,limit)).fetchall()]
        return {"query":q,"count":len(rows),"rows":rows}

@app.get("/api/product-codes/for-product")
def product_codes_for_product(name:str=Query(...,min_length=1),limit:int=Query(12,ge=1,le=50)):
    rows=_product_codes_for_name(name,limit)
    return {"product_name":name,"count":len(rows),"rows":rows}

@app.get("/api/product-codes/debug-name")
def product_codes_debug_name(name:str=Query(...,min_length=1)):
    codes=_catalog_codes_for_exact_product_name(name,50)
    return {
        "db":str(PRODUCT_CODES_DB),
        "exists":PRODUCT_CODES_DB.exists(),
        "product_name":name,
        "codes":codes,
        "count":len(codes),
    }

@app.get("/api/product-codes/stats")
def product_codes_stats():
    if not PRODUCT_CODES_DB.exists(): return {"exists":False,"codes":0,"lot_links":0,"updated_at":None}
    with ro(PRODUCT_CODES_DB) as c:
        ts=set(tables(c)); codes=c.execute("select count(*) from product_codes").fetchone()[0] if "product_codes" in ts else 0
        lots=c.execute("select count(*) from product_code_lots").fetchone()[0] if "product_code_lots" in ts else 0
        updated=None
        if "product_codes" in ts:
            try: updated=c.execute("select max(coalesce(updated_at,last_seen_at)) from product_codes").fetchone()[0]
            except Exception: pass
        return {"exists":True,"codes":int(codes),"lot_links":int(lots),"updated_at":updated}

@app.get("/api/products/search")
def fast_product_search(q:str=Query(...,min_length=2),limit:int=Query(20,ge=1,le=50)):
    qn=_normalize_product_query(q)
    if len(qn)<2:return []
    with _product_search_ro() as c:
        has_fts=c.execute("select 1 from sqlite_master where name='product_search_fts'").fetchone()
        if has_fts:
            rows=c.execute("""select p.product_id,p.display_name,p.normalized_name,p.deals_count,p.markets_count,
             p.first_date,p.last_date,p.avg_unit_price,p.min_unit_price,p.max_unit_price,p.median_unit_price
             from product_search_fts f join product_search p on p.product_id=f.rowid
             where product_search_fts match ?
             order by case when p.normalized_name like ? then 0 else 1 end,p.deals_count desc,p.display_name limit ?""",
             (_fts_query(qn),qn+"%",limit)).fetchall()
            return [dict(r) for r in rows]
        return [dict(r) for r in c.execute("""select product_id,display_name,normalized_name,deals_count,markets_count,
         first_date,last_date,avg_unit_price,min_unit_price,max_unit_price,median_unit_price
         from product_search where normalized_name like ? order by deals_count desc,display_name limit ?""",(qn+"%",limit)).fetchall()]

def _xt_norm_text(v):
    return " ".join(str(v or "").strip().lower().split())

def _xt_code_norm(v):
    return str(v or "").strip().replace("-", "_")

def _repair_xt_auction_recent(rows):
    """Repair XT-Xarid auction recent rows from the source DB.

    fact_deals historically stored auction_items.last_price in deal_amount for some XT rows,
    so generic unit_price=deal_amount/quantity makes the unit price too small.  The source
    auction_items table already contains the correct per-unit last_price.
    """
    if not rows:
        return rows
    try:
        db=platform("xt_auction")["db"]
    except Exception:
        return rows
    if not db.exists():
        return rows

    try:
        with ro(db) as xc:
            ts=set(tables(xc))
            if "auction_items" not in ts:
                return rows

            item_cache={}
            contract_cache={}
            procedure_cache={}

            for r in rows:
                if str(r.get("platform_key") or r.get("market") or "") != "xt_auction":
                    continue

                proc_id=r.get("deal_id") or r.get("trade_id") or r.get("source_id")
                try:
                    proc_id=int(str(proc_id))
                except Exception:
                    continue

                if proc_id not in item_cache:
                    item_cache[proc_id]=[dict(x) for x in xc.execute(
                        "select * from auction_items where proc_id=?", (proc_id,)
                    ).fetchall()]
                candidates=item_cache[proc_id]
                if not candidates:
                    continue

                target_code=_xt_code_norm(r.get("product_code"))
                target_name=_xt_norm_text(r.get("product_name"))
                target_desc=_xt_norm_text(r.get("description"))
                try: target_qty=float(r.get("quantity"))
                except Exception: target_qty=None

                def score(it):
                    sc=0
                    if target_code and _xt_code_norm(it.get("product_code"))==target_code: sc+=40
                    if target_name and _xt_norm_text(it.get("product_name"))==target_name: sc+=20
                    if target_desc and _xt_norm_text(it.get("spec"))==target_desc: sc+=100
                    try:
                        q=float(it.get("quantity"))
                        if target_qty is not None and abs(q-target_qty)<1e-9: sc+=30
                    except Exception:
                        pass
                    return sc

                best=max(candidates, key=score)
                if score(best)<=0:
                    continue

                # Procedure enrichment is authoritative for the public XT lot id and region.
                if proc_id not in procedure_cache:
                    pr=None
                    if "auction_procedures" in ts:
                        try:
                            pr=xc.execute(
                                "select * from auction_procedures where proc_id=? limit 1",
                                (proc_id,)
                            ).fetchone()
                        except Exception:
                            pr=None
                    procedure_cache[proc_id]=dict(pr) if pr else None
                pr=procedure_cache[proc_id]

                public_lot_id = None
                if pr:
                    public_lot_id = pr.get("real_lot_id")
                if public_lot_id in (None, ""):
                    public_lot_id = proc_id

                # Correct identifiers and item values from source auction_items.
                r["lot_no"]=str(public_lot_id)
                r["lot_number"]=str(public_lot_id)
                r["source_id"]=proc_id
                r["proc_id"]=proc_id

                if pr:
                    if pr.get("delivery_address"):
                        r["delivery_address"]=pr.get("delivery_address")
                    if pr.get("region_area_id"):
                        r["region_area_id"]=pr.get("region_area_id")
                    if pr.get("region_name"):
                        r["region_name"]=pr.get("region_name")
                        r["region"]=pr.get("region_name")
                    if pr.get("region_source"):
                        r["region_source"]=pr.get("region_source")
                    if pr.get("real_lot_id"):
                        r["real_lot_id"]=pr.get("real_lot_id")
                if best.get("quantity") is not None: r["quantity"]=best.get("quantity")
                if best.get("unit_name"): r["unit_name"]=best.get("unit_name")
                if best.get("product_code"): r["product_code"]=best.get("product_code")

                # XT Auction UI product-code priority:
                # 1) manually verified override
                # 2) unique exact-name MASTER
                # 3) original XT Auction source code
                product_name_for_code=(
                    r.get("product_name") or best.get("product_name")
                )

                manual_code=_manual_product_code_override(
                    product_name_for_code
                )

                if manual_code:
                    r["product_code"]=manual_code
                    r["product_code_source"]="manual_verified"
                else:
                    master_codes=_catalog_codes_for_exact_product_name(
                        product_name_for_code, 12
                    )
                    master_unique=[]
                    master_seen=set()

                    for mc in master_codes or []:
                        mc=str(mc or "").strip()
                        if mc and mc not in master_seen:
                            master_seen.add(mc)
                            master_unique.append(mc)

                    if len(master_unique)==1:
                        r["product_code"]=master_unique[0]
                        r["product_code_source"]="product_codes.sqlite3"

                if best.get("spec"): r["description"]=best.get("spec")
                if best.get("product_properties_json"): r["properties_json"]=best.get("product_properties_json")

                try: qty=float(best.get("quantity") or 0)
                except Exception: qty=0.0
                try: last_price=float(best.get("last_price"))
                except Exception: last_price=None
                if last_price is not None:
                    r["unit_price"]=last_price
                    r["deal_amount"]=last_price*qty if qty>0 else None
                    r["total_value"]=last_price*qty if qty>0 else None

                # Contract is authoritative for buyer / winner and completion date.
                if proc_id not in contract_cache:
                    ct=None
                    if "contracts" in ts:
                        ct=xc.execute(
                            "select * from contracts where proc_id=? order by updated_at desc limit 1",
                            (proc_id,)
                        ).fetchone()
                    contract_cache[proc_id]=dict(ct) if ct else None
                ct=contract_cache[proc_id]
                if ct:
                    if ct.get("org_company_title"):
                        r["customer_name"]=ct.get("org_company_title")
                        r["customer"]=ct.get("org_company_title")
                    if ct.get("winner_company_title"):
                        r["provider_name"]=ct.get("winner_company_title")
                        r["supplier"]=ct.get("winner_company_title")
                    d=ct.get("contract_close_date") or ct.get("contract_close_at")
                    if d:
                        r["deal_date"]=str(d)[:10]
                        r["date"]=str(d)[:10]
                    r["contract_number"]=ct.get("contract_number")
                    r["winner_inn"]=ct.get("winner_inn")
                    r["customer_inn"]=ct.get("org_inn")

                r["xt_source_repaired"]=True
                if pr and pr.get("region_name"):
                    r["xt_region_repaired"]=True
            return rows
    except Exception as e:
        # Never break product analytics if the source DB is temporarily unavailable.
        for r in rows:
            if str(r.get("platform_key") or r.get("market") or "") == "xt_auction":
                r["xt_repair_error"]=str(e)
        return rows


def _recent_product_code_from_source(row:dict):
    """Recover a missing product code for UZEX shop/national/auction rows.

    The unified analytics fact table can contain a null product_code even when the
    source DB has the code.  We only accept source rows tied to the same product
    name and one of the same lot/deal identifiers to avoid assigning a code from
    another similarly named product.
    """
    key=str(row.get("platform_key") or row.get("market") or "")
    if key not in {"shop","national","auction"}:
        return None

    try:
        v=platform(key)
    except Exception:
        return None
    if not v["db"].exists():
        return None

    target_name=str(row.get("product_name") or "").strip()
    if not target_name:
        return None

    # Values that can identify the concrete historical lot/deal.
    ids=[]
    for k in ("lot_no","lot_number","lot_id","trade_id","deal_id","source_id"):
        val=row.get(k)
        if val not in (None,""):
            s=str(val).strip()
            if s and s not in ids:
                ids.append(s)
    if not ids:
        return None

    code_names=["product_code","code","mxik_code","ikpu_code","catalog_code","classifier_code","product_classifier_code"]
    name_names=["product_name","name","title","category_name","subject"]
    id_names=["lot_display_no","lot_number","lot_no","lot_id","deal_id","trade_id","proc_id","source_id","id"]

    try:
        with ro(v["db"]) as c:
            ts=set(tables(c))
            candidates=[]
            if key=="auction":
                candidates += [t for t in ("auction_products","completed_products","products","completed_deals") if t in ts]
            else:
                candidates += [t for t in ("completed_deals","products","deal_products","items") if t in ts]
            # Include any other likely product table as a final source-side fallback.
            for t in ts:
                low=t.lower()
                if t not in candidates and ("product" in low or "deal" in low):
                    candidates.append(t)

            for t in candidates:
                cs=cols(c,t)
                codec=_pick(cs,code_names)
                namec=_pick(cs,name_names)
                if not codec or not namec:
                    continue
                idcs=[]
                low={x.lower():x for x in cs}
                for n in id_names:
                    if n.lower() in low and low[n.lower()] not in idcs:
                        idcs.append(low[n.lower()])
                if not idcs:
                    continue

                # Exact product-name match + exact concrete identifier match.
                id_expr=" or ".join(f"cast({ident(col)} as text)=?" for col in idcs for _ in ids)
                params=[target_name]
                for col in idcs:
                    params.extend(ids)
                sql=f"""select {ident(codec)} product_code
                        from {ident(t)}
                        where lower(trim({ident(namec)}))=lower(trim(?))
                          and ({id_expr})
                          and coalesce(trim(cast({ident(codec)} as text)),'')<>''
                        limit 1"""
                try:
                    rr=c.execute(sql,params).fetchone()
                except Exception:
                    rr=None
                if rr:
                    code=str(rr["product_code"] or "").strip()
                    if code:
                        return code
    except Exception:
        return None
    return None


def _recent_product_code_from_catalog(row:dict):
    """Second exact-lot fallback using product_codes.sqlite3 / product_code_lots."""
    if not PRODUCT_CODES_DB.exists():
        return None
    target_name=str(row.get("product_name") or "").strip()
    if not target_name:
        return None
    ids=[]
    for k in ("lot_no","lot_number","lot_id","trade_id","deal_id","source_id"):
        v=row.get(k)
        if v not in (None,""):
            s=str(v).strip()
            if s and s not in ids: ids.append(s)
    if not ids:
        return None

    try:
        with ro(PRODUCT_CODES_DB) as c:
            if "product_code_lots" not in set(tables(c)):
                return None
            cs=cols(c,"product_code_lots")
            codec=_pick(cs,["product_code","code"])
            namec=_pick(cs,["product_name","name"])
            if not codec or not namec:
                return None
            idcs=[]
            low={x.lower():x for x in cs}
            for n in ("lot_display_no","lot_number","lot_no","lot_id","deal_id","trade_id","proc_id","source_id","id"):
                if n in low and low[n] not in idcs:
                    idcs.append(low[n])
            if not idcs:
                return None
            id_expr=" or ".join(f"cast({ident(col)} as text)=?" for col in idcs for _ in ids)
            params=[target_name]
            for col in idcs:
                params.extend(ids)
            rr=c.execute(f"""select {ident(codec)} product_code
                             from product_code_lots
                             where lower(trim({ident(namec)}))=lower(trim(?))
                               and ({id_expr})
                               and coalesce(trim(cast({ident(codec)} as text)),'')<>''
                             order by rowid desc limit 1""",params).fetchone()
            if rr:
                return str(rr["product_code"] or "").strip() or None
    except Exception:
        return None
    return None


def _catalog_codes_for_exact_product_name(name:str, limit:int=12):
    """Read product_codes.sqlite3 by product name as the final SHOP/NATIONAL fallback."""
    name=str(name or "").strip()
    if not name or not PRODUCT_CODES_DB.exists():
        return []

    target_norm=_normalize_product_query(name)
    rows=[]
    try:
        with ro(PRODUCT_CODES_DB) as c:
            ts=set(tables(c))

            if "product_codes" in ts:
                cs=cols(c,"product_codes")
                codec=_pick(cs,["product_code","code"])
                namec=_pick(cs,["product_name","name"])
                if codec and namec:
                    exact=c.execute(
                        f"""select {ident(codec)} product_code,{ident(namec)} product_name
                            from product_codes
                            where lower(trim({ident(namec)}))=lower(trim(?))
                              and coalesce(trim(cast({ident(codec)} as text)),'')<>''
                            order by rowid desc limit ?""",
                        (name,limit*3)
                    ).fetchall()
                    rows += [dict(x) for x in exact]

                    if not rows:
                        first=(name.split() or [name])[0]
                        cand=c.execute(
                            f"""select {ident(codec)} product_code,{ident(namec)} product_name
                                from product_codes
                                where lower({ident(namec)}) like lower(?)
                                  and coalesce(trim(cast({ident(codec)} as text)),'')<>''
                                order by rowid desc limit 500""",
                            (f"%{first}%",)
                        ).fetchall()
                        for x in cand:
                            d=dict(x)
                            if _normalize_product_query(d.get("product_name"))==target_norm:
                                rows.append(d)

            if "product_code_lots" in ts:
                cs=cols(c,"product_code_lots")
                codec=_pick(cs,["product_code","code"])
                namec=_pick(cs,["product_name","name"])
                if codec and namec:
                    extra=c.execute(
                        f"""select {ident(codec)} product_code,{ident(namec)} product_name
                            from product_code_lots
                            where lower(trim({ident(namec)}))=lower(trim(?))
                              and coalesce(trim(cast({ident(codec)} as text)),'')<>''
                            order by rowid desc limit ?""",
                        (name,limit*3)
                    ).fetchall()
                    rows += [dict(x) for x in extra]
    except Exception:
        return []

    out=[];seen=set()
    for r in rows:
        code=str(r.get("product_code") or "").strip()
        if not code or code in seen:
            continue
        seen.add(code);out.append(code)
        if len(out)>=limit:
            break
    return out


def _repair_recent_uzex_product_codes(rows:list[dict]):
    for r in rows:
        key=str(r.get("platform_key") or r.get("market") or "")
        if key not in {"shop","national","auction","xt_national"}:
            continue

        # Manually verified code has highest priority.
        manual_code=_manual_product_code_override(r.get("product_name"))
        if manual_code:
            r["product_code"]=manual_code
            r["product_code_repaired"]=True
            r["product_code_source"]="manual_verified"
            continue

        # Completed Auction and XT National use the central MASTER catalog next.
        if key in {"auction","xt_national"}:
            master_codes=_catalog_codes_for_exact_product_name(
                r.get("product_name"),12
            )

            master_unique=[]
            seen=set()

            for x in master_codes or []:
                code=str(x or "").strip()
                if code and code not in seen:
                    seen.add(code)
                    master_unique.append(code)

            if len(master_unique)==1:
                r["product_code"]=master_unique[0]
                r["product_code_repaired"]=True
                r["product_code_source"]="product_codes.sqlite3"
                continue

            # If MASTER is ambiguous or absent, preserve the existing source code.
            if str(r.get("product_code") or "").strip():
                continue

        else:
            # Preserve existing SHOP/NATIONAL behavior.
            if str(r.get("product_code") or "").strip():
                continue

        code=_recent_product_code_from_source(r)
        if not code:
            code=_recent_product_code_from_catalog(r)

        if not code:
            codes=_catalog_codes_for_exact_product_name(
                r.get("product_name"),12
            )
            if codes:
                code=" | ".join(codes)
                r["product_code_candidates"]=codes

        if code:
            r["product_code"]=code
            r["product_code_repaired"]=True
            r["product_code_source"]="product_codes.sqlite3" if r.get("product_code_candidates") else "exact_lot_or_source"

    return rows



def _fast_recent_stats(display_name,last_date,recent_days):
    """Exact grouped statistics for the selected recent period.

    Unlike the recent-card list, this query is NOT limited by recent_limit.
    It returns all final deals in the period grouped by trading market.
    """
    if not ANALYTICS_DB.exists():
        return {"count":0,"amount":0.0,"breakdown":[]}
    with ro(ANALYTICS_DB) as c:
        cs=cols(c,"fact_deals")
        if "product_name" not in cs or "deal_date" not in cs:
            return {"count":0,"amount":0.0,"breakdown":[]}

        base=(last_date or datetime.now().strftime("%Y-%m-%d"))[:10]
        try:
            dt=datetime.strptime(base,"%Y-%m-%d")
        except Exception:
            dt=datetime.now()
        date_from=(dt-timedelta(days=recent_days-1)).strftime("%Y-%m-%d")

        market_expr=ident("market") if "market" in cs else "''"
        amount_expr=ident("deal_amount") if "deal_amount" in cs else "0"
        final_sql="and is_final=1" if "is_final" in cs else ""

        rr=c.execute(
            f"""select {market_expr} market,
                       count(*) deals,
                       coalesce(sum({amount_expr}),0) amount
                from fact_deals
                where product_name=? and deal_date>=? {final_sql}
                  and market in ('shop','national','auction','xt_auction','xt_national')
                group by {market_expr}
                order by deals desc, market""",
            (display_name,date_from)
        ).fetchall()

        breakdown=[]
        total_count=0
        total_amount=0.0
        for x in rr:
            m=x["market"]
            cnt=int(x["deals"] or 0)
            amt=_safe_float(x["amount"])
            total_count+=cnt
            total_amount+=amt
            breakdown.append({
                "key":m,
                "title":FAST_MARKET_LABELS.get(m,m or "Р В Р вЂ Р В РІР‚С™Р Р†Р вЂљРЎСљ"),
                "source":FAST_MARKET_SOURCES.get(m,m or "Р В Р вЂ Р В РІР‚С™Р Р†Р вЂљРЎСљ"),
                "logo":FAST_MARKET_LOGOS.get(m),
                "count":cnt,
                "amount":amt,
            })
        return {
            "date_from":date_from,
            "date_to":base,
            "count":total_count,
            "amount":total_amount,
            "breakdown":breakdown,
        }



# ==================== SHOP RECENT BATCH ENRICHMENT V5 ====================
SHOP_CARD_REQUIRED_INDEXES = {
    "completed_deals": {"idx_shop_completed_deal_id", "idx_shop_completed_lot_id", "idx_shop_completed_lot_display_no"},
    "product_details": {"idx_shop_product_details_lot_id"},
    "product_properties": {"idx_shop_product_properties_record_key"},
}

def _shop_card_indexes_ready(c):
    """Never allow a product search to trigger an unindexed 2.3M-row SHOP scan."""
    try:
        ts=set(tables(c))
        for table, needed in SHOP_CARD_REQUIRED_INDEXES.items():
            if table not in ts:
                if table=="product_properties":
                    continue
                return False
            have={str(r["name"]) for r in c.execute(f"pragma index_list({ident(table)})").fetchall()}
            if not needed.issubset(have):
                return False
        return True
    except Exception:
        return False


def _shop_recent_batch_enrich(rows):
    """
    Enrich all SHOP recent rows in ONE indexed local batch while product analytics loads.
    No external HTTP/API request. If the one-time indexes are not installed, enrichment
    is skipped immediately rather than making the user wait.
    """
    if not rows:
        return rows

    shop_rows=[r for r in rows if str(r.get("platform_key") or r.get("market") or "")=="shop"]
    if not shop_rows:
        return rows

    db=PLATFORMS["shop"]["db"]
    if not db.exists():
        return rows

    def uniq(values):
        out=[]; seen=set()
        for v in values:
            if v in (None,""):
                continue
            s=str(v).strip()
            if not s or s in seen:
                continue
            seen.add(s); out.append(v)
        return out

    try:
        with ro(db) as c:
            # Critical guard: no indexes => no potentially-minute-long scan.
            if not _shop_card_indexes_ready(c):
                for r in shop_rows:
                    r["shop_card_enrichment"]="indexes_required"
                return rows

            ts=set(tables(c))
            cs=cols(c,"completed_deals")

            wanted=[
                "deal_id","lot_id","lot_display_no","deal_date","product_name","category_name",
                "amount","unit_name","unit_name_enriched","deal_cost",
                "customer_name","customer_inn","customer_region_name",
                "provider_name","provider_inn","deal_status_name"
            ]
            selected=[x for x in wanted if x in cs]

            deal_ids=[]
            for r in shop_rows:
                v=r.get("deal_id")
                if v not in (None,""):
                    try: deal_ids.append(int(str(v).strip()))
                    except Exception: pass
            deal_ids=uniq(deal_ids)

            src_rows=[]
            if deal_ids:
                for pos in range(0,len(deal_ids),400):
                    chunk=deal_ids[pos:pos+400]
                    ph=",".join("?" for _ in chunk)
                    src_rows.extend(dict(x) for x in c.execute(
                        f"select {','.join(ident(x) for x in selected)} "
                        f"from completed_deals where deal_id in ({ph})",
                        chunk
                    ).fetchall())

            # Very small fallback for any row whose deal_id was not found.
            by_deal={str(x.get("deal_id")):x for x in src_rows if x.get("deal_id") not in (None,"")}
            missing=[r for r in shop_rows if str(r.get("deal_id") or "") not in by_deal]

            if missing:
                lot_ids=[]
                for r in missing:
                    v=r.get("lot_id")
                    if v not in (None,""):
                        try: lot_ids.append(int(str(v).strip()))
                        except Exception: pass
                lot_ids=uniq(lot_ids)
                for pos in range(0,len(lot_ids),400):
                    chunk=lot_ids[pos:pos+400]
                    if not chunk: continue
                    ph=",".join("?" for _ in chunk)
                    src_rows.extend(dict(x) for x in c.execute(
                        f"select {','.join(ident(x) for x in selected)} "
                        f"from completed_deals where lot_id in ({ph})",
                        chunk
                    ).fetchall())

            by_deal={str(x.get("deal_id")):x for x in src_rows if x.get("deal_id") not in (None,"")}
            by_lot={str(x.get("lot_id")):x for x in src_rows if x.get("lot_id") not in (None,"")}
            by_display={str(x.get("lot_display_no")):x for x in src_rows if x.get("lot_display_no") not in (None,"")}

            # Identify exact source row and merge authoritative local fields.
            source_for_recent={}
            source_lot_ids=[]
            for r in shop_rows:
                s=(by_deal.get(str(r.get("deal_id") or ""))
                   or by_lot.get(str(r.get("lot_id") or ""))
                   or by_display.get(str(r.get("lot_no") or r.get("lot_number") or "")))
                if not s:
                    continue
                source_for_recent[id(r)]=s
                if s.get("lot_id") not in (None,""):
                    source_lot_ids.append(s.get("lot_id"))

                r["lot_no"]=s.get("lot_display_no") or r.get("lot_no")
                r["lot_id"]=s.get("lot_id") or r.get("lot_id")
                r["deal_id"]=s.get("deal_id") or r.get("deal_id")
                r["category_name"]=s.get("category_name") or r.get("category_name")
                r["customer"]=s.get("customer_name") or r.get("customer")
                r["customer_name"]=s.get("customer_name") or r.get("customer_name")
                r["customer_inn"]=s.get("customer_inn")
                r["region"]=s.get("customer_region_name") or r.get("region")
                r["region_name"]=s.get("customer_region_name") or r.get("region_name")
                r["supplier"]=s.get("provider_name") or r.get("supplier")
                r["provider_name"]=s.get("provider_name") or r.get("provider_name")
                r["supplier_inn"]=s.get("provider_inn")
                r["provider_inn"]=s.get("provider_inn")
                r["status"]=s.get("deal_status_name") or r.get("status") or ""
                r["source_status"]=s.get("deal_status_name") or r.get("source_status") or ""

                # Keep Product Intelligence product_code; take local quantities/prices.
                if s.get("amount") is not None:
                    r["quantity"]=s.get("amount")
                r["unit_name"]=s.get("unit_name_enriched") or s.get("unit_name") or r.get("unit_name")
                if s.get("deal_cost") is not None:
                    r["total_value"]=s.get("deal_cost")
                    r["deal_amount"]=s.get("deal_cost")
                    try:
                        q=float(s.get("amount") or 0)
                        r["unit_price"]=float(s.get("deal_cost"))/q if q else r.get("unit_price")
                    except Exception:
                        pass

            # One indexed read from local product_details for all selected SHOP lots.
            source_lot_ids=uniq(source_lot_ids)
            details={}
            record_keys=[]
            if "product_details" in ts and source_lot_ids:
                for pos in range(0,len(source_lot_ids),400):
                    chunk=source_lot_ids[pos:pos+400]
                    ph=",".join("?" for _ in chunk)
                    for x in c.execute(
                        f"select record_key,lot_id,product_description,detail_raw_json,loaded_at "
                        f"from product_details where lot_id in ({ph}) "
                        f"order by loaded_at desc",
                        chunk
                    ).fetchall():
                        d=dict(x)
                        lk=str(d.get("lot_id"))
                        if lk not in details:  # first row is newest because ORDER BY loaded_at DESC
                            details[lk]=d
                            if d.get("record_key"):
                                record_keys.append(d.get("record_key"))

            # Optional properties, also one indexed batch.
            props_by_key={}
            record_keys=uniq(record_keys)
            if "product_properties" in ts and record_keys:
                for pos in range(0,len(record_keys),300):
                    chunk=record_keys[pos:pos+300]
                    ph=",".join("?" for _ in chunk)
                    for x in c.execute(
                        f"select record_key,property_group,property_name,property_value,unit_name,sort_order "
                        f"from product_properties where record_key in ({ph}) "
                        f"order by record_key,sort_order,id",
                        chunk
                    ).fetchall():
                        d=dict(x)
                        props_by_key.setdefault(str(d.get("record_key")),[]).append(d)

            for r in shop_rows:
                s=source_for_recent.get(id(r))
                if not s: continue
                det=details.get(str(s.get("lot_id")))
                if det:
                    obj={}
                    raw=det.get("detail_raw_json")
                    if raw:
                        try: obj=json.loads(raw) if isinstance(raw,str) else dict(raw)
                        except Exception: obj={}
                    desc=(det.get("product_description")
                          or obj.get("condition")
                          or obj.get("description")
                          or obj.get("product_description"))
                    if desc:
                        r["description"]=desc
                    if not r.get("category_name") and obj.get("category_name"):
                        r["category_name"]=obj.get("category_name")
                    rk=det.get("record_key")
                    if rk and props_by_key.get(str(rk)):
                        r["properties_json"]=props_by_key[str(rk)]

                r["shop_card_enrichment"]="ready"

        return rows
    except Exception as e:
        # Product analytics must remain usable even if optional enrichment fails.
        for r in shop_rows:
            r["shop_card_enrichment"]="error"
            r["shop_card_enrichment_error"]=str(e)
        return rows


# ================== /SHOP RECENT BATCH ENRICHMENT V5 ==================


# ==================== NATIONAL RECENT BATCH ENRICHMENT V7 ====================
NATIONAL_CARD_REQUIRED_INDEXES = {
    "completed_deals": {
        "idx_national_completed_deal_id",
        "idx_national_completed_lot_id",
        "idx_national_completed_lot_display_no",
    },
}

def _national_card_indexes_ready(c):
    """Avoid any unindexed scan while Product Intelligence is loading."""
    try:
        ts=set(tables(c))
        if "completed_deals" not in ts:
            return False
        have={str(r["name"]) for r in c.execute(
            f"pragma index_list({ident('completed_deals')})"
        ).fetchall()}
        return NATIONAL_CARD_REQUIRED_INDEXES["completed_deals"].issubset(have)
    except Exception:
        return False


def _national_recent_batch_enrich(rows):
    """
    Enrich all xarid.uzex.uz Р В Р вЂ Р В РІР‚С™Р Р†Р вЂљРЎСљ Р В Р’В Р РЋРЎв„ўР В Р’В Р РЋРІР‚ВР В Р’В Р вЂ™Р’В»Р В Р’В Р вЂ™Р’В»Р В Р’В Р РЋРІР‚ВР В Р’В Р Р†РІР‚С›РІР‚вЂњ Р В Р’В Р СћРІР‚ВР В Р Р‹Р РЋРІР‚С”Р В Р’В Р РЋРІР‚СњР В Р’В Р РЋРІР‚СћР В Р’В Р В РІР‚В¦ recent cards in ONE local indexed batch.
    No external HTTP/API call. Card click itself performs no DB request.
    """
    if not rows:
        return rows

    national_rows=[
        r for r in rows
        if str(r.get("platform_key") or r.get("market") or "")=="national"
    ]
    if not national_rows:
        return rows

    db=PLATFORMS["national"]["db"]
    if not db.exists():
        return rows

    def uniq(values):
        out=[]; seen=set()
        for v in values:
            if v in (None,""):
                continue
            s=str(v).strip()
            if not s or s in seen:
                continue
            seen.add(s); out.append(v)
        return out

    try:
        with ro(db) as c:
            # Never make the user wait on an unindexed historical scan.
            if not _national_card_indexes_ready(c):
                for r in national_rows:
                    r["national_card_enrichment"]="indexes_required"
                return rows

            ts=set(tables(c))
            cs=cols(c,"completed_deals")

            wanted=[
                "deal_id","lot_id","lot_display_no","deal_date",
                "product_name","category_name","amount",
                "unit_name","unit_name_enriched","deal_cost",
                "customer_name","customer_inn","customer_region_name",
                "provider_name","provider_inn","deal_status_name"
            ]
            selected=[x for x in wanted if x in cs]

            # First: exact deal_id batch.
            deal_ids=[]
            for r in national_rows:
                v=r.get("deal_id")
                if v not in (None,""):
                    try: deal_ids.append(int(str(v).strip()))
                    except Exception: pass
            deal_ids=uniq(deal_ids)

            src_rows=[]
            for pos in range(0,len(deal_ids),400):
                chunk=deal_ids[pos:pos+400]
                if not chunk:
                    continue
                ph=",".join("?" for _ in chunk)
                src_rows.extend(dict(x) for x in c.execute(
                    f"select {','.join(ident(x) for x in selected)} "
                    f"from completed_deals where deal_id in ({ph})",
                    chunk
                ).fetchall())

            # Fallback: lot_id only for unresolved rows.
            by_deal={
                str(x.get("deal_id")):x for x in src_rows
                if x.get("deal_id") not in (None,"")
            }
            missing=[
                r for r in national_rows
                if str(r.get("deal_id") or "") not in by_deal
            ]

            lot_ids=[]
            for r in missing:
                v=r.get("lot_id")
                if v not in (None,""):
                    try: lot_ids.append(int(str(v).strip()))
                    except Exception: pass
            lot_ids=uniq(lot_ids)

            for pos in range(0,len(lot_ids),400):
                chunk=lot_ids[pos:pos+400]
                if not chunk:
                    continue
                ph=",".join("?" for _ in chunk)
                src_rows.extend(dict(x) for x in c.execute(
                    f"select {','.join(ident(x) for x in selected)} "
                    f"from completed_deals where lot_id in ({ph})",
                    chunk
                ).fetchall())

            by_deal={
                str(x.get("deal_id")):x for x in src_rows
                if x.get("deal_id") not in (None,"")
            }
            by_lot={
                str(x.get("lot_id")):x for x in src_rows
                if x.get("lot_id") not in (None,"")
            }
            by_display={
                str(x.get("lot_display_no")):x for x in src_rows
                if x.get("lot_display_no") not in (None,"")
            }

            source_for_recent={}
            source_lot_ids=[]

            for r in national_rows:
                s=(by_deal.get(str(r.get("deal_id") or ""))
                   or by_lot.get(str(r.get("lot_id") or ""))
                   or by_display.get(str(r.get("lot_no") or r.get("lot_number") or "")))
                if not s:
                    continue

                source_for_recent[id(r)]=s
                if s.get("lot_id") not in (None,""):
                    source_lot_ids.append(s.get("lot_id"))

                r["lot_no"]=s.get("lot_display_no") or r.get("lot_no")
                r["lot_id"]=s.get("lot_id") or r.get("lot_id")
                r["deal_id"]=s.get("deal_id") or r.get("deal_id")
                r["category_name"]=s.get("category_name") or r.get("category_name")

                r["customer"]=s.get("customer_name") or r.get("customer")
                r["customer_name"]=s.get("customer_name") or r.get("customer_name")
                r["customer_inn"]=s.get("customer_inn")

                r["region"]=s.get("customer_region_name") or r.get("region")
                r["region_name"]=s.get("customer_region_name") or r.get("region_name")

                r["supplier"]=s.get("provider_name") or r.get("supplier")
                r["provider_name"]=s.get("provider_name") or r.get("provider_name")
                r["supplier_inn"]=s.get("provider_inn")
                r["provider_inn"]=s.get("provider_inn")

                r["status"]=s.get("deal_status_name") or r.get("status") or ""
                r["source_status"]=s.get("deal_status_name") or r.get("source_status") or ""

                if s.get("amount") is not None:
                    r["quantity"]=s.get("amount")

                r["unit_name"]=(
                    s.get("unit_name_enriched")
                    or s.get("unit_name")
                    or r.get("unit_name")
                )

                if s.get("deal_cost") is not None:
                    r["total_value"]=s.get("deal_cost")
                    r["deal_amount"]=s.get("deal_cost")
                    try:
                        q=float(s.get("amount") or 0)
                        r["unit_price"]=(
                            float(s.get("deal_cost"))/q if q else r.get("unit_price")
                        )
                    except Exception:
                        pass

            # National completed DB may also contain local details/properties.
            # Use them when present; otherwise keep fact_deals description/properties.
            source_lot_ids=uniq(source_lot_ids)
            details={}
            record_keys=[]

            if "product_details" in ts and source_lot_ids:
                dcs=cols(c,"product_details")
                if "lot_id" in dcs:
                    select_cols=[
                        x for x in (
                            "record_key","lot_id","product_description",
                            "detail_raw_json","loaded_at"
                        )
                        if x in dcs
                    ]
                    if "lot_id" in select_cols:
                        for pos in range(0,len(source_lot_ids),400):
                            chunk=source_lot_ids[pos:pos+400]
                            ph=",".join("?" for _ in chunk)
                            order=" order by loaded_at desc" if "loaded_at" in dcs else ""
                            for x in c.execute(
                                f"select {','.join(ident(x) for x in select_cols)} "
                                f"from product_details where lot_id in ({ph}){order}",
                                chunk
                            ).fetchall():
                                d=dict(x)
                                lk=str(d.get("lot_id"))
                                if lk not in details:
                                    details[lk]=d
                                    if d.get("record_key"):
                                        record_keys.append(d.get("record_key"))

            props_by_key={}
            record_keys=uniq(record_keys)
            if "product_properties" in ts and record_keys:
                pcs=cols(c,"product_properties")
                if "record_key" in pcs:
                    select_cols=[
                        x for x in (
                            "record_key","property_group","property_name",
                            "property_value","unit_name","sort_order","id"
                        )
                        if x in pcs
                    ]
                    order_parts=[
                        x for x in ("record_key","sort_order","id")
                        if x in pcs
                    ]
                    order=(" order by "+",".join(order_parts)) if order_parts else ""
                    for pos in range(0,len(record_keys),300):
                        chunk=record_keys[pos:pos+300]
                        ph=",".join("?" for _ in chunk)
                        for x in c.execute(
                            f"select {','.join(ident(x) for x in select_cols)} "
                            f"from product_properties where record_key in ({ph}){order}",
                            chunk
                        ).fetchall():
                            d=dict(x)
                            props_by_key.setdefault(
                                str(d.get("record_key")),[]
                            ).append(d)

            for r in national_rows:
                s=source_for_recent.get(id(r))
                if not s:
                    continue

                det=details.get(str(s.get("lot_id")))
                if det:
                    obj={}
                    raw=det.get("detail_raw_json")
                    if raw:
                        try:
                            obj=json.loads(raw) if isinstance(raw,str) else dict(raw)
                        except Exception:
                            obj={}

                    desc=(
                        det.get("product_description")
                        or obj.get("condition")
                        or obj.get("description")
                        or obj.get("product_description")
                    )
                    if desc:
                        r["description"]=desc

                    if not r.get("category_name") and obj.get("category_name"):
                        r["category_name"]=obj.get("category_name")

                    rk=det.get("record_key")
                    if rk and props_by_key.get(str(rk)):
                        r["properties_json"]=props_by_key[str(rk)]

                r["national_card_enrichment"]="ready"

        return rows

    except Exception as e:
        for r in national_rows:
            r["national_card_enrichment"]="error"
            r["national_card_enrichment_error"]=str(e)
        return rows


# ================== /NATIONAL RECENT BATCH ENRICHMENT V7 ==================

def _fast_recent(display_name,last_date,recent_days,recent_limit):
    if not ANALYTICS_DB.exists(): return []
    with ro(ANALYTICS_DB) as c:
        cs=cols(c,"fact_deals")
        fields=[("market","market"),("deal_id","deal_id"),("trade_id","trade_id"),("lot_id","lot_id"),("lot_number","lot_number"),
        ("deal_date","deal_date"),("product_name","product_name"),("product_code","product_code"),("quantity","quantity"),("unit_name","unit_name"),("unit_price","unit_price"),
        ("deal_amount","deal_amount"),
        ("category_name","category_name"),
        ("customer_name","customer_name"),("customer_inn","customer_inn"),
        ("region_name","region_name"),
        ("provider_name","provider_name"),("provider_inn","provider_inn"),
        ("description","description"),("properties_json","properties_json"),("attributes_json","attributes_json"),("source_status","source_status")]
        sel=[f"{ident(a)} {ident(b)}" if a in cs else f"null {ident(b)}" for a,b in fields]
        base=(last_date or datetime.now().strftime("%Y-%m-%d"))[:10]
        try:dt=datetime.strptime(base,"%Y-%m-%d")
        except:dt=datetime.now()
        date_from=(dt-timedelta(days=recent_days-1)).strftime("%Y-%m-%d")
        final_sql="and is_final=1" if "is_final" in cs else ""
        rr=c.execute(f"""select {','.join(sel)} from fact_deals where product_name=? and deal_date>=? {final_sql}
          and market in ('shop','national','auction','xt_auction','xt_national')
        order by deal_date desc limit ?""",(display_name,date_from,recent_limit)).fetchall()
        out=[]
        for x in rr:
            r=dict(x);m=r.get("market")
            r.update({"platform_key":m,"platform":FAST_MARKET_LABELS.get(m,m or "Р В Р вЂ Р В РІР‚С™Р Р†Р вЂљРЎСљ"),
            "lot_no":r.get("lot_number") or r.get("lot_id") or r.get("trade_id") or r.get("deal_id"),
            "date":r.get("deal_date"),"total_value":r.get("deal_amount"),
            "category":r.get("category_name"),
            "customer":r.get("customer_name"),
            "customer_inn":r.get("customer_inn"),
            "supplier":r.get("provider_name"),
            "supplier_inn":r.get("provider_inn"),
            "region":r.get("region_name"),
            "status":r.get("source_status") or ""})
            out.append(r)
        return _national_recent_batch_enrich(_shop_recent_batch_enrich(_repair_xt_auction_recent(_repair_recent_uzex_product_codes(out))))



def _median_num(values):
    vals=sorted(float(x) for x in values if x not in (None, "") and _safe_float(x)>0)
    if not vals:return None
    n=len(vals); mid=n//2
    return vals[mid] if n%2 else (vals[mid-1]+vals[mid])/2.0

def _percentile_num(values,pct):
    vals=sorted(float(x) for x in values if x not in (None, "") and _safe_float(x)>0)
    if not vals:return None
    if len(vals)==1:return vals[0]
    pos=(len(vals)-1)*float(pct)
    lo=int(pos); hi=min(lo+1,len(vals)-1); frac=pos-lo
    return vals[lo]*(1-frac)+vals[hi]*frac

def _product_decision_v2(display_name,last_date,seasonality):
    """Deep, explainable product BI signals from local fact_deals only.

    No external API. All windows are anchored to the selected product's latest deal date
    so historical products are evaluated against their own data freshness.
    """
    empty={
      "as_of":(last_date or "")[:10],
      "demand":{},"price":{},"regions":[],"season":{},"capacity":{},
      "data_quality":{"source":"uzbi_analytics.fact_deals","exact_product_name":True}
    }
    if not ANALYTICS_DB.exists():return empty
    base=(last_date or datetime.now().strftime("%Y-%m-%d"))[:10]
    try:dt=datetime.strptime(base,"%Y-%m-%d")
    except Exception:dt=datetime.now();base=dt.strftime("%Y-%m-%d")
    def ds(days_back):return (dt-timedelta(days=days_back-1)).strftime("%Y-%m-%d")
    def end_exclusive(end_date):
        return (datetime.strptime(end_date,"%Y-%m-%d")+timedelta(days=1)).strftime("%Y-%m-%d")
    def prev_start(days):return (dt-timedelta(days=days*2-1)).strftime("%Y-%m-%d")
    def prev_end(days):return (dt-timedelta(days=days)).strftime("%Y-%m-%d")
    with ro(ANALYTICS_DB) as c:
        cs=cols(c,"fact_deals")
        if "product_name" not in cs or "deal_date" not in cs:return empty
        final_sql="and is_final=1" if "is_final" in cs else ""
        amount_expr=ident("deal_amount") if "deal_amount" in cs else "0"
        price_expr=ident("unit_price") if "unit_price" in cs else "null"
        region_expr=ident("region_name") if "region_name" in cs else "null"
        def agg(start,end):
            r=c.execute(f"""select count(*) cnt,coalesce(sum({amount_expr}),0) amount
              from fact_deals where product_name=? and deal_date>=? and deal_date<? {final_sql}
              and market in ('shop','national','auction','xt_auction','xt_national')""",(display_name,start,end_exclusive(end))).fetchone()
            return {"count":int(r[0] or 0),"amount":_safe_float(r[1])}
        w7=agg(ds(7),base);p7=agg(prev_start(7),prev_end(7))
        w30=agg(ds(30),base);p30=agg(prev_start(30),prev_end(30))
        w90=agg(ds(90),base);p90=agg(prev_start(90),prev_end(90))
        def change(cur,prev):
            return ((cur-prev)/prev*100.0) if prev else (100.0 if cur else None)
        demand={
          "last7":w7,"prev7":p7,"change7_pct":change(w7["count"],p7["count"]),
          "last30":w30,"prev30":p30,"change30_pct":change(w30["count"],p30["count"]),
          "last90":w90,"prev90":p90,"change90_pct":change(w90["count"],p90["count"]),
          "daily_velocity30":w30["count"]/30.0,"daily_velocity90":w90["count"]/90.0
        }
        # Exact unit-price distributions for recent and comparison windows.
        prices={}
        if "unit_price" in cs:
            def price_vals(start,end):
                rr=c.execute(f"""select {price_expr} from fact_deals where product_name=? and deal_date>=? and deal_date<?
                  and {price_expr} is not null and {price_expr}>0 {final_sql}
                  and market in ('shop','national','auction','xt_auction','xt_national')""",(display_name,start,end_exclusive(end))).fetchall()
                return [r[0] for r in rr]
            v30=price_vals(ds(30),base); vp30=price_vals(prev_start(30),prev_end(30)); v90=price_vals(ds(90),base)
            m30=_median_num(v30);mp30=_median_num(vp30);m90=_median_num(v90)
            prices={
              "median30":m30,"median_prev30":mp30,"median90":m90,
              "change30_pct":((m30-mp30)/mp30*100.0) if (m30 and mp30) else None,
              "p25_90":_percentile_num(v90,.25),"p75_90":_percentile_num(v90,.75),
              "min90":min(v90) if v90 else None,"max90":max(v90) if v90 else None,
              "samples30":len(v30),"samples90":len(v90)
            }
        # Top demand geography for the latest 90-day window.
        # Normalize Russian and Uzbek region spellings into one Uzbek Cyrillic BI standard.
        # share_pct is against ALL 90-day deals, not only rows with a known region.
        regions=[]
        if "region_name" in cs:
            region_map={
              "\u0414\u0436\u0438\u0437\u0430\u043a\u0441\u043a\u0430\u044f \u043e\u0431\u043b\u0430\u0441\u0442\u044c":"\u0416\u0438\u0437\u0437\u0430\u0445 \u0432\u0438\u043b\u043e\u044f\u0442\u0438",
              "\u0416\u0438\u0437\u0437\u0430\u0445 \u0432\u0438\u043b\u043e\u044f\u0442\u0438":"\u0416\u0438\u0437\u0437\u0430\u0445 \u0432\u0438\u043b\u043e\u044f\u0442\u0438",

              "\u0411\u0443\u0445\u0430\u0440\u0441\u043a\u0430\u044f \u043e\u0431\u043b\u0430\u0441\u0442\u044c":"\u0411\u0443\u0445\u043e\u0440\u043e \u0432\u0438\u043b\u043e\u044f\u0442\u0438",
              "\u0411\u0443\u0445\u043e\u0440\u043e \u0432\u0438\u043b\u043e\u044f\u0442\u0438":"\u0411\u0443\u0445\u043e\u0440\u043e \u0432\u0438\u043b\u043e\u044f\u0442\u0438",

              "\u0433\u043e\u0440\u043e\u0434 \u0422\u0430\u0448\u043a\u0435\u043d\u0442":"\u0422\u043e\u0448\u043a\u0435\u043d\u0442 \u0448\u0430\u04b3\u0440\u0438",
              "\u0422\u043e\u0448\u043a\u0435\u043d\u0442 \u0448\u0430\u04b3\u0440\u0438":"\u0422\u043e\u0448\u043a\u0435\u043d\u0442 \u0448\u0430\u04b3\u0440\u0438",

              "\u0422\u0430\u0448\u043a\u0435\u043d\u0442\u0441\u043a\u0430\u044f \u043e\u0431\u043b\u0430\u0441\u0442\u044c":"\u0422\u043e\u0448\u043a\u0435\u043d\u0442 \u0432\u0438\u043b\u043e\u044f\u0442\u0438",
              "\u0422\u043e\u0448\u043a\u0435\u043d\u0442 \u0432\u0438\u043b\u043e\u044f\u0442\u0438":"\u0422\u043e\u0448\u043a\u0435\u043d\u0442 \u0432\u0438\u043b\u043e\u044f\u0442\u0438",

              "\u0421\u0430\u043c\u0430\u0440\u043a\u0430\u043d\u0434\u0441\u043a\u0430\u044f \u043e\u0431\u043b\u0430\u0441\u0442\u044c":"\u0421\u0430\u043c\u0430\u0440\u049b\u0430\u043d\u0434 \u0432\u0438\u043b\u043e\u044f\u0442\u0438",
              "\u0421\u0430\u043c\u0430\u0440\u049b\u0430\u043d\u0434 \u0432\u0438\u043b\u043e\u044f\u0442\u0438":"\u0421\u0430\u043c\u0430\u0440\u049b\u0430\u043d\u0434 \u0432\u0438\u043b\u043e\u044f\u0442\u0438",

              "\u0410\u043d\u0434\u0438\u0436\u0430\u043d\u0441\u043a\u0430\u044f \u043e\u0431\u043b\u0430\u0441\u0442\u044c":"\u0410\u043d\u0434\u0438\u0436\u043e\u043d \u0432\u0438\u043b\u043e\u044f\u0442\u0438",
              "\u0410\u043d\u0434\u0438\u0436\u043e\u043d \u0432\u0438\u043b\u043e\u044f\u0442\u0438":"\u0410\u043d\u0434\u0438\u0436\u043e\u043d \u0432\u0438\u043b\u043e\u044f\u0442\u0438",

              "\u0421\u0443\u0440\u0445\u0430\u043d\u0434\u0430\u0440\u044c\u0438\u043d\u0441\u043a\u0430\u044f \u043e\u0431\u043b\u0430\u0441\u0442\u044c":"\u0421\u0443\u0440\u0445\u043e\u043d\u0434\u0430\u0440\u0451 \u0432\u0438\u043b\u043e\u044f\u0442\u0438",
              "\u0421\u0443\u0440\u0445\u043e\u043d\u0434\u0430\u0440\u0451 \u0432\u0438\u043b\u043e\u044f\u0442\u0438":"\u0421\u0443\u0440\u0445\u043e\u043d\u0434\u0430\u0440\u0451 \u0432\u0438\u043b\u043e\u044f\u0442\u0438",

              "\u041a\u0430\u0448\u043a\u0430\u0434\u0430\u0440\u044c\u0438\u043d\u0441\u043a\u0430\u044f \u043e\u0431\u043b\u0430\u0441\u0442\u044c":"\u049a\u0430\u0448\u049b\u0430\u0434\u0430\u0440\u0451 \u0432\u0438\u043b\u043e\u044f\u0442\u0438",
              "\u049a\u0430\u0448\u049b\u0430\u0434\u0430\u0440\u0451 \u0432\u0438\u043b\u043e\u044f\u0442\u0438":"\u049a\u0430\u0448\u049b\u0430\u0434\u0430\u0440\u0451 \u0432\u0438\u043b\u043e\u044f\u0442\u0438",

              "\u0424\u0435\u0440\u0433\u0430\u043d\u0441\u043a\u0430\u044f \u043e\u0431\u043b\u0430\u0441\u0442\u044c":"\u0424\u0430\u0440\u0493\u043e\u043d\u0430 \u0432\u0438\u043b\u043e\u044f\u0442\u0438",
              "\u0424\u0430\u0440\u0493\u043e\u043d\u0430 \u0432\u0438\u043b\u043e\u044f\u0442\u0438":"\u0424\u0430\u0440\u0493\u043e\u043d\u0430 \u0432\u0438\u043b\u043e\u044f\u0442\u0438",

              "\u0425\u043e\u0440\u0435\u0437\u043c\u0441\u043a\u0430\u044f \u043e\u0431\u043b\u0430\u0441\u0442\u044c":"\u0425\u043e\u0440\u0430\u0437\u043c \u0432\u0438\u043b\u043e\u044f\u0442\u0438",
              "\u0425\u043e\u0440\u0430\u0437\u043c \u0432\u0438\u043b\u043e\u044f\u0442\u0438":"\u0425\u043e\u0440\u0430\u0437\u043c \u0432\u0438\u043b\u043e\u044f\u0442\u0438",

              "\u041d\u0430\u043c\u0430\u043d\u0433\u0430\u043d\u0441\u043a\u0430\u044f \u043e\u0431\u043b\u0430\u0441\u0442\u044c":"\u041d\u0430\u043c\u0430\u043d\u0433\u0430\u043d \u0432\u0438\u043b\u043e\u044f\u0442\u0438",
              "\u041d\u0430\u043c\u0430\u043d\u0433\u0430\u043d \u0432\u0438\u043b\u043e\u044f\u0442\u0438":"\u041d\u0430\u043c\u0430\u043d\u0433\u0430\u043d \u0432\u0438\u043b\u043e\u044f\u0442\u0438",

              "\u041d\u0430\u0432\u043e\u0438\u0439\u0441\u043a\u0430\u044f \u043e\u0431\u043b\u0430\u0441\u0442\u044c":"\u041d\u0430\u0432\u043e\u0438\u0439 \u0432\u0438\u043b\u043e\u044f\u0442\u0438",
              "\u041d\u0430\u0432\u043e\u0438\u0439 \u0432\u0438\u043b\u043e\u044f\u0442\u0438":"\u041d\u0430\u0432\u043e\u0438\u0439 \u0432\u0438\u043b\u043e\u044f\u0442\u0438",

              "\u0421\u044b\u0440\u0434\u0430\u0440\u044c\u0438\u043d\u0441\u043a\u0430\u044f \u043e\u0431\u043b\u0430\u0441\u0442\u044c":"\u0421\u0438\u0440\u0434\u0430\u0440\u0451 \u0432\u0438\u043b\u043e\u044f\u0442\u0438",
              "\u0421\u0438\u0440\u0434\u0430\u0440\u0451 \u0432\u0438\u043b\u043e\u044f\u0442\u0438":"\u0421\u0438\u0440\u0434\u0430\u0440\u0451 \u0432\u0438\u043b\u043e\u044f\u0442\u0438",

              "\u0420\u0435\u0441\u043f\u0443\u0431\u043b\u0438\u043a\u0430 \u041a\u0430\u0440\u0430\u043a\u0430\u043b\u043f\u0430\u043a\u0441\u0442\u0430\u043d":"\u049a\u043e\u0440\u0430\u049b\u0430\u043b\u043f\u043e\u0493\u0438\u0441\u0442\u043e\u043d \u0420\u0435\u0441\u043f\u0443\u0431\u043b\u0438\u043a\u0430\u0441\u0438",
              "\u049a\u043e\u0440\u0430\u049b\u0430\u043b\u043f\u043e\u0493\u0438\u0441\u0442\u043e\u043d \u0420\u0435\u0441\u043f\u0443\u0431\u043b\u0438\u043a\u0430\u0441\u0438":"\u049a\u043e\u0440\u0430\u049b\u0430\u043b\u043f\u043e\u0493\u0438\u0441\u0442\u043e\u043d \u0420\u0435\u0441\u043f\u0443\u0431\u043b\u0438\u043a\u0430\u0441\u0438",

              "\u0420\u0435\u0441\u043f\u0443\u0431\u043b\u0438\u043a\u0430 \u0423\u0437\u0431\u0435\u043a\u0438\u0441\u0442\u0430\u043d":"\u04b2\u0443\u0434\u0443\u0434 \u0430\u043d\u0438\u049b\u043b\u0430\u043d\u043c\u0430\u0433\u0430\u043d"
            }

            rr=c.execute(f"""select {region_expr} region,count(*) cnt,coalesce(sum({amount_expr}),0) amount
              from fact_deals where product_name=? and deal_date>=? and deal_date<? {final_sql}
              and market in ('shop','national','auction','xt_auction','xt_national')
              and coalesce(trim({region_expr}),'')<>''
              group by {region_expr}""",
              (display_name,ds(90),end_exclusive(base))).fetchall()

            grouped_regions={}
            for r in rr:
                raw_region=str(r[0] or "").strip()
                canonical=region_map.get(raw_region,raw_region)
                z=grouped_regions.setdefault(canonical,{"count":0,"amount":0.0})
                z["count"]+=int(r[1] or 0)
                z["amount"]+=_safe_float(r[2])

            total=max(0,int(w90.get("count") or 0))
            regions=[
              {
                "region":name,
                "count":int(z["count"]),
                "amount":_safe_float(z["amount"]),
                "share_pct":(int(z["count"])/total*100.0 if total else 0)
              }
              for name,z in sorted(
                grouped_regions.items(),
                key=lambda item:(-int(item[1]["count"]),item[0])
              )[:8]
            ]

        # TOP-5 MULTI-PLATFORM V2.2
        # Explicit scope: all five completed product trading directions.
        participant_markets=("shop","national","auction","xt_auction","xt_national")
        market_placeholders=",".join("?" for _ in participant_markets)
        market_expr=ident("market") if "market" in cs else "null"
        customer_name_expr=ident("customer_name") if "customer_name" in cs else "null"
        customer_inn_expr=ident("customer_inn") if "customer_inn" in cs else "null"
        provider_name_expr=ident("provider_name") if "provider_name" in cs else "null"
        provider_inn_expr=ident("provider_inn") if "provider_inn" in cs else "null"

        def _participant_norm_name(v):
            s=str(v or "").strip().casefold().replace("Р В Р Р‹Р Р†Р вЂљР’В","Р В Р’В Р вЂ™Р’Вµ")
            s=re.sub(r"[\"'`Р В Р вЂ Р В РІР‚С™Р вЂ™Р’ВР В Р вЂ Р В РІР‚С™Р Р†РІР‚С›РЎС›Р В Р вЂ Р В РІР‚С™Р РЋРЎв„ўР В Р вЂ Р В РІР‚С™Р РЋРЎС™Р В РІР‚в„ўР вЂ™Р’В«Р В РІР‚в„ўР вЂ™Р’В»]+"," ",s)
            s=re.sub(r"[^0-9a-zР В Р’В Р вЂ™Р’В°-Р В Р Р‹Р В Р РЏР В Р Р‹Р Р†Р вЂљР’ВР В Р Р‹Р РЋРІР‚С”Р В РЎС›Р Р†Р вЂљРЎвЂќР В РЎС›Р Р†Р вЂљРЎС™Р В РЎС›Р РЋРІР‚вЂњР В Р Р‹Р РЋРІР‚С”\- ]+"," ",s,flags=re.I)
            s=re.sub(r"\s+"," ",s).strip()
            return s

        def _participant_norm_inn(v):
            s=re.sub(r"\D+","",str(v or ""))
            return s if len(s)>=6 else ""

        participant_sql=f"""select {market_expr} market,
                                   {customer_name_expr} customer_name,{customer_inn_expr} customer_inn,
                                   {provider_name_expr} provider_name,{provider_inn_expr} provider_inn,
                                   {amount_expr} deal_amount
                            from fact_deals
                            where product_name=? and deal_date>=? and deal_date<? {final_sql}
                              and {market_expr} in ({market_placeholders})"""
        participant_rows=c.execute(
            participant_sql,
            (display_name,ds(90),end_exclusive(base),*participant_markets)
        ).fetchall()

        market_counts={m:0 for m in participant_markets}
        market_amounts={m:0.0 for m in participant_markets}
        for rr in participant_rows:
            m=str(rr["market"] or "")
            if m in market_counts:
                market_counts[m]+=1
                market_amounts[m]+=_safe_float(rr["deal_amount"])
        all5_count=sum(market_counts.values())
        all5_amount=sum(market_amounts.values())

        def top_participants_multi(kind):
            is_customer=(kind=="customer")
            name_key="customer_name" if is_customer else "provider_name"
            inn_key="customer_inn" if is_customer else "provider_inn"
            grouped={}
            for rr in participant_rows:
                name=str(rr[name_key] or "").strip()
                inn=_participant_norm_inn(rr[inn_key])
                norm=_participant_norm_name(name)
                if not name and not inn:
                    continue
                key=("inn",inn) if inn else ("name",norm)
                if not key[1]:
                    continue
                z=grouped.setdefault(key,{
                    "name":name or inn,"inn":inn or None,"count":0,"amount":0.0,
                    "platforms":{},"name_variants":set()
                })
                if name:
                    z["name_variants"].add(name)
                    if len(name)>len(str(z.get("name") or "")):
                        z["name"]=name
                z["count"]+=1
                z["amount"]+=_safe_float(rr["deal_amount"])
                m=str(rr["market"] or "")
                pm=z["platforms"].setdefault(m,{"count":0,"amount":0.0})
                pm["count"]+=1
                pm["amount"]+=_safe_float(rr["deal_amount"])
            arr=[]
            for z in grouped.values():
                plist=[]
                for m in participant_markets:
                    pm=z["platforms"].get(m)
                    if pm:
                        plist.append({
                            "key":m,"title":FAST_MARKET_LABELS.get(m,m),
                            "source":FAST_MARKET_SOURCES.get(m,m),
                            "count":int(pm["count"]),"amount":_safe_float(pm["amount"])
                        })
                z["platforms"]=plist
                z["platform_count"]=len(plist)
                z["share_pct"]=(z["count"]/all5_count*100.0 if all5_count else 0.0)
                z["amount_share_pct"]=(z["amount"]/all5_amount*100.0 if all5_amount else 0.0)
                z["name_variants_count"]=len(z.pop("name_variants",set()))
                arr.append(z)
            arr.sort(key=lambda x:(-int(x["count"]),-_safe_float(x["amount"]),str(x["name"])))
            return arr[:5]

        top_customers=top_participants_multi("customer")
        top_suppliers=top_participants_multi("supplier")
        customer_top5_share=sum(float(x.get("share_pct") or 0) for x in top_customers)
        supplier_top5_share=sum(float(x.get("share_pct") or 0) for x in top_suppliers)
        customer_top5_amount_share=sum(float(x.get("amount_share_pct") or 0) for x in top_customers)
        supplier_top5_amount_share=sum(float(x.get("amount_share_pct") or 0) for x in top_suppliers)
        participant_platforms=[
            {"key":m,"title":FAST_MARKET_LABELS.get(m,m),"source":FAST_MARKET_SOURCES.get(m,m),
             "count":int(market_counts.get(m,0)),"amount":_safe_float(market_amounts.get(m,0))}
            for m in participant_markets
        ]
        # Season profile already aggregates all historical years by calendar month.
        # Historical seasonality uses the normalized annual SHOP share.
        # `deals` is retained as a compatible field, while `share_pct`
        # is the actual seasonality signal.
        sm={int(x.get("month") or 0):float(x.get("share_pct") if x.get("share_pct") is not None else 0)
            for x in (seasonality or [])}
        cur=dt.month;n1=1 if cur==12 else cur+1;n2=((cur+1)%12)+1
        maxv=max(sm.values()) if sm else 0
        strongest=sorted(
            ([{"month":m,"share_pct":round(v,2)} for m,v in sm.items()]),
            key=lambda x:-x["share_pct"]
        )[:3]
        season={
          "current_month":cur,
          "current_deals":sm.get(cur,0),
          "current_share_pct":sm.get(cur,0),
          "current_strength_pct":(sm.get(cur,0)/maxv*100.0 if maxv else None),
          "next_month":n1,
          "next_month_deals":sm.get(n1,0),
          "next_month_share_pct":sm.get(n1,0),
          "next_month_strength_pct":(sm.get(n1,0)/maxv*100.0 if maxv else None),
          "next2_month":n2,
          "next2_month_deals":sm.get(n2,0),
          "next2_month_share_pct":sm.get(n2,0),
          "next2_month_strength_pct":(sm.get(n2,0)/maxv*100.0 if maxv else None),
          "strongest_months":strongest
        }
        count90=w90["count"]
        capacity={"level":"high" if count90>=300 else "medium" if count90>=60 else "niche","deals90":count90,"amount90":w90["amount"]}
        # Explainability / confidence and price anomaly flags.
        samples90=int(prices.get("samples90") or 0) if prices else 0
        if count90>=60 and samples90>=30:
            confidence={
                "level":"high",
                "label":"?????",
                "reason":"90 ?????? ???????? ?? ???? ?????????? ??????"
            }
        else:
            confidence={
                "level":"low",
                "label":"????",
                "reason":"??????? ?????? ???? ??????????? ?????? ????"
            }
        price_change=prices.get("change30_pct") if prices else None
        price_anomaly=bool(price_change is not None and abs(float(price_change))>=35)
        return {"as_of":base,"demand":demand,"price":prices,"regions":regions,"season":season,"capacity":capacity,
                "top_customers":top_customers,"top_suppliers":top_suppliers,
                "concentration":{"customers_top5_share_pct":customer_top5_share,
                                 "suppliers_top5_share_pct":supplier_top5_share,
                                 "customers_top5_amount_share_pct":customer_top5_amount_share,
                                 "suppliers_top5_amount_share_pct":supplier_top5_amount_share},
                "participant_scope":{"window_days":90,"markets":participant_platforms,
                                     "market_keys":list(participant_markets),
                                     "deals_count":all5_count,"amount":all5_amount,
                                     "merge_key":"INN/STIR first; normalized name fallback"},
                "confidence":confidence,"price_anomaly":price_anomaly,
                "data_quality":{"source":"uzbi_analytics.fact_deals","exact_product_name":True,"final_only":"is_final" in cs,
                                "top5_scope":"all_5_product_markets","participant_merge":"inn_then_normalized_name"}}

@app.get("/api/product-fast/{product_id}")
def product_fast_intelligence(product_id:int,recent_days:int=Query(30,ge=1,le=366),recent_limit:int=Query(500,ge=10,le=500)):
    with _product_search_ro() as c:
        p=c.execute("select * from product_search where product_id=?",(product_id,)).fetchone()
        if not p: raise HTTPException(404,"Р В Р’В Р РЋРЎвЂєР В Р’В Р РЋРІР‚СћР В Р’В Р В РІР‚В Р В Р’В Р вЂ™Р’В°Р В Р Р‹Р В РІР‚С™ Р В Р Р‹Р Р†Р вЂљРЎв„ўР В Р’В Р РЋРІР‚СћР В Р’В Р РЋРІР‚вЂќР В Р’В Р РЋРІР‚ВР В Р’В Р вЂ™Р’В»Р В Р’В Р РЋР’ВР В Р’В Р вЂ™Р’В°Р В Р’В Р СћРІР‚ВР В Р’В Р РЋРІР‚В")
        p=dict(p)
        markets=[dict(r) for r in c.execute("select * from product_market_stats where product_id=? order by deals_count desc,market",(product_id,)).fetchall()]
        month_rows=[dict(r) for r in c.execute("select * from product_monthly_stats where product_id=? order by ym,market",(product_id,)).fetchall()]
    mm={}
    season=[{"month":m,"deals":0,"amount":0.0} for m in range(1,13)]

    # Historical seasonality:
    # use only complete SHOP years 2023-2025.
    # Do not mix NATIONAL/AUCTION/XT markets or partial 2026 data.
    shop_hist={}
    for r in month_rows:
        z=mm.setdefault(r["ym"],{"ym":r["ym"],"deals":0,"amount":0.0,"qty":0.0})
        z["deals"]+=int(r.get("deals_count") or 0)
        z["amount"]+=float(r.get("total_amount") or 0)
        z["qty"]+=float(r.get("total_quantity") or 0)

        try:
            ym=str(r["ym"])
            year=int(ym[:4])
            mo=int(ym[5:7])
            market=str(r.get("market") or "").lower()
            if market=="shop" and 2023 <= year <= 2025:
                y=shop_hist.setdefault(year, {})
                y[mo]={
                    "deals":int(r.get("deals_count") or 0),
                    "amount":float(r.get("total_amount") or 0)
                }
        except Exception:
            pass

    # Average the monthly SHOP share within each complete year.
    # This prevents years with different total deal volume from dominating.
    year_totals={}
    for year,months in shop_hist.items():
        year_totals[year]=sum(int(v.get("deals") or 0) for v in months.values())

    for mo in range(1,13):
        shares=[]
        amounts=[]
        for year in (2023,2024,2025):
            total=year_totals.get(year,0)
            row=shop_hist.get(year,{}).get(mo,{})
            deals=int(row.get("deals") or 0)
            amount=float(row.get("amount") or 0)
            if total>0:
                shares.append(deals/total*100.0)
                amounts.append(amount)

        season[mo-1]["deals"]=round(sum(shares)/len(shares),2) if shares else 0
        season[mo-1]["share_pct"]=round(sum(shares)/len(shares),2) if shares else None
        season[mo-1]["amount"]=round(sum(amounts)/len(amounts),2) if amounts else 0.0
        season[mo-1]["source"]="shop_2023_2025_complete_years"
    monthly=[]
    for ym in sorted(mm):
        z=mm[ym];z["unit_price"]=z["amount"]/z["qty"] if z["qty"]>0 else None;monthly.append(z)
    platforms=[]
    for r in markets:
        m=r.get("market")
        platforms.append({"key":m,"title":FAST_MARKET_LABELS.get(m,m or "Р В Р вЂ Р В РІР‚С™Р Р†Р вЂљРЎСљ"),"icon":FAST_MARKET_ICONS.get(m,"Р В Р вЂ Р В РІР‚С™Р РЋРЎвЂє"),"logo":FAST_MARKET_LOGOS.get(m),"source":FAST_MARKET_SOURCES.get(m,m or "Р В Р вЂ Р В РІР‚С™Р Р†Р вЂљРЎСљ"),
        "count":int(r.get("deals_count") or 0),"amount":float(r.get("total_amount") or 0),"quantity":float(r.get("total_quantity") or 0),
        "unit_price":{"min":r.get("min_unit_price"),"max":r.get("max_unit_price"),"avg":r.get("avg_unit_price"),
        "median":r.get("median_unit_price"),"count":int(r.get("deals_count") or 0)}})
    return {"query":p["display_name"],"product_id":product_id,"match_mode":"fast_cache_exact",
    "summary":{"deals":int(p.get("deals_count") or 0),"amount":float(p.get("total_amount") or 0),"quantity":float(p.get("total_quantity") or 0),
    "platforms":int(p.get("markets_count") or 0),"unit_min":p.get("min_unit_price"),"unit_max":p.get("max_unit_price"),
    "unit_avg":p.get("avg_unit_price"),"unit_median":p.get("median_unit_price"),"first_date":p.get("first_date"),"last_date":p.get("last_date")},
    "platforms":platforms,"seasonality":season,"monthly":monthly[-24:],
    "decision_v2":_product_decision_v2(p["display_name"],p.get("last_date"),season),
    "recent_summary":_fast_recent_stats(p["display_name"],p.get("last_date"),recent_days),
    "recent":_fast_recent(p["display_name"],p.get("last_date"),recent_days,recent_limit),
    "codes":_product_codes_for_name(p["display_name"],12)}

def _auction_lot_completed_details(deal_id=None, lot_id=None, lot_no=None):
    """Return ALL products and properties belonging to one completed UZEX auction lot."""
    v=platform("auction")
    if not v["db"].exists():
        return {"products":[],"properties":[]}

    ids=[]
    for x in (deal_id,lot_id,lot_no):
        if x not in (None,""):
            s=str(x).strip()
            if s and s not in ids: ids.append(s)
    if not ids:
        return {"products":[],"properties":[]}

    with ro(v["db"]) as c:
        ts=set(tables(c))
        products=[]
        props=[]

        if "auction_products" in ts:
            cs=cols(c,"auction_products")
            idcs=[x for x in ("deal_id","lot_id","lot_number","lot_display_no","id") if x in cs]
            if idcs:
                where=" or ".join(f"cast({ident(col)} as text)=?" for col in idcs for _ in ids)
                params=[]
                for col in idcs: params.extend(ids)
                order=[]
                for col in ("product_no","item_no","id"):
                    if col in cs: order.append(ident(col))
                sql=f"select * from auction_products where ({where})"
                if order: sql+=" order by "+",".join(order)
                products=[dict(x) for x in c.execute(sql,params).fetchall()]

        # Dynamic property-table discovery.
        prop_tables=[]
        for tn in ("auction_product_properties","product_properties","completed_product_properties","auction_properties"):
            if tn in ts and tn not in prop_tables: prop_tables.append(tn)
        for tn in ts:
            low=tn.lower()
            if "propert" in low and tn not in prop_tables:
                prop_tables.append(tn)

        for tn in prop_tables:
            cs=cols(c,tn)
            idcs=[x for x in ("deal_id","lot_id","lot_number","lot_display_no") if x in cs]
            if not idcs:
                continue
            where=" or ".join(f"cast({ident(col)} as text)=?" for col in idcs for _ in ids)
            params=[]
            for col in idcs: params.extend(ids)
            order=[]
            for col in ("product_no","item_no","id","prop_id"):
                if col in cs: order.append(ident(col))
            sql=f"select * from {ident(tn)} where ({where})"
            if order: sql+=" order by "+",".join(order)
            try:
                props.extend(dict(x) for x in c.execute(sql,params).fetchall())
            except Exception:
                pass

        return {"products":products,"properties":props}


@app.get("/api/product-recent/auction-lot-detail")
def product_recent_auction_lot_detail(
    deal_id:Optional[str]=None,
    lot_id:Optional[str]=None,
    lot_no:Optional[str]=None
):
    data=_auction_lot_completed_details(deal_id,lot_id,lot_no)

    # BIData MASTER product-code policy for completed Auction cards:
    # exact normalized product name + exactly one unique code in product_codes.sqlite3.
    # Existing source code remains only as fallback when MASTER has no unique match.
    products=data.get("products") or []

    if products and PRODUCT_CODES_DB.exists():
        try:
            master={}

            with ro(PRODUCT_CODES_DB) as pc:
                if "product_codes" in set(tables(pc)):
                    rows=pc.execute("""
                        select product_name, product_code
                        from product_codes
                        where product_name is not null
                          and trim(product_name)<>''
                          and product_code is not null
                          and trim(product_code)<>''
                    """).fetchall()

                    for row in rows:
                        name=str(row["product_name"] or "").strip().casefold()
                        code=str(row["product_code"] or "").strip()

                        if not name or not code:
                            continue

                        master.setdefault(name,set()).add(code)

            for item in products:
                # Highest priority: manually verified canonical code.
                manual_code=_manual_product_code_override(
                    item.get("product_name")
                )

                if manual_code:
                    item["product_code"]=manual_code
                    item["product_code_source"]="manual_verified"
                    continue

                # Second priority: unique exact-name MASTER code.
                name=str(item.get("product_name") or "").strip().casefold()
                codes=master.get(name)

                if codes and len(codes)==1:
                    item["product_code"]=next(iter(codes))
                    item["product_code_source"]="product_codes.sqlite3"

        except Exception:
            # Never break the lot card if the MASTER catalog is temporarily unavailable.
            pass

    return {
        "platform_key":"auction",
        "title":PLATFORMS["auction"]["title"],
        "count":len(products),
        "products":products,
        "properties":data["properties"],
    }


# ================== /v10.5 FAST PRODUCT SEARCH / BI ==================

# ==================== v10.2 PRODUCT INTELLIGENCE ====================
def _safe_float(v):
    try: return float(v or 0)
    except: return 0.0

def _pick(cs, names):
    low={x.lower():x for x in cs}
    for n in names:
        if n.lower() in low:return low[n.lower()]
    return None

def _product_source(c, kind):
    ts=tables(c)
    candidates=[]
    if kind=="auction" and "auction_products" in ts: candidates.append("auction_products")
    if kind=="xt" and "items" in ts: candidates.append("items")
    candidates += [x for x in ("completed_deals","tenders","lots","items","deals") if x in ts and x not in candidates]
    for t in candidates:
        cs=cols(c,t)
        pc=_pick(cs,["product_name","service_type","name","title","category_name","subject"])
        if pc:return t,cs,pc
    return None,[],None

def _product_profile_one(key, q, recent_days=30, recent_limit=100):
    v=platform(key)
    if not v["db"].exists(): return None
    with ro(v["db"]) as c:
        t,cs,pc=_product_source(c,v["kind"])
        if not t or not pc:return None
        dc=_pick(cs,["deal_date","confirmed_date","confirmed_at","created_at","publicated_at","published_at","placement_date","start_date"])
        lotc=_pick(cs,["lot_display_no","lot_number","display_no","number","proc_number","lot_id","proc_id","tender_id","deal_id","id"])
        qtyc=_pick(cs,["amount","quantity","qty","count","total_count"])
        totalc=_pick(cs,["deal_cost","total_amount","winner_offer","winner_amount","contract_price","total_price","price","sum","start_price"])
        unitc=_pick(cs,["unit_name","unit","measure_name"])
        custc=_pick(cs,["customer_name","buyer_name","organization_name","org_name","customer"])
        suppc=_pick(cs,["provider_name","supplier_name","winner_name","provider"])
        regc=_pick(cs,["customer_region_name","region_name_uz","region_name","customer_region","region"])
        statc=_pick(cs,["deal_status_name","status_name","status","state"])
        idc=_pick(cs,["deal_id","id","tender_id","lot_id","proc_id"])
        codec=_pick(cs,["product_code","code","mxik_code","ikpu_code","catalog_code","classifier_code","product_classifier_code"])
        # Exact selected product, case-insensitive and whitespace-tolerant.
        where=f"lower(trim({ident(pc)}))=lower(trim(?))"
        p=[q]
        count=c.execute(f"select count(*) from {ident(t)} where {where}",p).fetchone()[0]
        if not count:return {"key":key,"title":v["title"],"icon":v["icon"],"count":0}
        total=0.0; qty=0.0
        if totalc:
            total=_safe_float(c.execute(f"select coalesce(sum(cast({ident(totalc)} as real)),0) from {ident(t)} where {where}",p).fetchone()[0])
        if qtyc:
            qty=_safe_float(c.execute(f"select coalesce(sum(cast({ident(qtyc)} as real)),0) from {ident(t)} where {where}",p).fetchone()[0])

        # Unit-price stats are calculated row-by-row as total / quantity.
        unit_stats={"min":None,"max":None,"avg":None,"median":None,"count":0}
        prices=[]
        if totalc and qtyc and totalc!=qtyc:
            rows=c.execute(f"""select cast({ident(totalc)} as real) tot,cast({ident(qtyc)} as real) qty
                               from {ident(t)} where {where}
                               and cast({ident(qtyc)} as real)>0 and cast({ident(totalc)} as real)>=0""",p).fetchall()
            prices=[_safe_float(r["tot"])/_safe_float(r["qty"]) for r in rows if _safe_float(r["qty"])>0]
            prices=[x for x in prices if x>=0]
            if prices:
                ps=sorted(prices); n=len(ps)
                med=ps[n//2] if n%2 else (ps[n//2-1]+ps[n//2])/2
                unit_stats={"min":min(ps),"max":max(ps),"avg":sum(ps)/n,"median":med,"count":n}

        # Monthly seasonality and price dynamics.
        monthly=[]
        if dc:
            sel=[f"substr({ident(dc)},1,7) ym","count(*) deals"]
            sel.append(f"coalesce(sum(cast({ident(totalc)} as real)),0) amount" if totalc else "0 amount")
            sel.append(f"coalesce(sum(cast({ident(qtyc)} as real)),0) qty" if qtyc else "0 qty")
            monthly=[dict(r) for r in c.execute(f"""select {','.join(sel)} from {ident(t)}
                         where {where} and coalesce({ident(dc)},'')<>''
                         group by substr({ident(dc)},1,7) order by ym""",p).fetchall()]
            for r in monthly:
                r["unit_price"]=(_safe_float(r["amount"])/_safe_float(r["qty"])) if _safe_float(r["qty"])>0 else None

        season=[{"month":m,"deals":0,"amount":0.0} for m in range(1,13)]
        for r in monthly:
            try:m=int(str(r["ym"])[5:7])
            except:continue
            if 1<=m<=12:
                season[m-1]["deals"]+=int(r["deals"] or 0);season[m-1]["amount"]+=_safe_float(r["amount"])

        recent=[]
        if dc:
            maxd=c.execute(f"select max(substr({ident(dc)},1,10)) from {ident(t)} where {where}",p).fetchone()[0]
            if maxd:
                fields=[
                    f"{ident(dc)} date",
                    f"{ident(lotc)} lot_no" if lotc else "'' lot_no",
                    f"{ident(pc)} product_name",
                    f"{ident(qtyc)} quantity" if qtyc else "null quantity",
                    f"{ident(unitc)} unit_name" if unitc else "'' unit_name",
                    f"{ident(totalc)} total_value" if totalc else "null total_value",
                    f"{ident(custc)} customer" if custc else "'' customer",
                    f"{ident(suppc)} supplier" if suppc else "'' supplier",
                    f"{ident(regc)} region" if regc else "'' region",
                    f"{ident(statc)} status" if statc else "'' status",
                    f"{ident(idc)} source_id" if idc else "'' source_id",
                    f"{ident(codec)} product_code" if codec else "null product_code",
                ]
                rp=[q,maxd,f"-{int(recent_days)} day",int(recent_limit)]
                recent=[dict(r) for r in c.execute(f"""select {','.join(fields)} from {ident(t)}
                    where {where} and substr({ident(dc)},1,10)>=date(? ,?)
                    order by {ident(dc)} desc limit ?""",rp).fetchall()]
                for r in recent:
                    qq=_safe_float(r.get("quantity")); tv=_safe_float(r.get("total_value"))
                    r["unit_price"]=tv/qq if qq>0 else None
                    r["platform"]=v["title"];r["platform_key"]=key
        return {"key":key,"title":v["title"],"icon":v["icon"],"count":count,"amount":total,"quantity":qty,
                "unit_price":unit_stats,"monthly":monthly,"seasonality":season,"recent":recent}

@app.get("/api/product-intelligence")
def product_intelligence(q:str=Query(...,min_length=2),recent_days:int=Query(30,ge=1,le=366),recent_limit:int=Query(120,ge=10,le=500)):
    result=[]; errors=[]
    for key in ("shop","national","auction","xt_auction"):
        try:
            x=_product_profile_one(key,q,recent_days,recent_limit)
            if x: result.append(x)
        except Exception as e:
            errors.append({"key":key,"error":str(e)})
    active=[x for x in result if x.get("count",0)>0]
    prices=[]
    for x in active:
        u=x.get("unit_price") or {}
        if u.get("count") and u.get("avg") is not None:
            prices.append((u["avg"],u["count"]))
    total_price_n=sum(n for _,n in prices)
    weighted_avg=sum(p*n for p,n in prices)/total_price_n if total_price_n else None
    mins=[x["unit_price"]["min"] for x in active if x.get("unit_price",{}).get("min") is not None]
    maxs=[x["unit_price"]["max"] for x in active if x.get("unit_price",{}).get("max") is not None]
    recent=sorted([r for x in active for r in x.get("recent",[])],key=lambda r:str(r.get("date") or ""),reverse=True)[:recent_limit]
    recent=_repair_xt_auction_recent(_repair_recent_uzex_product_codes(recent))
    season=[{"month":m,"deals":0,"amount":0.0} for m in range(1,13)]
    for x in active:
        for r in x.get("seasonality",[]):
            season[r["month"]-1]["deals"]+=int(r["deals"] or 0);season[r["month"]-1]["amount"]+=_safe_float(r["amount"])
    monthly_map={}
    for x in active:
        for r in x.get("monthly",[]):
            z=monthly_map.setdefault(r["ym"],{"ym":r["ym"],"deals":0,"amount":0.0,"qty":0.0})
            z["deals"]+=int(r["deals"] or 0);z["amount"]+=_safe_float(r["amount"]);z["qty"]+=_safe_float(r["qty"])
    monthly=[]
    for ym in sorted(monthly_map):
        z=monthly_map[ym];z["unit_price"]=z["amount"]/z["qty"] if z["qty"]>0 else None;monthly.append(z)
    return {
      "query":q,"match_mode":"exact_product_name",
      "summary":{"deals":sum(x["count"] for x in active),"amount":sum(_safe_float(x.get("amount")) for x in active),
                 "quantity":sum(_safe_float(x.get("quantity")) for x in active),"platforms":len(active),
                 "unit_min":min(mins) if mins else None,"unit_max":max(maxs) if maxs else None,
                 "unit_avg":weighted_avg},
      "platforms":active,"seasonality":season,"monthly":monthly[-24:],"recent":recent,"errors":errors
    }
# ================== /v10.2 PRODUCT INTELLIGENCE ==================

def _parse_pending_dt(v):
    """Best-effort parser for pending lot deadlines (local/Uzb server time semantics)."""
    if v is None: return None
    s=str(v).strip()
    if not s: return None
    s=s.replace("Z","+00:00")
    try:
        dt=datetime.fromisoformat(s)
        # Compare timezone-aware values in local naive form only when offset is present.
        if dt.tzinfo is not None:
            dt=dt.astimezone().replace(tzinfo=None)
        return dt
    except Exception:
        pass
    for fmt in ("%Y-%m-%d %H:%M:%S","%Y-%m-%d %H:%M","%Y-%m-%d","%d.%m.%Y %H:%M:%S","%d.%m.%Y %H:%M","%d.%m.%Y"):
        try:return datetime.strptime(s[:19],fmt)
        except Exception:pass
    return None


def _truthy_db(v):
    if v is None:return None
    if isinstance(v,(int,float)):return bool(v)
    s=str(v).strip().casefold()
    if s in {"1","true","yes","y","ha","Р В Р’В Р СћРІР‚ВР В Р’В Р вЂ™Р’В°","active","open"}:return True
    if s in {"0","false","no","n","yo'q","Р В Р’В Р Р†РІР‚С›РІР‚вЂњР В Р Р‹Р РЋРІР‚С”Р В РЎС›Р Р†Р вЂљРЎвЂќ","Р В Р’В Р В РІР‚В¦Р В Р’В Р вЂ™Р’ВµР В Р Р‹Р Р†Р вЂљРЎв„ў","inactive","closed"}:return False
    return None


def _pending_parent_is_active(base:dict, parent_cols:list[str], m:dict) -> bool:
    """A row is shown as active only when its master parent itself is still active/current."""
    now=datetime.now()
    low={c.lower():c for c in parent_cols}

    # Strongest signal: explicit active/is_active/open flag in the master table.
    active_col=None
    for name in ("active","is_active","is_open","open","enabled"):
        if name in low:
            active_col=low[name]; break
    if active_col:
        flag=_truthy_db(base.get(active_col))
        if flag is False:return False

    # Expired deadline is never an active lot, regardless of stale textual status.
    deadline_col=m.get("deadline")
    deadline=base.get(deadline_col) if deadline_col else None
    dt=_parse_pending_dt(deadline)
    if dt is not None and dt < now:
        return False

    # Reject explicit terminal/cancelled statuses where a platform keeps history in the same table.
    status_col=m.get("status")
    st=str(base.get(status_col) if status_col else "").strip().casefold()
    bad_tokens=(
        "cancel", "cancelled", "canceled", "closed", "finished",
        "completed", "expired", "rejected",
        "bekor", "muddat",
        "\u043e\u0442\u043c\u0435\u043d",
        "\u0437\u0430\u043a\u0440\u044b\u0442",
        "\u0437\u0430\u0432\u0435\u0440\u0448",
        "\u0438\u0441\u0442\u0435\u043a",
        "\u0440\u0430\u0441\u0442\u043e\u0440\u0433",
        "\u0430\u043d\u043d\u0443\u043b",
    )
    if st and any(x in st for x in bad_tokens):
        return False

    # If there is neither an explicit active flag nor a deadline, trust the dedicated pending master table.
    return True


def _canonical_active_product_code(product_name, source_code=None):
    """
    BIData active-market product-code policy:
      1. manually verified override
      2. unique exact-name MASTER code
      3. source code normalized to exactly 14 digits
      4. None
    """
    product_name = str(product_name or "").strip()

    # 1) Highest priority: manually verified mapping.
    manual_code = _manual_product_code_override(product_name)
    if manual_code:
        digits = re.sub(r"\D", "", str(manual_code))
        if len(digits) == 14:
            return digits, "manual_verified"

    # 2) Unique exact-name MASTER mapping.
    if product_name:
        master_codes = _catalog_codes_for_exact_product_name(
            product_name, 12
        )

        unique = []
        seen = set()

        for code in master_codes or []:
            digits = re.sub(r"\D", "", str(code or ""))
            if len(digits) == 14 and digits not in seen:
                seen.add(digits)
                unique.append(digits)

        if len(unique) == 1:
            return unique[0], "product_codes.sqlite3"

    # 3) Preserve the source meaning, but normalize old punctuation.
    digits = re.sub(r"\D", "", str(source_code or ""))

    if len(digits) == 14:
        return digits, "source_canonicalized"

    return None, None


def _active_product_matches(q:str, limit_each:int=80, only_keys=None):
    only_keys=set(only_keys) if only_keys else None
    allrows=[]; errors=[]
    q_norm=str(q or "").strip().casefold()

    # SHOP active offers use their own schema; keep auction/xt_national logic untouched.
    shopv=PENDING.get("shop")
    if (only_keys is None or "shop" in only_keys) and shopv and shopv["db"].exists():
        try:
            with ro(shopv["db"]) as c:
                if "shop_offers" in set(tables(c)):
                    shop_cols=set(cols(c,"shop_offers"))
                    where_parts=[]
                    params=[]

                    if "product_name" in shop_cols:
                        ftsdb = DATA / "bidata_active_search_fts.sqlite3"
                        c.execute("ATTACH DATABASE ? AS activefts", (str(ftsdb),))

                        active_parts=[]
                        if "is_current" in shop_cols:
                            active_parts.append("coalesce(s.is_current,0)=1")
                        if "date_end" in shop_cols:
                            active_parts.append("(s.date_end is null or trim(s.date_end)='' or s.date_end>=datetime('now','localtime'))")

                        sql = """
                            select s.*
                            from activefts.shop_active_fts f
                            join shop_offers s on s.offer_id=f.offer_id
                            where f.product_name_norm like ?
                        """
                        if active_parts:
                            sql += " and " + " and ".join(active_parts)
                        sql += " order by coalesce(s.date_end,s.date_ini,'') asc limit ?"

                        params = [f"%{q_norm}%", limit_each]

                        try:
                            matched_rows = c.execute(sql,tuple(params)).fetchall()
                        finally:
                            c.execute("DETACH DATABASE activefts")

                        for rr in matched_rows:
                            d=dict(rr)
                            allrows.append({
                                "platform_key":"shop",
                                "platform":shopv["title"],
                                "id":d.get("offer_id"),
                                "lot_no":d.get("display_no") or d.get("offer_id"),
                                "date":d.get("date_ini"),
                                "deadline":d.get("date_end"),
                                "product":d.get("product_name"),
                                "product_code":d.get("product_code"),
                                "category":d.get("category_name"),
                                "amount":d.get("total_value"),
                                "quantity":d.get("amount"),
                                "unit_name":d.get("unit_name"),
                                "unit_price":d.get("price"),
                                "customer":d.get("customer_name"),
                                "customer_inn":d.get("customer_inn"),
                                "region":d.get("customer_region_name"),
                                "status":d.get("status_name"),
                                "source_url":d.get("source_url"),
                                "mark_name":d.get("mark_name"),
                                "manufacturer_name":d.get("manufacturer_name"),
                                "description":d.get("product_description") or d.get("condition_text"),
                            })
        except Exception as e:
            errors.append({"key":"shop","error":str(e)})

    # NATIONAL SHOP active offers
    nationalv=PENDING.get("national")
    if (only_keys is None or "national" in only_keys) and nationalv and nationalv["db"].exists():
        try:
            with ro(nationalv["db"]) as c:
                if "national_offers" in set(tables(c)):
                    national_cols=set(cols(c,"national_offers"))
                    where_parts=[]
                    params=[]

                    if "product_name" in national_cols:
                        ftsdb = DATA / "bidata_active_search_fts.sqlite3"
                        c.execute("ATTACH DATABASE ? AS activefts", (str(ftsdb),))

                        active_parts=[]
                        if "is_current" in national_cols:
                            active_parts.append("coalesce(s.is_current,0)=1")
                        if "date_end" in national_cols:
                            active_parts.append("(s.date_end is null or trim(s.date_end)='' or s.date_end>=datetime('now','localtime'))")

                        sql = """
                            select s.*
                            from activefts.national_active_fts f
                            join national_offers s on s.offer_id=f.offer_id
                            where f.product_name_norm like ?
                        """
                        if active_parts:
                            sql += " and " + " and ".join(active_parts)
                        sql += " order by coalesce(s.date_end,s.date_ini,'') asc limit ?"

                        params = [f"%{q_norm}%", limit_each]

                        try:
                            matched_rows = c.execute(sql,tuple(params)).fetchall()
                        finally:
                            c.execute("DETACH DATABASE activefts")

                        for rr in matched_rows:
                            d=dict(rr)
                            allrows.append({
                                "platform_key":"national",
                                "platform":nationalv["title"],
                                "id":d.get("offer_id"),
                                "lot_no":d.get("display_no") or d.get("offer_id"),
                                "date":d.get("date_ini"),
                                "deadline":d.get("date_end"),
                                "product":d.get("product_name"),
                                "product_code":d.get("product_code"),
                                "category":d.get("category_name"),
                                "amount":d.get("total_value"),
                                "quantity":d.get("amount"),
                                "unit_name":d.get("unit_name"),
                                "unit_price":d.get("price"),
                                "customer":d.get("customer_name"),
                                "customer_inn":d.get("customer_inn"),
                                "region":d.get("customer_region_name"),
                                "status":d.get("status_name"),
                                "source_url":d.get("source_url"),
                                "mark_name":d.get("mark_name"),
                                "manufacturer_name":d.get("manufacturer_name"),
                                "description":d.get("product_description") or d.get("condition_text"),
                            })
        except Exception as e:
            errors.append({"key":"national","error":str(e)})

    for key in ("auction","xt_auction","xt_national"):
        if only_keys is not None and key not in only_keys:
            continue
        v=PENDING[key]
        if not v["db"].exists(): continue
        try:
            with ro(v["db"]) as c:
                t,m=pending_table_and_map(c,v["kind"])
                parent_cols=cols(c,t)
                parent_colset=set(parent_cols)

                # Search the master pending table and its product/item children.
                # IMPORTANT: a child hit is accepted ONLY if its parent still exists in the pending master.
                candidates=[]
                pc=m.get("product")
                if pc:
                    candidates.append((t,pc,None,"main"))
                ts=set(tables(c))
                for child in ("pending_products","products","items","lots","tender_items"):
                    if child in ts and child!=t:
                        cc=cols(c,child)
                        cp=_pick(cc,["product_name","name","title","category_name","subject"])
                        joinid=_pick(cc,["lot_id","proc_id","tender_id","tenderid","parent_id","id"])
                        if cp and joinid:candidates.append((child,cp,joinid,"child"))

                seen=set()
                for ct,cp,cid,mode in candidates:

                    if key=="xt_national" and mode=="main" and ct=="national_lots" and cp=="product_name":
                        ftsdb = DATA / "bidata_active_search_fts.sqlite3"
                        c.execute("ATTACH DATABASE ? AS activefts", (str(ftsdb),))
                        try:
                            rows = c.execute("""
                                select s.*
                                from activefts.xt_national_active_fts f
                                join national_lots s
                                  on s.proc_id=f.proc_id
                                where f.product_name_norm like ?
                                limit ?
                            """,(f"%{q_norm}%",limit_each)).fetchall()
                        finally:
                            c.execute("DETACH DATABASE activefts")
                    elif key in ("auction","xt_auction"):
                        # SQLite LIKE is not Unicode case-insensitive for Cyrillic.
                        # These pending product tables are small, so use Python casefold.
                        scan_rows=c.execute(
                            f"""select * from {ident(ct)}"""
                        ).fetchall()
                        rows=[
                            rr for rr in scan_rows
                            if q_norm in str(dict(rr).get(cp) or "").strip().casefold()
                        ][:limit_each]
                    else:
                        rows=c.execute(f"""select * from {ident(ct)}
                            where {ident(cp)} like ?
                            limit ?""",(f"%{q}%",limit_each)).fetchall()
                    for rr in rows:
                        d=dict(rr)
                        if mode=="main":
                            base=d
                            rid=str(base.get(m.get("id")) if m.get("id") else base.get("id") or "")
                        else:
                            rid=str(d.get(cid) or "")
                            if not rid:
                                continue

                            # Match child FK to the same-named parent key first (lot_id->lot_id etc.).
                            # Falling back to mapped id is safe only when that parent column exists.
                            parent_key=cid if cid in parent_colset else m.get("id")
                            if not parent_key or parent_key not in parent_colset:
                                continue
                            pr=c.execute(
                                f"select * from {ident(t)} where {ident(parent_key)}=? limit 1",
                                (rid,)
                            ).fetchone()
                            # Critical: orphan/history child rows must NEVER appear as active lots.
                            if not pr:
                                continue
                            base=dict(pr)

                        # Parent is authoritative for activity/deadline/status.
                        if not _pending_parent_is_active(base,parent_cols,m):
                            continue

                        def val(alias):
                            col=m.get(alias); return base.get(col) if col else None

                        parent_id=val("id") or rid
                        display=val("display") or parent_id

                        # One matching CHILD PRODUCT = one active result/card.
                        # Do not collapse two different products/items that belong
                        # to the same lot number.
                        if mode=="child":
                            child_row_id=(
                                d.get("id") or
                                d.get("item_id") or
                                d.get("product_id") or
                                d.get("detail_id")
                            )

                            # Product POSITION identity:
                            # same lot may contain the same product name more than once
                            # with different unit / quantity / price.
                            child_token="|".join(str(x or "") for x in (
                                child_row_id,
                                d.get(cp),
                                d.get("product_code") or d.get("code"),
                                d.get("quantity") or d.get("qty"),
                                d.get("unit_name") or d.get("unit") or d.get("measure_name"),
                                d.get("start_unit_price") or d.get("start_price") or
                                d.get("initial_unit_price") or d.get("initial_price") or d.get("price"),
                                d.get("start_cost") or d.get("initial_cost") or
                                d.get("start_amount") or d.get("initial_amount")
                            ))

                            ded=(key,str(parent_id),str(display),"child",child_token)
                        else:
                            child_row_id=None
                            ded=(key,str(parent_id),str(display),"main")

                        if ded in seen:
                            continue
                        seen.add(ded)

                        def dval(*names):
                            for name in names:
                                if name in d and d.get(name) is not None:
                                    return d.get(name)
                            return None

                        matched_qty=dval("quantity","qty")
                        matched_unit=dval("unit_name","unit","measure_name")
                        matched_price=dval(
                            "start_unit_price","start_price",
                            "initial_unit_price","initial_price","price"
                        )
                        
                        if key=="xt_national" and mode=="main":
                            matched_qty=base.get("amount")
                            matched_unit=base.get("unit")
                            matched_price=base.get("price")
                        matched_amount=dval(
                            "start_cost","initial_cost",
                            "start_amount","initial_amount",
                            "total_amount","total_price"
                        )

                        if matched_amount is None and matched_qty is not None and matched_price is not None:
                            try:
                                matched_amount=float(matched_qty)*float(matched_price)
                            except Exception:
                                matched_amount=None

                        result_id = display if key=="xt_auction" else parent_id
                        proc_id = base.get("proc_id") if key=="xt_auction" else None

                        allrows.append({
                          "platform_key":key,"platform":v["title"],"id":result_id,
                          "proc_id":proc_id,
                          "lot_no":display,"date":val("date"),"deadline":val("deadline"),
                          "product":d.get(cp) or val("product"),
                          "amount":matched_amount if matched_amount is not None else val("amount"),
                          "quantity":matched_qty if matched_qty is not None else None,
                          "unit_name":matched_unit if matched_unit is not None else None,
                          "unit_price":matched_price if matched_price is not None else None,
                          "matched_item_id":child_row_id,
                          "customer":val("customer"),
                          "customer_inn":(
                              val("customer_inn")
                              or (_organization_stir_lookup(val("customer")) if key=="auction" else None)
                          ),
                          "region":val("region"),"status":val("status")
                        })
        except Exception as e:
            errors.append({"key":key,"error":str(e)})
    # BIData canonical product-code policy for all 5 active market platforms.
    for r in allrows:
        product_name = r.get("product") or r.get("product_name")
        source_code = r.get("product_code")

        code, code_source = _canonical_active_product_code(
            product_name,
            source_code
        )

        if code:
            r["product_code"] = code
            r["product_code_source"] = code_source

    allrows.sort(key=lambda x:str(x.get("deadline") or x.get("date") or ""))
    return allrows,errors


def _active_product_aggregate(q):
    """
    Full active-product summary without using the UI row limit.
    Large FTS-backed platforms use SQL COUNT/SUM.
    Small auction platforms reuse the existing activity/dedupe logic.
    """
    q_norm=str(q or "").strip().casefold()
    result={}
    errors=[]

    def put(key,count,amount):
        meta=PENDING.get(key,{})
        result[key]={
            "key":key,
            "title":meta.get("title") or key,
            "source":meta.get("source") or "",
            "trade_type":meta.get("trade_type") or "",
            "logo":meta.get("logo"),
            "count":int(count or 0),
            "amount":_safe_float(amount),
        }

    ftsdb=DATA / "bidata_active_search_fts.sqlite3"

    # SHOP
    shopv=PENDING.get("shop")
    if shopv and shopv["db"].exists() and ftsdb.exists():
        try:
            with ro(shopv["db"]) as c:
                c.execute("ATTACH DATABASE ? AS activefts",(str(ftsdb),))
                try:
                    r=c.execute("""
                        select
                            count(*) n,
                            coalesce(sum(coalesce(s.total_value,0)),0) amount
                        from activefts.shop_active_fts f
                        join shop_offers s on s.offer_id=f.offer_id
                        where f.product_name_norm like ?
                          and coalesce(s.is_current,0)=1
                          and (
                                s.date_end is null
                                or trim(s.date_end)=''
                                or s.date_end >= datetime('now','localtime')
                              )
                    """,(f"%{q_norm}%",)).fetchone()
                finally:
                    c.execute("DETACH DATABASE activefts")
                put("shop",r["n"],r["amount"])
        except Exception as e:
            errors.append("shop aggregate: "+str(e))

    # NATIONAL
    natv=PENDING.get("national")
    if natv and natv["db"].exists() and ftsdb.exists():
        try:
            with ro(natv["db"]) as c:
                c.execute("ATTACH DATABASE ? AS activefts",(str(ftsdb),))
                try:
                    r=c.execute("""
                        select
                            count(*) n,
                            coalesce(sum(coalesce(s.total_value,0)),0) amount
                        from activefts.national_active_fts f
                        join national_offers s on s.offer_id=f.offer_id
                        where f.product_name_norm like ?
                          and coalesce(s.is_current,0)=1
                          and (
                                s.date_end is null
                                or trim(s.date_end)=''
                                or s.date_end >= datetime('now','localtime')
                              )
                    """,(f"%{q_norm}%",)).fetchone()
                finally:
                    c.execute("DETACH DATABASE activefts")
                put("national",r["n"],r["amount"])
        except Exception as e:
            errors.append("national aggregate: "+str(e))

    # XT NATIONAL
    xtnv=PENDING.get("xt_national")
    if xtnv and xtnv["db"].exists() and ftsdb.exists():
        try:
            with ro(xtnv["db"]) as c:
                c.execute("ATTACH DATABASE ? AS activefts",(str(ftsdb),))
                try:
                    r=c.execute("""
                        select
                            count(*) n,
                            coalesce(
                                sum(coalesce(s.price,0)*coalesce(s.amount,0)),
                                0
                            ) amount
                        from activefts.xt_national_active_fts f
                        join national_lots s on s.proc_id=f.proc_id
                        where f.product_name_norm like ?
                          and coalesce(s.is_current,0)=1
                          and (
                                s.close_at is null
                                or trim(s.close_at)=''
                                or datetime(
                                    replace(replace(s.close_at,'T',' '),'Z','')
                                ) >= datetime('now','localtime')
                              )
                    """,(f"%{q_norm}%",)).fetchone()
                finally:
                    c.execute("DETACH DATABASE activefts")
                put("xt_national",r["n"],r["amount"])
        except Exception as e:
            errors.append("xt_national aggregate: "+str(e))

    # AUCTION + XT AUCTION
    # Their pending product tables are small. Reuse the proven parent,
    # activity and dedupe rules, but exclude the three large FTS platforms.
    demand_detail_rows=[]
    try:
        full_rows,match_errors=_active_product_matches(
            q,100000,only_keys=("auction","xt_auction")
        )
        demand_detail_rows=full_rows
        errors.extend(match_errors or [])
        for key in ("auction","xt_auction"):
            rr=[x for x in full_rows if x.get("platform_key")==key]
            put(
                key,
                len(rr),
                sum(_safe_float(x.get("amount")) for x in rr)
            )
    except Exception as e:
        errors.append("auction aggregate: "+str(e))

    # Keep zero-count platforms in the breakdown as well; endpoint can
    # decide whether to expose or hide them.
    for key in ("shop","national","auction","xt_auction","xt_national"):
        if key not in result:
            put(key,0,0.0)

    breakdown=sorted(
        result.values(),
        key=lambda x:(-int(x.get("count") or 0),str(x.get("title") or ""))
    )

    supply_keys={"shop","national","xt_national"}
    demand_keys={"auction","xt_auction"}

    supply_rows=[x for x in breakdown if x.get("key") in supply_keys]
    demand_rows=[x for x in breakdown if x.get("key") in demand_keys]

    supply_detail_rows=[]
    try:
        supply_detail_rows,supply_match_errors=_active_product_matches(
            q,100000,only_keys=("shop","national","xt_national")
        )
        errors.extend(supply_match_errors or [])
    except Exception as e:
        errors.append("supply detail: "+str(e))

    supply={
        "count":sum(int(x.get("count") or 0) for x in supply_rows),
        "amount":sum(_safe_float(x.get("amount")) for x in supply_rows),
        "breakdown":supply_rows,
        "rows":supply_detail_rows,
    }
    demand={
        "count":sum(int(x.get("count") or 0) for x in demand_rows),
        "amount":sum(_safe_float(x.get("amount")) for x in demand_rows),
        "breakdown":demand_rows,
        "rows":demand_detail_rows,
    }

    return {
        "count":sum(int(x.get("count") or 0) for x in breakdown),
        "amount":sum(_safe_float(x.get("amount")) for x in breakdown),
        "breakdown":breakdown,
        "supply":supply,
        "demand":demand,
        "errors":errors,
    }


@app.get("/api/product-active-lots")
def product_active_lots(
    q:str=Query(...,min_length=2),
    limit:int=Query(150,ge=10,le=500),
    offset:int=Query(0,ge=0,le=500)
):
    import time as _time
    _t0 = _time.perf_counter()

    fetch_limit=min(500,offset+limit)
    rows,errors=_active_product_matches(q,fetch_limit)
    _t1 = _time.perf_counter()

    aggregate=_active_product_aggregate(q)

    visible=rows[offset:offset+limit]

    return {
        "query":q,
        "count":aggregate["count"],
        "visible_count":len(visible),
        "offset":offset,
        "limit":limit,
        "has_more":(offset+len(visible)) < aggregate["count"],
        "amount":aggregate["amount"],
        "breakdown":aggregate["breakdown"],
        "supply":aggregate["supply"],
        "demand":aggregate["demand"],
        "rows":visible,
        "errors":errors + aggregate.get("errors",[])
    }


@app.get("/api/auction/analytics-check")
def auction_analytics_check():
    v=platform("auction")
    with ro(v["db"]) as c:
        ts=tables(c)
        result={"status":"ok","db":str(v["db"]),"tables":ts,"checks":{}}
        for dim in ("products","categories","customers","providers","regions","statuses"):
            try:
                rows=_auction_dimension_query(c,dim,3,None)
                result["checks"][dim]={"ok":True,"rows":rows}
            except HTTPException as e:
                result["checks"][dim]={"ok":False,"error":str(e.detail)}
        return result

@app.get("/api/{key}/dimension/{dim}")
def dimension_rows(
    key:str,
    dim:str,
    limit:int=Query(100,ge=1,le=500),
    q:Optional[str]=None,
    product_q:Optional[str]=None,
    date_from:Optional[str]=None,
    date_to:Optional[str]=None,
    region:Optional[str]=None,amount_min:Optional[str]=None,amount_max:Optional[str]=None):
    amount_min,amount_max=_market_amount_bounds(key,amount_min,amount_max)
    """
    Detailed dimension browser for Products / Customers / Providers / Regions / Statuses.
    Returns name, deals count and amount when available.
    """
    if dim not in {"products","categories","regions","customers","providers","statuses"}:
        raise HTTPException(400,"dimension")
    v=platform(key)
    with ro(v["db"]) as c:
        if v["kind"]=="completed":
            if amount_min is not None or amount_max is not None:
                return _amount_completed_dimension(c,key,dim,limit,q,product_q,date_from,date_to,region,amount_min,amount_max)
            # SHOP / NATIONAL DIMENSION CACHE V1
            if key in ("shop","national"):
                cache_db=Path(
                    r"C:\Qudrat\xarid_uz\UZBI\data"
                    rf"\{key}_dimension_cache.sqlite3"
                )

                if cache_db.exists():
                    try:
                        with ro(cache_db) as cc:
                            meta=dict(
                                cc.execute(
                                    "select key,value from cache_meta"
                                ).fetchall()
                            )

                            data_min=meta.get("data_date_min")
                            data_max=meta.get("data_date_max")

                            cache_period_ok=(
                                date_from
                                and date_to
                                and date_from[:4]=="2026"
                                and date_to[:4]=="2026"
                            )

                            if cache_period_ok:
                                full_period=(
                                    data_min
                                    and data_max
                                    and date_from[:10]<=data_min
                                    and date_to[:10]>=data_max
                                    and not region
                                    and not product_q
                                )

                                summary_tables={
                                    "products":
                                        "dimension_summary_products",
                                    "categories":
                                        "dimension_summary_categories",
                                    "regions":
                                        "dimension_summary_regions",
                                    "customers":
                                        "dimension_summary_customers",
                                    "providers":
                                        "dimension_summary_providers",
                                    "statuses":
                                        "dimension_summary_statuses",
                                }

                                fact_columns={
                                    "products":"product_name",
                                    "categories":"category_name",
                                    "regions":"region_name",
                                    "customers":"customer_name",
                                    "providers":"provider_name",
                                }

                                q_clean=(q or "").strip()
                                q_digits="".join(
                                    ch for ch in q_clean
                                    if ch.isdigit()
                                )
                                q_is_tin=(
                                    q_clean.isdigit()
                                    and len(q_digits)>=8
                                )

                                if (
                                    full_period
                                    and not q_is_tin
                                    and dim not in (
                                        "products",
                                        "customers",
                                        "providers",
                                        "statuses",
                                    )
                                ):
                                    table=summary_tables[dim]
                                    where=["1=1"]
                                    params=[]

                                    if q:
                                        where.append("name like ?")
                                        params.append(
                                            f"%{q.strip()}%"
                                        )

                                    sql=f"""
                                        select *
                                        from {table}
                                        where {' and '.join(where)}
                                        order by amount desc
                                        limit ?
                                    """

                                    return [
                                        dict(row)
                                        for row in cc.execute(
                                            sql,
                                            params+[limit],
                                        )
                                    ]

                                if dim=="statuses":
                                    status_full_period=(
                                        full_period
                                        and not q
                                    )

                                    if status_full_period:
                                        sql="""
                                            select *
                                            from dimension_summary_statuses
                                            order by amount desc
                                            limit ?
                                        """

                                        return [
                                            dict(row)
                                            for row in cc.execute(
                                                sql,
                                                [limit],
                                            )
                                        ]

                                    where=[
                                        "deal_date>=?",
                                        "deal_date<=?",
                                        "status_name<>''",
                                    ]
                                    params=[
                                        date_from[:10],
                                        date_to[:10],
                                    ]

                                    if region:
                                        where.append(
                                            "region_name=?"
                                        )
                                        params.append(region)

                                    if product_q:
                                        where.append(
                                            "product_name like ?"
                                        )
                                        params.append(
                                            f"%{product_q.strip()}%"
                                        )

                                    if q:
                                        q_clean=q.strip()
                                        q_is_tin=(
                                            q_clean.isdigit()
                                            and len(q_clean)>=8
                                        )

                                        if key=="shop" and q_is_tin:
                                            where.append("""(
                                                customer_inn=? or
                                                provider_inn=?
                                            )""")
                                            params.extend([
                                                q_clean,
                                                q_clean,
                                            ])
                                        else:
                                            where.append(
                                                "status_name like ?"
                                            )
                                            params.append(
                                                f"%{q_clean}%"
                                            )

                                    where_sql=" and ".join(where)

                                    sql=f"""
                                        select
                                            status_name name,
                                            count(*) deals,
                                            coalesce(
                                                sum(amount),0
                                            ) amount,
                                            coalesce(
                                                avg(amount),0
                                            ) avg_amount
                                        from dimension_status_facts
                                        where {where_sql}
                                        group by status_name
                                        order by amount desc
                                        limit ?
                                    """

                                    return [
                                        dict(row)
                                        for row in cc.execute(
                                            sql,
                                            params+[limit],
                                        )
                                    ]

                                col=fact_columns[dim]
                                where=[
                                    "deal_date>=?",
                                    "deal_date<=?",
                                    f"{col}<>''",
                                ]
                                params=[
                                    date_from[:10],
                                    date_to[:10],
                                ]

                                if region:
                                    where.append("region_name=?")
                                    params.append(region)

                                if product_q:
                                    where.append(
                                        "product_name like ?"
                                    )
                                    params.append(
                                        f"%{product_q.strip()}%"
                                    )

                                if q:
                                    q_clean=q.strip()
                                    q_digits="".join(
                                        ch for ch in q_clean
                                        if ch.isdigit()
                                    )
                                    q_is_tin=(
                                        q_clean.isdigit()
                                        and len(q_digits)>=8
                                    )
                                    qq=f"%{q_clean}%"

                                    if q_is_tin:
                                        where.append("""(
                                            customer_inn=? or
                                            provider_inn=?
                                        )""")
                                        params.extend([
                                            q_clean,
                                            q_clean,
                                        ])
                                    else:
                                        if key=="shop" and dim=="products":
                                            product_names=_resolve_shop_product_names(
                                                q_clean,
                                                limit=50
                                            )

                                            if product_names:
                                                placeholders=",".join(
                                                    "?" for _ in product_names
                                                )
                                                where.append(
                                                    f"product_name in ({placeholders})"
                                                )
                                                params.extend(product_names)
                                            else:
                                                where.append("1=0")
                                        else:
                                            where.append(f"{col} like ?")
                                            params.append(qq)

                                where_sql=" and ".join(where)

                                if dim=="products":
                                    sql=f"""
                                        select
                                            deal_id,
                                            lot_id,
                                            lot_display_no,
                                            deal_date,
                                            product_name name,
                                            unit_name,
                                            quantity,
                                            case
                                                when coalesce(
                                                    quantity,0
                                                )<>0
                                                then amount/quantity
                                            end unit_price,
                                            amount,
                                            region_name,
                                            customer_name,
                                            customer_inn,
                                            provider_name,
                                            provider_inn
                                        from dimension_facts
                                        where {where_sql}
                                        order by amount desc
                                        limit ?
                                    """
                                elif dim=="categories":
                                    sql=f"""
                                        select
                                            category_name name,
                                            count(*) deals,
                                            coalesce(
                                                sum(amount),0
                                            ) amount,
                                            coalesce(
                                                avg(amount),0
                                            ) avg_amount
                                        from dimension_facts
                                        where {where_sql}
                                        group by category_name
                                        order by amount desc
                                        limit ?
                                    """
                                elif dim=="regions":
                                    sql=f"""
                                        select
                                            region_name name,
                                            count(*) deals,
                                            coalesce(
                                                sum(amount),0
                                            ) amount,
                                            coalesce(
                                                avg(amount),0
                                            ) avg_amount
                                        from dimension_facts
                                        where {where_sql}
                                        group by region_name
                                        order by amount desc
                                        limit ?
                                    """
                                elif dim=="customers":
                                    sql=f"""
                                        select
                                            deal_id,
                                            lot_id,
                                            lot_display_no,
                                            deal_date,
                                            customer_name name,
                                            customer_inn inn,
                                            unit_name,
                                            quantity,
                                            amount,
                                            region_name
                                        from dimension_facts
                                        where {where_sql}
                                        order by amount desc
                                        limit ?
                                    """
                                else:
                                    sql=f"""
                                        select
                                            deal_id,
                                            lot_id,
                                            lot_display_no,
                                            deal_date,
                                            provider_name name,
                                            provider_inn inn,
                                            unit_name,
                                            quantity,
                                            amount,
                                            region_name
                                        from dimension_facts
                                        where {where_sql}
                                        order by amount desc
                                        limit ?
                                    """

                                return [
                                    dict(row)
                                    for row in cc.execute(
                                        sql,
                                        params+[limit],
                                    )
                                ]

                    except sqlite3.Error:
                        # Safe fallback to the original SHOP DB.
                        pass

            mp={"products":"product_name","categories":"category_name","regions":"customer_region_name",
                "customers":"customer_name","providers":"provider_name","statuses":"deal_status_name"}
            col=mp[dim]
            where=f"coalesce({col},'')<>''"
            p=[]
            if q:
                q_clean=q.strip()
                q_digits="".join(ch for ch in q_clean if ch.isdigit())
                q_is_tin=q_clean.isdigit() and len(q_digits)>=8
                qq=f"%{q_clean}%"

                if key=="shop" and q_is_tin:
                    where += """ and (
                        customer_inn like ? or
                        provider_inn like ?
                    )"""
                    p.extend([qq,qq])
                else:
                    where += f" and {col} like ?"
                    p.append(qq)
            sql=f"""select {col} name,count(*) deals,coalesce(sum(deal_cost),0) amount,
                    coalesce(avg(deal_cost),0) avg_amount
                    from completed_deals where {where}
                    group by 1 order by amount desc limit ?"""
            p.append(limit)
            return [dict(x) for x in _market_amount_execute(c,key,amount_min,amount_max,date_from,date_to,region,sql,p)]
        if v["kind"]=="auction":
            success=(
                "\u041f\u043e\u0441\u0442\u0430\u0432\u043a\u0430 \u0441\u043e\u0432\u0435\u0440\u0448\u0435\u043d\u0430",
                "\u041e\u043f\u043b\u0430\u0447\u0435\u043d\u0430",
                "\u041f\u043e\u0441\u0442\u0430\u0432\u043a\u0430 \u0441\u043e\u0432\u0435\u0440\u0448\u0435\u043d\u0430 (\u0412\u043e\u0441\u0441\u0442\u0430\u043d\u043e\u0432\u043b\u0435\u043d\u0430)",
            )

            # Status browser intentionally keeps all statuses.
            if dim=="statuses":
                where=["coalesce(d.status,'')<>''"]
                p=[]
                if date_from:
                    where.append("substr(d.deal_date,1,10)>=?")
                    p.append(date_from[:10])
                if date_to:
                    where.append("substr(d.deal_date,1,10)<=?")
                    p.append(date_to[:10])
                if region:
                    where.append("d.region_name=?")
                    p.append(region)

                if product_q and "auction_products" in tables(c):
                    where.append("""exists (
                        select 1
                        from auction_products ap
                        where ap.completed_deal_id=d.id
                          and ap.product_name like ?
                    )""")
                    p.append(f"%{product_q.strip()}%")

                if q:
                    qq=f"%{q.strip()}%"
                    where.append("""(
                        coalesce(d.lot_number,'') like ? or
                        coalesce(d.category_name,'') like ? or
                        coalesce(d.customer_name,'') like ? or
                        coalesce(d.customer_tin,'') like ? or
                        coalesce(d.provider_name,'') like ? or
                        coalesce(d.provider_tin,'') like ? or
                        exists (
                            select 1
                            from auction_products aq
                            where aq.completed_deal_id=d.id
                              and coalesce(aq.product_name,'') like ?
                        )
                    )""")
                    p.extend([qq,qq,qq,qq,qq,qq,qq])

                w=" and ".join(where)
                sql=f"""select d.status name,
                               count(*) deals,
                               coalesce(sum(d.contract_price),0) amount,
                               coalesce(avg(d.contract_price),0) avg_amount
                        from completed_deals d
                        where {w}
                        group by d.status
                        order by amount desc
                        limit ?"""
                return [dict(x) for x in _market_amount_execute(c,key,amount_min,amount_max,date_from,date_to,region,sql,p+[limit])]

            where=["d.status in (?,?,?)"]
            p=list(success)

            if date_from:
                where.append("substr(d.deal_date,1,10)>=?")
                p.append(date_from[:10])
            if date_to:
                where.append("substr(d.deal_date,1,10)<=?")
                p.append(date_to[:10])
            if region:
                where.append("d.region_name=?")
                p.append(region)

            # AUCTION MAIN SEARCH:
            # digits -> customer/provider TIN
            # text   -> product name
            q_is_tin = False

            if q:
                q_clean = q.strip()
                q_digits = "".join(ch for ch in q_clean if ch.isdigit())
                q_is_tin = q_clean.isdigit() and len(q_digits) >= 8
                qq = f"%{q_clean}%"

                if q_is_tin:
                    # TIN search: find deals belonging to customer/provider.
                    # Products dimension will then show ALL products
                    # from those matching deals.
                    where.append("""(
                        coalesce(d.customer_tin,'') like ? or
                        coalesce(d.provider_tin,'') like ?
                    )""")
                    p.extend([qq, qq])

                elif dim=="products":
                    # Product-name search.
                    where.append("""exists (
                        select 1
                        from auction_products aq
                        where aq.completed_deal_id=d.id
                          and coalesce(aq.product_name,'') like ?
                    )""")
                    p.append(qq)

                else:
                    # For other dimensions a text query means product name.
                    where.append("""exists (
                        select 1
                        from auction_products aq
                        where aq.completed_deal_id=d.id
                          and coalesce(aq.product_name,'') like ?
                    )""")
                    p.append(qq)

            w=" and ".join(where)

            if dim=="products" and "auction_products" in tables(c):
                extra=""
                pp=list(p)

                # For product-name search filter the actual product row too.
                # For TIN search do NOT filter product_name:
                # show every product belonging to matching customer/provider.
                if q and not q_is_tin:
                    extra += " and coalesce(p.product_name,'') like ?"
                    pp.append(f"%{q.strip()}%")

                sql=f"""select
                               d.id deal_id,
                               d.lot_number lot_display_no,
                               d.deal_date deal_date,
                               p.product_name name,
                               coalesce(nullif(trim(p.unit_name),''),'') unit_name,
                               p.quantity quantity,
                               case
                                   when coalesce(p.quantity,0)<>0
                                   then coalesce(p.total_amount,0)/p.quantity
                                   else null
                               end unit_price,
                               coalesce(p.total_amount,0) amount,
                               d.region_name region_name,
                               d.customer_name customer_name,
                               d.customer_tin customer_inn,
                               d.provider_name provider_name,
                               d.provider_tin provider_inn
                        from auction_products p
                        join completed_deals d on d.id=p.completed_deal_id
                        where {w}
                          and coalesce(p.product_name,'')<>''
                          {extra}
                        order by amount desc
                        limit ?"""
                return [dict(x) for x in _market_amount_execute(c,key,amount_min,amount_max,date_from,date_to,region,sql,pp+[limit])]

            if dim=="categories":
                extra=""
                pp=list(p)

                sql=f"""select d.category_name name,
                               count(*) deals,
                               coalesce(sum(d.contract_price),0) amount,
                               coalesce(avg(d.contract_price),0) avg_amount
                        from completed_deals d
                        where {w}
                          and coalesce(d.category_name,'')<>''
                          {extra}
                        group by d.category_name
                        order by amount desc
                        limit ?"""
                return [dict(x) for x in _market_amount_execute(c,key,amount_min,amount_max,date_from,date_to,region,sql,pp+[limit])]

            if dim=="customers":
                extra=""
                pp=list(p)

                if product_q and "auction_products" in tables(c):
                    extra += """ and exists (
                        select 1
                        from auction_products ap
                        where ap.completed_deal_id=d.id
                          and ap.product_name like ?
                    )"""
                    pp.append(f"%{product_q.strip()}%")


                sql=f"""select d.customer_name name,
                               coalesce(nullif(trim(d.customer_tin),''),'') customer_tin,
                               count(*) deals,
                               coalesce(sum(d.contract_price),0) amount,
                               coalesce(avg(d.contract_price),0) avg_amount
                        from completed_deals d
                        where {w}
                          and coalesce(d.customer_name,'')<>''
                          {extra}
                        group by d.customer_name,
                                 coalesce(nullif(trim(d.customer_tin),''),'')
                        order by amount desc
                        limit ?"""
                return [dict(x) for x in _market_amount_execute(c,key,amount_min,amount_max,date_from,date_to,region,sql,pp+[limit])]

            if dim=="providers":
                extra=""
                pp=list(p)

                if product_q and "auction_products" in tables(c):
                    extra += """ and exists (
                        select 1
                        from auction_products ap
                        where ap.completed_deal_id=d.id
                          and ap.product_name like ?
                    )"""
                    pp.append(f"%{product_q.strip()}%")


                sql=f"""select d.provider_name name,
                               coalesce(nullif(trim(d.provider_tin),''),'') provider_tin,
                               count(*) deals,
                               coalesce(sum(d.contract_price),0) amount,
                               coalesce(avg(d.contract_price),0) avg_amount
                        from completed_deals d
                        where {w}
                          and coalesce(d.provider_name,'')<>''
                          {extra}
                        group by d.provider_name,
                                 coalesce(nullif(trim(d.provider_tin),''),'')
                        order by amount desc
                        limit ?"""
                return [dict(x) for x in _market_amount_execute(c,key,amount_min,amount_max,date_from,date_to,region,sql,pp+[limit])]

            mp={
                "regions":"region_name",
            }
            col=mp.get(dim)
            if not col:
                return []

            extra=""
            pp=list(p)

            if dim=="regions" and product_q and "auction_products" in tables(c):
                extra += """ and exists (
                    select 1
                    from auction_products ap
                    where ap.completed_deal_id=d.id
                      and ap.product_name like ?
                )"""
                pp.append(f"%{product_q.strip()}%")


            sql=f"""select d.{col} name,
                           count(*) deals,
                           coalesce(sum(d.contract_price),0) amount,
                           coalesce(avg(d.contract_price),0) avg_amount
                    from completed_deals d
                    where {w}
                      and coalesce(d.{col},'')<>''
                      {extra}
                    group by d.{col}
                    order by amount desc
                    limit ?"""
            return [dict(x) for x in _market_amount_execute(c,key,amount_min,amount_max,date_from,date_to,region,sql,pp+[limit])]

        if v["kind"]=="etender":
            mp={"products":"category_name","categories":"category_name","customers":"customer_name",
                "providers":"provider_name","statuses":"deal_status_name"}
            col=mp.get(dim)
            if not col:return []
            where=f"coalesce({col},'')<>''";p=[]
            if q: where+=f" and {col} like ?";p.append(f"%{q}%")
            p.append(limit)
            return [dict(x) for x in _market_amount_execute(c,key,amount_min,amount_max,date_from,date_to,region,
                f"""select {col} name,count(*) deals,coalesce(sum(deal_cost),0) amount,
                    coalesce(avg(deal_cost),0) avg_amount
                    from completed_deals where {where} group by 1 order by amount desc limit ?""",p)]
        if v["kind"]=="shaffof":
            mp={"products":"service_type","categories":"object_type","regions":"region","customers":"customer_name",
                "providers":"winner_name","statuses":"status_name"}
            col=mp.get(dim)
            if not col:return []
            where=f"coalesce({col},'')<>''";p=[]
            if q: where+=f" and {col} like ?";p.append(f"%{q}%")
            p.append(limit)
            return [dict(x) for x in _market_amount_execute(c,key,amount_min,amount_max,date_from,date_to,region,
                f"""select {col} name,count(*) deals,coalesce(sum(coalesce(winner_offer,start_price)),0) amount,
                    coalesce(avg(coalesce(winner_offer,start_price)),0) avg_amount
                    from tenders where {where} group by 1 order by amount desc limit ?""",p)]
        if v["kind"]=="xt_auction_completed":
            return top(
                key=key,
                dim=dim,
                limit=limit,
                date_from=date_from,
                date_to=date_to,
                region=region,
                q=q,amount_min=amount_min,amount_max=amount_max,
            )

        if v["kind"]=="xt_national_completed":
            if dim=="products":
                where = "coalesce(product_name,'')<>''"
                params = []

                if date_from:
                    where += " and close_date>=?"
                    params.append(date_from)

                if date_to:
                    where += " and close_date<=?"
                    params.append(date_to)

                if region:
                    where += " and region_name=?"
                    params.append(region)

                if q:
                    where += " and product_name like ?"
                    params.append(f"%{q}%")

                params.append(limit)

                return [dict(x) for x in _market_amount_execute(c,key,amount_min,amount_max,date_from,date_to,region,
                    f"""select
                        product_name name,
                        unit_name,
                        count(distinct proc_id) deals,
                        coalesce(sum(quantity),0) quantity,
                        coalesce(sum(total_amount),0) amount
                    from completed_lots
                    where {where}
                    group by product_name,unit_name
                    order by amount desc
                    limit ?""",
                    params
                )]

            if dim=="providers":
                where = "nullif(trim(winner_name),'') is not null"
                params = []
                if date_from:
                    where += " and close_date>=?"
                    params.append(date_from)
                if date_to:
                    where += " and close_date<=?"
                    params.append(date_to)
                if region:
                    where += " and region_name=?"
                    params.append(region)
                if q:
                    where += " and CASEFOLD(winner_name) like CASEFOLD(?)"
                    params.append(f"%{q.strip()}%")
                params.append(limit)
                return [dict(x) for x in _market_amount_execute(c,key,amount_min,amount_max,date_from,date_to,region,
                    f"""select trim(winner_name) name,
                        count(distinct proc_id) deals,
                        coalesce(sum(total_amount),0) amount,
                        coalesce(sum(total_amount),0) / nullif(count(distinct proc_id),0) avg_amount
                    from completed_lots
                    where {where}
                    group by trim(winner_name)
                    order by amount desc
                    limit ?""", params)]

            if dim=="regions":
                where = "nullif(trim(region_name),'') is not null"
                params = []
                if date_from:
                    where += " and close_date>=?"
                    params.append(date_from)
                if date_to:
                    where += " and close_date<=?"
                    params.append(date_to)
                if region:
                    where += " and region_name=?"
                    params.append(region)
                if q:
                    where += " and CASEFOLD(region_name) like CASEFOLD(?)"
                    params.append(f"%{q.strip()}%")
                params.append(limit)
                return [dict(x) for x in _market_amount_execute(c,key,amount_min,amount_max,date_from,date_to,region,
                    f"""select trim(region_name) name,
                        count(distinct proc_id) deals,
                        coalesce(sum(total_amount),0) amount,
                        coalesce(sum(total_amount),0) / nullif(count(distinct proc_id),0) avg_amount
                    from completed_lots
                    where {where}
                    group by trim(region_name)
                    order by amount desc
                    limit ?""", params)]

            if dim=="statuses":
                where = "nullif(trim(status),'') is not null"
                params = []
                if date_from:
                    where += " and close_date>=?"
                    params.append(date_from)
                if date_to:
                    where += " and close_date<=?"
                    params.append(date_to)
                if region:
                    where += " and region_name=?"
                    params.append(region)
                if q:
                    where += " and CASEFOLD(status) like CASEFOLD(?)"
                    params.append(f"%{q.strip()}%")
                params.append(limit)
                return [dict(x) for x in _market_amount_execute(c,key,amount_min,amount_max,date_from,date_to,region,
                    f"""select trim(status) name,
                        count(distinct proc_id) deals,
                        coalesce(sum(total_amount),0) amount,
                        coalesce(sum(total_amount),0) / nullif(count(distinct proc_id),0) avg_amount
                    from completed_lots
                    where {where}
                    group by trim(status)
                    order by amount desc
                    limit ?""", params)]

            # Buyer data is hidden/dropped in the current XT National source rows.
            if dim=="customers":
                return []

            return []

        if v["kind"]=="xt":
            if dim in {"products","categories"}:
                col="product_name" if dim=="products" else "category_name"
                where=f"coalesce({col},'')<>''";p=[]
                if q: where+=f" and {col} like ?";p.append(f"%{q}%")
                p.append(limit)
                return [dict(x) for x in _market_amount_execute(c,key,amount_min,amount_max,date_from,date_to,region,
                    f"""select {col} name,count(*) deals,coalesce(sum(total_price),0) amount,
                        coalesce(avg(total_price),0) avg_amount
                        from items where {where} group by 1 order by amount desc limit ?""",p)]
            mp={"customers":"customer_name","statuses":"status"}
            col=mp.get(dim)
            if not col:return []
            where=f"coalesce({col},'')<>''";p=[]
            if q: where+=f" and {col} like ?";p.append(f"%{q}%")
            p.append(limit)
            return [dict(x) for x in _market_amount_execute(c,key,amount_min,amount_max,date_from,date_to,region,
                f"""select {col} name,count(*) deals,coalesce(sum(total_amount),0) amount,
                    coalesce(avg(total_amount),0) avg_amount
                    from tenders where {where} group by 1 order by amount desc limit ?""",p)]
        if v["kind"]=="mc":
            if dim=="products" and "tender_items" in tables(c):
                where="coalesce(product_name,'')<>''";p=[]
                if q:where+=" and product_name like ?";p.append(f"%{q}%")
                p.append(limit)
                return [dict(x) for x in _market_amount_execute(c,key,amount_min,amount_max,date_from,date_to,region,
                    """select product_name name,count(*) deals,coalesce(sum(total_price),0) amount,
                       coalesce(avg(total_price),0) avg_amount from tender_items
                       where """+where+""" group by 1 order by amount desc limit ?""",p)]
            mp={"categories":"service_type_name","regions":"region_name_uz","customers":"customer_name","statuses":"status"}
            col=mp.get(dim)
            if not col:return []
            where=f"coalesce({col},'')<>''";p=[]
            if q:where+=f" and {col} like ?";p.append(f"%{q}%")
            p.append(limit)
            return [dict(x) for x in _market_amount_execute(c,key,amount_min,amount_max,date_from,date_to,region,
                f"""select {col} name,count(*) deals,coalesce(sum(start_price),0) amount,
                    coalesce(avg(start_price),0) avg_amount
                    from tenders where {where} group by 1 order by amount desc limit ?""",p)]
    return []


@app.get("/api/shop/product-regions")
def shop_product_regions(
    product_name: str,
    unit_name: str = "",
    date_from: Optional[str] = None,
    date_to: Optional[str] = None,
):
    cache_db = Path(
        r"C:\Qudrat\xarid_uz\UZBI\data"
        r"\shop_dimension_cache.sqlite3"
    )

    if not cache_db.exists():
        raise HTTPException(503, "shop dimension cache unavailable")

    with ro(cache_db) as c:
        meta = dict(
            c.execute(
                "select key,value from cache_meta"
            ).fetchall()
        )

        data_min = meta.get("data_date_min")
        data_max = meta.get("data_date_max")

        full_period = (
            data_min
            and data_max
            and date_from
            and date_to
            and date_from[:10] <= data_min
            and date_to[:10] >= data_max
        )

        params = [product_name, unit_name]

        if full_period:
            sql = """
                select
                    region_name,
                    deals,
                    quantity,
                    amount,
                    case
                        when coalesce(quantity,0)<>0
                        then amount/quantity
                    end as unit_price
                from dimension_product_regions
                where product_name=?
                  and unit_name=?
                order by quantity desc, amount desc
            """
        else:
            where = [
                "product_name=?",
                "unit_name=?",
                "region_name<>''",
            ]

            if date_from:
                where.append("deal_date>=?")
                params.append(date_from[:10])

            if date_to:
                where.append("deal_date<=?")
                params.append(date_to[:10])

            sql = f"""
                select
                    region_name,
                    count(*) as deals,
                    coalesce(sum(quantity),0) as quantity,
                    coalesce(sum(amount),0) as amount,
                    case
                        when coalesce(sum(quantity),0)<>0
                        then sum(amount)/sum(quantity)
                    end as unit_price
                from dimension_facts
                where {' and '.join(where)}
                group by region_name
                order by quantity desc, amount desc
            """

        rows = [
            dict(row)
            for row in c.execute(sql, params)
        ]

        return {
            "product_name": product_name,
            "unit_name": unit_name,
            "regions": rows,
            "region_count": len(rows),
            "deals": sum(
                int(row.get("deals") or 0)
                for row in rows
            ),
            "quantity": sum(
                float(row.get("quantity") or 0)
                for row in rows
            ),
            "amount": sum(
                float(row.get("amount") or 0)
                for row in rows
            ),
        }


@app.get("/api/{key}/export.xlsx")
def export_xlsx(key:str,limit:int=Query(5000,ge=1,le=20000),date_from:Optional[str]=None,date_to:Optional[str]=None,region:Optional[str]=None,q:Optional[str]=None,amount_min:Optional[str]=None,amount_max:Optional[str]=None):
    amount_min,amount_max=_market_amount_bounds(key,amount_min,amount_max)
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
    except Exception as e:
        raise HTTPException(500,"openpyxl Р В Р Р‹Р РЋРІР‚С”Р В Р Р‹Р В РІР‚С™Р В Р’В Р В РІР‚В¦Р В Р’В Р вЂ™Р’В°Р В Р Р‹Р Р†Р вЂљРЎв„ўР В Р’В Р РЋРІР‚ВР В Р’В Р вЂ™Р’В»Р В Р’В Р РЋР’ВР В Р’В Р вЂ™Р’В°Р В Р’В Р РЋРІР‚вЂњР В Р’В Р вЂ™Р’В°Р В Р’В Р В РІР‚В¦")
    data=records(key,limit=limit,offset=0,q=q,date_from=date_from,date_to=date_to,region=region,amount_min=amount_min,amount_max=amount_max)
    wb=Workbook()
    ws=wb.active
    ws.title="BIData"
    ws.append(data["columns"])
    for cell in ws[1]:
        cell.font=Font(bold=True)
        cell.alignment=Alignment(horizontal="center")
    for row in data["rows"]:
        ws.append([row.get(c) for c in data["columns"]])
    ws.freeze_panes="A2"
    for col in ws.columns:
        letter=col[0].column_letter
        width=min(max(12,max(len(str(x.value or "")) for x in col[:200])+2),45)
        ws.column_dimensions[letter].width=width
    bio=io.BytesIO()
    wb.save(bio)
    bio.seek(0)
    return StreamingResponse(
        bio,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition":f'attachment; filename="uzbi_{key}_export.xlsx"'}
    )


def pending_source(key:str):
    if key not in PENDING:
        raise HTTPException(404,"Р В Р’В Р Р†Р вЂљР’ВР В Р Р‹Р РЋРІР‚Сљ Р В Р’В Р РЋРІР‚вЂќР В Р’В Р вЂ™Р’В»Р В Р’В Р вЂ™Р’В°Р В Р Р‹Р Р†Р вЂљРЎв„ўР В Р Р‹Р Р†Р вЂљРЎвЂєР В Р’В Р РЋРІР‚СћР В Р Р‹Р В РІР‚С™Р В Р’В Р РЋР’ВР В Р’В Р вЂ™Р’В° Р В Р Р‹Р РЋРІР‚СљР В Р Р‹Р Р†Р вЂљР Р‹Р В Р Р‹Р РЋРІР‚СљР В Р’В Р В РІР‚В¦ Р В Р Р‹Р Р†Р вЂљРЎвЂєР В Р’В Р вЂ™Р’В°Р В Р’В Р РЋРІР‚СћР В Р’В Р вЂ™Р’В» Р В Р’В Р вЂ™Р’В»Р В Р’В Р РЋРІР‚СћР В Р Р‹Р Р†Р вЂљРЎв„ў Р В Р’В Р РЋР’ВР В Р’В Р вЂ™Р’В°Р В Р’В Р В РІР‚В¦Р В Р’В Р вЂ™Р’В±Р В Р’В Р вЂ™Р’В°Р В Р Р‹Р В РЎвЂњР В Р’В Р РЋРІР‚В Р В Р Р‹Р В РЎвЂњР В Р’В Р РЋРІР‚СћР В Р’В Р вЂ™Р’В·Р В Р’В Р вЂ™Р’В»Р В Р’В Р вЂ™Р’В°Р В Р’В Р В РІР‚В¦Р В Р’В Р РЋР’ВР В Р’В Р вЂ™Р’В°Р В Р’В Р РЋРІР‚вЂњР В Р’В Р вЂ™Р’В°Р В Р’В Р В РІР‚В¦")
    v=PENDING[key]
    if not v["db"].exists():
        raise HTTPException(503,f"Pending DB Р В Р Р‹Р Р†Р вЂљРЎв„ўР В Р’В Р РЋРІР‚СћР В Р’В Р РЋРІР‚вЂќР В Р’В Р РЋРІР‚ВР В Р’В Р вЂ™Р’В»Р В Р’В Р РЋР’ВР В Р’В Р вЂ™Р’В°Р В Р’В Р СћРІР‚ВР В Р’В Р РЋРІР‚В: {v['db']}")
    return v

@app.get("/api/pending/platforms")
def pending_platforms():
    out=[]
    for k,v in PENDING.items():
        row={"key":k,"title":v["title"],"kind":v["kind"],"logo":v.get("logo"),"source":v.get("source"),"trade_type":v.get("trade_type"),"path":str(v["db"]),"exists":v["db"].exists()}
        if v["db"].exists():
            try:
                with ro(v["db"]) as c:
                    t=best_table(c); row["table"]=t
                    if v["kind"]=="etender_pending" and t=="pending_lots" and "is_active" in cols(c,t):
                        row["rows"]=c.execute(
                            "select count(*) from pending_lots where coalesce(is_active,0)=1"
                        ).fetchone()[0]
                    elif v["kind"]=="xt" and t=="tenders" and "record_type" in cols(c,t):
                        row["rows"]=c.execute(
                            "select count(*) from tenders where record_type='active'"
                        ).fetchone()[0]
                    elif v["kind"]=="shaffof_pending" and t=="tenders":
                        row["rows"]=c.execute(
                            "select count(*) from tenders where status_name='Faol lot'"
                        ).fetchone()[0]
                    elif v["kind"]=="mc" and t=="tenders":
                        row["rows"]=c.execute(
                            "select count(*) from tenders where cast(status as text)='2' and winner_id is null"
                        ).fetchone()[0]
                    else:
                        row["rows"]=c.execute(
                            f"select count(*) from {ident(t)}"
                        ).fetchone()[0] if t else 0
            except Exception as e: row["error"]=str(e)
        out.append(row)
    return out

def pending_table_and_map(c, kind):
    ts=tables(c)

    if kind=="shop_pending":
        t="shop_offers" if "shop_offers" in ts else best_table(c)
        cs=cols(c,t)
        return t,{
          "id":("offer_id" if "offer_id" in cs else detect(cs,"id")),
          "display":("display_no" if "display_no" in cs else detect(cs,"display")),
          "date":("date_ini" if "date_ini" in cs else detect(cs,"date")),
          "deadline":("date_end" if "date_end" in cs else detect(cs,"close")),
          "product":("product_name" if "product_name" in cs else detect(cs,"product")),
          "amount":("total_value" if "total_value" in cs else detect(cs,"amount")),
          "customer":("customer_name" if "customer_name" in cs else detect(cs,"customer")),
          "customer_inn":("customer_inn" if "customer_inn" in cs else detect(cs,"customer_inn")),
          "region":("customer_region_name" if "customer_region_name" in cs else detect(cs,"region")),
          "status":("status_name" if "status_name" in cs else detect(cs,"status"))
        }

    if kind=="xt_auction_pending":
        t="pending_lots" if "pending_lots" in ts else best_table(c)
        cs=cols(c,t)
        return t,{
          "id":"lot_id" if "lot_id" in cs else None,
          "display":"display_no" if "display_no" in cs else ("lot_id" if "lot_id" in cs else None),
          "date":"start_date" if "start_date" in cs else detect(cs,"date"),
          "deadline":"end_date" if "end_date" in cs else detect(cs,"close"),
          "product":"category_name" if "category_name" in cs else detect(cs,"product"),
          "amount":"current_cost" if "current_cost" in cs else ("start_cost" if "start_cost" in cs else detect(cs,"amount")),
          "customer":"customer_name" if "customer_name" in cs else detect(cs,"customer"),
          "customer_inn":"customer_inn" if "customer_inn" in cs else detect(cs,"customer_inn"),
          "region":"region_name" if "region_name" in cs else detect(cs,"region"),
          "status":"status_name" if "status_name" in cs else detect(cs,"status")
        }

    if kind=="auction_pending":
        t="pending_lots" if "pending_lots" in ts else best_table(c)
        cs=cols(c,t)
        return t,{
          "id":detect(cs,"id") or ("lot_id" if "lot_id" in cs else None),
          "display":detect(cs,"display") or ("display_no" if "display_no" in cs else None),
          "date":("start_date" if "start_date" in cs else detect(cs,"date")),
          "deadline":("end_date" if "end_date" in cs else detect(cs,"close")),
          "product":detect(cs,"product") or ("category_name" if "category_name" in cs else None),
          "amount":("current_cost" if "current_cost" in cs else ("start_cost" if "start_cost" in cs else detect(cs,"amount"))),
          "customer":detect(cs,"customer"),"customer_inn":detect(cs,"customer_inn"),
          "region":detect(cs,"region"),"status":detect(cs,"status") or ("status_name" if "status_name" in cs else None)
        }
    if kind=="national_pending":
        t="national_offers" if "national_offers" in ts else best_table(c)
        cs=cols(c,t)
        return t,{
          "id":("offer_id" if "offer_id" in cs else detect(cs,"id")),
          "display":("display_no" if "display_no" in cs else detect(cs,"display")),
          "date":("date_ini" if "date_ini" in cs else detect(cs,"date")),
          "deadline":("date_end" if "date_end" in cs else detect(cs,"close")),
          "product":("product_name" if "product_name" in cs else detect(cs,"product")),
          "amount":("total_value" if "total_value" in cs else detect(cs,"amount")),
          "customer":("customer_name" if "customer_name" in cs else detect(cs,"customer")),
          "customer_inn":("customer_inn" if "customer_inn" in cs else detect(cs,"customer_inn")),
          "region":("customer_region_name" if "customer_region_name" in cs else detect(cs,"region")),
          "status":("status_name" if "status_name" in cs else detect(cs,"status"))
        }

    if kind=="xt_national_pending":
        t="national_lots" if "national_lots" in ts else best_table(c); cs=cols(c,t)
        return t,{"id":"proc_id","display":"proc_id","date":"publicated_at","deadline":"close_at","product":"product_name","amount":"totalcost","customer":None,"customer_inn":None,"region":"last_region_seen","status":"status"}
    if kind=="etender_pending":
        t="pending_lots" if "pending_lots" in ts else best_table(c);cs=cols(c,t)
        return t,{k:detect(cs,k) for k in ["id","display","date","close","product","amount","customer","customer_inn","region","status"]} | {"deadline":detect(cs,"close")}
    if kind=="xt":
        return "tenders",{"id":"proc_id","display":"lot_number","date":"publicated_at","deadline":"close_at","product":"name","amount":"total_amount","customer":"customer_name","customer_inn":None,"region":None,"status":"status"}
    if kind=="shaffof_pending":
        t="tenders" if "tenders" in ts else best_table(c);cs=cols(c,t)
        return t,{"id":("id" if "id" in cs else detect(cs,"id")),"display":("unique_name" if "unique_name" in cs else detect(cs,"display")),
        "date":("created_at" if "created_at" in cs else detect(cs,"date")),"deadline":("deadline" if "deadline" in cs else ("end_date" if "end_date" in cs else detect(cs,"close"))),
        "product":("name" if "name" in cs else ("service_type" if "service_type" in cs else detect(cs,"product"))),
        "amount":("start_price" if "start_price" in cs else detect(cs,"amount")),"customer":("customer_name" if "customer_name" in cs else detect(cs,"customer")),
        "customer_inn":("customer_inn" if "customer_inn" in cs else detect(cs,"customer_inn")),"region":("region" if "region" in cs else detect(cs,"region")),
        "status":("status_name" if "status_name" in cs else detect(cs,"status"))}
    if kind=="mc":
        return "tenders",{"id":"id","display":"unique_name","date":"confirmed_date","deadline":"placement_term","product":"name","amount":"start_price","customer":"customer_name","customer_inn":"customer_inn","region":"region_name_uz","status":"status"}
    t=best_table(c);cs=cols(c,t)
    return t,{k:detect(cs,k) for k in ["id","display","date","product","amount","customer","customer_inn","region","status"]} | {"deadline":detect(cs,"close")}

@app.get("/api/pending/{key}/list")
def pending_list(key:str,limit:int=Query(50,ge=1,le=200),offset:int=Query(0,ge=0),q:Optional[str]=None):
    v=pending_source(key)
    with ro(v["db"]) as c:
        t,m=pending_table_and_map(c,v["kind"])
        select=[]
        aliases=["id","display","date","deadline","product","amount","customer","customer_inn","region","status"]
        for a in aliases:
            col=m.get(a)
            select.append(f"{ident(col)} {ident(a)}" if col else f"null {ident(a)}")
        where="1=1";p=[]

        # Pending endpoints must expose active records only.
        if v["kind"]=="etender_pending":
            where+=" and coalesce(is_active,0)=1"
        elif v["kind"]=="xt":
            where+=" and record_type='active'"
        elif v["kind"]=="shaffof_pending":
            where+=" and status_name='Faol lot'"

        if q:
            parts=[]
            for a in ("display","product","customer","customer_inn","status"):
                col=m.get(a)
                if col: parts.append(f"cast({ident(col)} as text) like ?");p.append(f"%{q}%")
            if parts: where+=" and ("+" or ".join(parts)+")"
        order=m.get("date") or m.get("deadline") or m.get("id")
        sql=f"select {','.join(select)} from {ident(t)} where {where}"
        if order: sql+=f" order by {ident(order)} desc"
        sql+=" limit ? offset ?";p += [limit,offset]
        rows=[dict(x) for x in c.execute(sql,p)]
        return {"title":v["title"],"table":t,"rows":rows}

@app.get("/api/pending/{key}/detail/{row_id}")
def pending_detail(key:str,row_id:str):
    v=pending_source(key)
    with ro(v["db"]) as c:
        t,m=pending_table_and_map(c,v["kind"]); idc=m.get("id")
        if not idc: raise HTTPException(500,"ID Р В Р’В Р РЋР’ВР В Р’В Р вЂ™Р’В°Р В Р’В Р Р†РІР‚С›РІР‚вЂњР В Р’В Р СћРІР‚ВР В Р’В Р РЋРІР‚СћР В Р’В Р В РІР‚В¦Р В Р’В Р РЋРІР‚В Р В Р’В Р вЂ™Р’В°Р В Р’В Р В РІР‚В¦Р В Р’В Р РЋРІР‚ВР В РЎС›Р Р†Р вЂљРЎвЂќР В Р’В Р вЂ™Р’В»Р В Р’В Р вЂ™Р’В°Р В Р’В Р В РІР‚В¦Р В Р’В Р РЋР’ВР В Р’В Р вЂ™Р’В°Р В Р’В Р СћРІР‚ВР В Р’В Р РЋРІР‚В")
        if v["kind"]=="etender_pending":
            row=c.execute(
                f"select * from {ident(t)} where {ident(idc)}=? and coalesce(is_active,0)=1 limit 1",
                (row_id,)
            ).fetchone()
        elif v["kind"]=="xt":
            row=c.execute(
                f"select * from {ident(t)} where {ident(idc)}=? and record_type='active' limit 1",
                (row_id,)
            ).fetchone()
        elif v["kind"]=="shaffof_pending":
            row=c.execute(
                f"select * from {ident(t)} where {ident(idc)}=? and status_name='Faol lot' limit 1",
                (row_id,)
            ).fetchone()
        elif v["kind"]=="mc":
            row=c.execute(
                f"select * from {ident(t)} where {ident(idc)}=? and cast(status as text)='2' and winner_id is null limit 1",
                (row_id,)
            ).fetchone()
        else:
            row=c.execute(
                f"select * from {ident(t)} where {ident(idc)}=? limit 1",
                (row_id,)
            ).fetchone()

        if not row: raise HTTPException(404,"Р В Р’В Р Р†Р вЂљРЎвЂќР В Р’В Р РЋРІР‚СћР В Р Р‹Р Р†Р вЂљРЎв„ў Р В Р Р‹Р Р†Р вЂљРЎв„ўР В Р’В Р РЋРІР‚СћР В Р’В Р РЋРІР‚вЂќР В Р’В Р РЋРІР‚ВР В Р’В Р вЂ™Р’В»Р В Р’В Р РЋР’ВР В Р’В Р вЂ™Р’В°Р В Р’В Р СћРІР‚ВР В Р’В Р РЋРІР‚В")
        d=dict(row)

        # BIData canonical product code for active-lot detail card.
        _pd_name = (
            d.get("product_name")
            or d.get("name")
            or d.get("product")
            or d.get("category_name")
        )
        _pd_source_code = (
            d.get("product_code")
            or d.get("code")
            or d.get("product_id")
        )

        _pd_code, _pd_code_source = _canonical_active_product_code(
            _pd_name,
            _pd_source_code
        )

        if _pd_code:
            d["product_code"] = _pd_code
            d["product_code_source"] = _pd_code_source

        extras={}
        ts=set(tables(c))
        if v["kind"] in ("auction_pending","xt_auction_pending"):
            if "pending_products" in ts:
                extras["products"]=[dict(x) for x in c.execute("select * from pending_products where lot_id=? order by product_no",(row_id,))]
            if "pending_product_properties" in ts:
                extras["properties"]=[dict(x) for x in c.execute("select * from pending_product_properties where lot_id=? order by product_no,id",(row_id,))]
        elif v["kind"]=="national_pending":
            # UZEX National Shop active offer details.
            offer_id = row_id

            extras["properties"] = [
                dict(x) for x in c.execute(
                    "select * from national_offer_properties "
                    "where offer_id=? order by rowid",
                    (offer_id,)
                )
            ] if "national_offer_properties" in ts else []

            extras["detail_fields"] = [
                dict(x) for x in c.execute(
                    "select * from national_offer_detail_fields "
                    "where offer_id=? order by rowid",
                    (offer_id,)
                )
            ] if "national_offer_detail_fields" in ts else []

            extras["documents"] = [
                dict(x) for x in c.execute(
                    "select * from national_offer_documents "
                    "where offer_id=? order by rowid",
                    (offer_id,)
                )
            ] if "national_offer_documents" in ts else []

            extras["districts"] = [
                dict(x) for x in c.execute(
                    "select * from national_delivery_districts "
                    "where offer_id=? order by rowid",
                    (offer_id,)
                )
            ] if "national_delivery_districts" in ts else []

            # Stable aliases for the generic active-lot modal.
            if d.get("mark_name") and not d.get("brand"):
                d["brand"] = d.get("mark_name")

            if d.get("manufacturer_name") and not d.get("manufacturer"):
                d["manufacturer"] = d.get("manufacturer_name")

            if d.get("product_description") and not d.get("description"):
                d["description"] = d.get("product_description")

            extras["product_details"] = {
                "brand": d.get("mark_name"),
                "manufacturer": d.get("manufacturer_name"),
                "description": (
                    d.get("product_description")
                    or d.get("specification_text")
                    or d.get("condition_text")
                ),
                "properties": extras["properties"],
            }

        elif v["kind"]=="xt_national_pending":
            extras["properties"]=[dict(x) for x in c.execute(
                "select * from national_product_properties where proc_id=? order by prop_id",
                (row_id,)
            )] if "national_product_properties" in ts else []

            # UZBI v13: XT National product passport.
            # Keep the raw property rows, but also expose a stable normalized structure
            # for the web card (Р В Р’В Р РЋРЎв„ўР В Р’В Р вЂ™Р’В°Р В Р Р‹Р В РІР‚С™Р В Р’В Р РЋРІР‚СњР В Р’В Р вЂ™Р’В° / Р В Р’В Р вЂ™Р’ВР В Р Р‹Р Р†РІР‚С™Р’В¬Р В Р’В Р вЂ™Р’В»Р В Р’В Р вЂ™Р’В°Р В Р’В Р вЂ™Р’В± Р В Р Р‹Р Р†Р вЂљР Р‹Р В Р’В Р РЋРІР‚ВР В РЎС›Р Р†Р вЂљРЎвЂќР В Р’В Р вЂ™Р’В°Р В Р Р‹Р В РІР‚С™Р В Р Р‹Р РЋРІР‚СљР В Р’В Р В РІР‚В Р В Р Р‹Р Р†Р вЂљР Р‹Р В Р’В Р РЋРІР‚В / Р В Р’В Р РЋРЎвЂєР В Р’В Р РЋРІР‚СћР В Р’В Р В РІР‚В Р В Р’В Р вЂ™Р’В°Р В Р Р‹Р В РІР‚С™ Р В Р Р‹Р Р†Р вЂљРЎв„ўР В Р’В Р вЂ™Р’В°Р В Р’В Р В РІР‚В Р В Р Р‹Р В РЎвЂњР В Р’В Р РЋРІР‚ВР В Р Р‹Р Р†Р вЂљРЎвЂєР В Р’В Р РЋРІР‚В / other props).
            def _xtn_first(obj, names):
                low={str(k).lower():k for k in obj.keys()}
                for name in names:
                    k=low.get(str(name).lower())
                    if k is not None:
                        val=obj.get(k)
                        if val not in (None,"") and str(val).strip():
                            return str(val).strip()
                return None

            normalized_props=[]
            for pr in extras["properties"]:
                pname=_xtn_first(pr,[
                    "prop_name","property_name","name","characteristic_name",
                    "attribute_name","param_name","title"
                ])
                pvalue=_xtn_first(pr,[
                    "val_name","value_name","prop_value","property_value","value",
                    "characteristic_value","attribute_value","param_value","text"
                ])
                if pname or pvalue:
                    normalized_props.append({
                        "name":pname or "Р В Р’В Р СћРЎвЂ™Р В Р Р‹Р РЋРІР‚СљР В Р Р‹Р В РЎвЂњР В Р Р‹Р РЋРІР‚СљР В Р Р‹Р В РЎвЂњР В Р’В Р РЋРІР‚ВР В Р Р‹Р В Р РЏР В Р Р‹Р Р†Р вЂљРЎв„ў",
                        "value":pvalue or "Р В Р вЂ Р В РІР‚С™Р Р†Р вЂљРЎСљ",
                        "prop_id":pr.get("prop_id"),
                        "raw":pr
                    })

            def _xtn_prop_value(tokens):
                for pr in normalized_props:
                    nm=str(pr.get("name") or "").casefold()
                    if any(tok in nm for tok in tokens):
                        val=str(pr.get("value") or "").strip()
                        if val and val!="Р В Р вЂ Р В РІР‚С™Р Р†Р вЂљРЎСљ":
                            return val
                return None

            brand=(
                _xtn_first(d,["brand","brand_name","mark","marka","model","product_brand"])
                or _xtn_prop_value(("Р В Р’В Р РЋР’ВР В Р’В Р вЂ™Р’В°Р В Р Р‹Р В РІР‚С™Р В Р’В Р РЋРІР‚СњР В Р’В Р вЂ™Р’В°","marka","brand","Р В Р’В Р вЂ™Р’В±Р В Р Р‹Р В РІР‚С™Р В Р’В Р вЂ™Р’ВµР В Р’В Р В РІР‚В¦Р В Р’В Р СћРІР‚В","model","Р В Р’В Р РЋР’ВР В Р’В Р РЋРІР‚СћР В Р’В Р СћРІР‚ВР В Р’В Р вЂ™Р’ВµР В Р’В Р вЂ™Р’В»Р В Р Р‹Р В Р вЂ°"))
            )
            manufacturer=(
                _xtn_first(d,[
                    "manufacturer","manufacturer_name","producer","producer_name",
                    "maker","vendor","ishlab_chiqaruvchi"
                ])
                or _xtn_prop_value((
                    "Р В Р’В Р РЋРІР‚ВР В Р Р‹Р Р†РІР‚С™Р’В¬Р В Р’В Р вЂ™Р’В»Р В Р’В Р вЂ™Р’В°Р В Р’В Р вЂ™Р’В± Р В Р Р‹Р Р†Р вЂљР Р‹Р В Р’В Р РЋРІР‚ВР В РЎС›Р Р†Р вЂљРЎвЂќР В Р’В Р вЂ™Р’В°Р В Р Р‹Р В РІР‚С™Р В Р Р‹Р РЋРІР‚СљР В Р’В Р В РІР‚В Р В Р Р‹Р Р†Р вЂљР Р‹Р В Р’В Р РЋРІР‚В","ishlab chiqaruvchi","Р В Р’В Р РЋРІР‚вЂќР В Р Р‹Р В РІР‚С™Р В Р’В Р РЋРІР‚СћР В Р’В Р РЋРІР‚ВР В Р’В Р вЂ™Р’В·Р В Р’В Р В РІР‚В Р В Р’В Р РЋРІР‚СћР В Р’В Р СћРІР‚ВР В Р’В Р РЋРІР‚ВР В Р Р‹Р Р†Р вЂљРЎв„ўР В Р’В Р вЂ™Р’ВµР В Р’В Р вЂ™Р’В»Р В Р Р‹Р В Р вЂ°",
                    "manufacturer","producer"
                ))
            )
            description=(
                _xtn_first(d,[
                    "description","product_description","spec","specification",
                    "description_uz","description_ru","technical_specification",
                    "characteristics","product_spec"
                ])
                or _xtn_prop_value((
                    "Р В Р Р‹Р Р†Р вЂљРЎв„ўР В Р’В Р РЋРІР‚СћР В Р’В Р В РІР‚В Р В Р’В Р вЂ™Р’В°Р В Р Р‹Р В РІР‚С™Р В Р’В Р В РІР‚В¦Р В Р’В Р РЋРІР‚ВР В Р’В Р В РІР‚В¦Р В Р’В Р РЋРІР‚вЂњ Р В Р Р‹Р Р†Р вЂљРЎв„ўР В Р’В Р вЂ™Р’В°Р В Р’В Р В РІР‚В Р В Р Р‹Р В РЎвЂњР В Р’В Р РЋРІР‚ВР В Р Р‹Р Р†Р вЂљРЎвЂєР В Р’В Р РЋРІР‚В","Р В Р Р‹Р Р†Р вЂљРЎв„ўР В Р’В Р РЋРІР‚СћР В Р’В Р В РІР‚В Р В Р’В Р вЂ™Р’В°Р В Р Р‹Р В РІР‚С™ Р В Р Р‹Р Р†Р вЂљРЎв„ўР В Р’В Р вЂ™Р’В°Р В Р’В Р В РІР‚В Р В Р Р‹Р В РЎвЂњР В Р’В Р РЋРІР‚ВР В Р Р‹Р Р†Р вЂљРЎвЂєР В Р’В Р РЋРІР‚В","tovarning tavsifi",
                    "tovar tavsifi","Р В Р’В Р РЋРІР‚СћР В Р’В Р РЋРІР‚вЂќР В Р’В Р РЋРІР‚ВР В Р Р‹Р В РЎвЂњР В Р’В Р вЂ™Р’В°Р В Р’В Р В РІР‚В¦Р В Р’В Р РЋРІР‚ВР В Р’В Р вЂ™Р’Вµ Р В Р Р‹Р Р†Р вЂљРЎв„ўР В Р’В Р РЋРІР‚СћР В Р’В Р В РІР‚В Р В Р’В Р вЂ™Р’В°Р В Р Р‹Р В РІР‚С™Р В Р’В Р вЂ™Р’В°","description","specification"
                ))
            )

            extras["product_details"]={
                "brand":brand,
                "manufacturer":manufacturer,
                "description":description,
                "properties":[
                    {"name":x["name"],"value":x["value"],"prop_id":x.get("prop_id")}
                    for x in normalized_props
                ]
            }

            # Stable aliases also live on record so an existing generic modal can
            # display them even before the frontend switches to product_details.
            if brand and not d.get("brand"):
                d["brand"]=brand
            if manufacturer and not d.get("manufacturer"):
                d["manufacturer"]=manufacturer
            if description and not d.get("description"):
                d["description"]=description

            extras["regions"]=[dict(x) for x in c.execute(
                "select * from national_lot_regions where proc_id=?",
                (row_id,)
            )] if "national_lot_regions" in ts else []

            # Resolve XT National region hierarchy from national_regions.
            region_code = (
                d.get("owner_legal_area_id")
                or d.get("last_region_seen")
                or (
                    extras["regions"][0].get("area_id")
                    if extras.get("regions")
                    else None
                )
            )

            region_path = []
            if region_code and "national_regions" in ts:
                parts = str(region_code).split(".")
                candidate_ids = [
                    ".".join(parts[:i])
                    for i in range(1, len(parts)+1)
                ]

                for area_id in candidate_ids:
                    rr = c.execute(
                        """select area_id,parent_id,area_name,depth
                           from national_regions
                           where area_id=?
                           limit 1""",
                        (area_id,)
                    ).fetchone()

                    if rr:
                        region_path.append(dict(rr))

            extras["region_path"] = region_path
            extras["region_name"] = (
                " / ".join(
                    str(x.get("area_name") or "").strip()
                    for x in region_path
                    if str(x.get("area_name") or "").strip()
                )
                or None
            )
            extras["region_code"] = region_code

        elif v["kind"]=="xt":
            extras["lots"]=[dict(x) for x in c.execute("select * from lots where proc_id=?",(row_id,))] if "lots" in ts else []
            extras["items"]=[dict(x) for x in c.execute("select * from items where proc_id=?",(row_id,))] if "items" in ts else []
            extras["properties"]=[dict(x) for x in c.execute("select * from item_properties where proc_id=?",(row_id,))] if "item_properties" in ts else []
        elif v["kind"]=="mc":
            extras["items"]=[dict(x) for x in c.execute("select * from tender_items where tender_id=?",(row_id,))] if "tender_items" in ts else []
            extras["properties"]=[dict(x) for x in c.execute("select * from tender_item_props where tender_id=?",(row_id,))] if "tender_item_props" in ts else []
            extras["files"]=[dict(x) for x in c.execute("select * from tender_files where tender_id=?",(row_id,))] if "tender_files" in ts else []
        elif v["kind"]=="etender_pending":
            for tn in ("pending_products","products","items"):
                if tn in ts:
                    cc=cols(c,tn)
                    lotc="lot_id" if "lot_id" in cc else ("tender_id" if "tender_id" in cc else None)
                    if lotc:
                        extras["products"]=[dict(x) for x in c.execute(f"select * from {ident(tn)} where {ident(lotc)}=?",(row_id,))]
                        break
        # BIData organization reference: UZEX Auction customer STIR.
        # Source pending DB remains untouched.
        if key=="auction" and isinstance(d,dict):
            current_stir=d.get("customer_inn") or d.get("customer_stir")
            customer_name=(
                d.get("customer_name")
                or d.get("customer")
                or d.get("buyer_name")
                or d.get("organization_name")
            )
            if not current_stir and customer_name:
                recovered_stir=_organization_stir_lookup(customer_name)
                if recovered_stir:
                    d["customer_inn"]=recovered_stir
                    d["customer_stir_reference"]="bidata_organizations"

        return {"key":key,"title":v["title"],"table":t,"record":d,"extras":extras}

@app.get("/api/pending/{key}/detail/{row_id}/pdf")
def pending_pdf(key:str,row_id:str):
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak
    except Exception:
        raise HTTPException(500,"reportlab Р В Р Р‹Р РЋРІР‚С”Р В Р Р‹Р В РІР‚С™Р В Р’В Р В РІР‚В¦Р В Р’В Р вЂ™Р’В°Р В Р Р‹Р Р†Р вЂљРЎв„ўР В Р’В Р РЋРІР‚ВР В Р’В Р вЂ™Р’В»Р В Р’В Р РЋР’ВР В Р’В Р вЂ™Р’В°Р В Р’В Р РЋРІР‚вЂњР В Р’В Р вЂ™Р’В°Р В Р’В Р В РІР‚В¦")
    data=pending_detail(key,row_id)
    bio=io.BytesIO()
    doc=SimpleDocTemplate(bio,pagesize=landscape(A4),leftMargin=24,rightMargin=24,topMargin=24,bottomMargin=24)
    styles=getSampleStyleSheet(); story=[]
    story.append(Paragraph(f"UZBI Р В Р вЂ Р В РІР‚С™Р Р†Р вЂљРЎСљ {data['title']}",styles["Title"]))
    story.append(Paragraph(f"ID: {row_id}",styles["Heading2"]))
    story.append(Spacer(1,10))
    rec=data["record"]
    rows=[["Р В Р’В Р РЋРЎв„ўР В Р’В Р вЂ™Р’В°Р В Р’В Р Р†РІР‚С›РІР‚вЂњР В Р’В Р СћРІР‚ВР В Р’В Р РЋРІР‚СћР В Р’В Р В РІР‚В¦","Р В РЎС›Р РЋРІвЂћСћР В Р’В Р РЋРІР‚ВР В Р’В Р Р†РІР‚С›РІР‚вЂњР В Р’В Р РЋР’ВР В Р’В Р вЂ™Р’В°Р В Р Р‹Р Р†Р вЂљРЎв„ў"]]
    for k,v in rec.items():
        if "json" in k.lower(): continue
        s="" if v is None else str(v)
        if len(s)>400:s=s[:400]+"Р В Р вЂ Р В РІР‚С™Р вЂ™Р’В¦"
        rows.append([str(k),s])
    table=Table(rows,colWidths=[170,570],repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND",(0,0),(-1,0),colors.HexColor("#0B9275")),
        ("TEXTCOLOR",(0,0),(-1,0),colors.white),
        ("FONTNAME",(0,0),(-1,0),"Helvetica-Bold"),
        ("FONTSIZE",(0,0),(-1,-1),8),
        ("GRID",(0,0),(-1,-1),0.25,colors.HexColor("#DCE4EA")),
        ("VALIGN",(0,0),(-1,-1),"TOP"),
        ("ROWBACKGROUNDS",(0,1),(-1,-1),[colors.white,colors.HexColor("#F7FAFC")]),
    ]))
    story.append(table)
    for section,items in data["extras"].items():
        if not items: continue
        story.append(PageBreak());story.append(Paragraph(section.upper(),styles["Heading2"]));story.append(Spacer(1,8))
        cols0=[k for k in items[0].keys() if "json" not in k.lower()][:10]
        rr=[cols0]
        for item in items[:150]:
            rr.append([str(item.get(c,"") or "")[:120] for c in cols0])
        tt=Table(rr,repeatRows=1)
        tt.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),colors.HexColor("#14213D")),("TEXTCOLOR",(0,0),(-1,0),colors.white),("FONTSIZE",(0,0),(-1,-1),6),("GRID",(0,0),(-1,-1),0.2,colors.HexColor("#DCE4EA")),("VALIGN",(0,0),(-1,-1),"TOP")]))
        story.append(tt)
    doc.build(story);bio.seek(0)
    return StreamingResponse(bio,media_type="application/pdf",headers={"Content-Disposition":f'attachment; filename="uzbi_{key}_{row_id}.pdf"'})

@app.get("/api/{key}/schema")
def schema(key:str):
    v=platform(key)
    with ro(v["db"]) as c:
        return {"path":str(v["db"]),"tables":{t:cols(c,t) for t in tables(c)}}

app.include_router(product_catalog_router)


# === UZBI WEB DAILY DIGEST API V1 ===

def _digest_db_path(market_key: str) -> Path:
    mapping = {
        "shop": UZBI_ROOT / "data" / "shop_completed_deals.sqlite3",
        "national": UZBI_ROOT / "data" / "national_completed_deals.sqlite3",
    }
    if market_key not in mapping:
        raise HTTPException(400, "market must be shop or national")
    path = mapping[market_key]
    if not path.exists():
        raise HTTPException(404, f"Digest DB not found: {path}")
    return path


def _digest_num(value):
    try:
        return float(value or 0)
    except (TypeError, ValueError):
        return 0.0


def _digest_change(current, previous):
    current = float(current or 0)
    previous = float(previous or 0)
    if previous <= 0:
        return None
    return round((current - previous) / previous * 100.0, 2)


def _build_web_digest(market_key: str, target_date: date):
    db_path = _digest_db_path(market_key)

    day_from = target_date.isoformat()
    day_to = (target_date + timedelta(days=1)).isoformat()
    prev_from = (target_date - timedelta(days=1)).isoformat()

    with sqlite3.connect(db_path, timeout=30) as conn:
        conn.row_factory = sqlite3.Row
        conn.execute("PRAGMA busy_timeout=30000")
        conn.execute("PRAGMA query_only=ON")

        rows = conn.execute(
            """
            SELECT
                deal_id,
                lot_id,
                lot_display_no,
                deal_date,
                category_name,
                product_name,
                amount,
                unit_name,
                deal_cost,
                customer_inn,
                customer_name,
                customer_region_name,
                provider_inn,
                provider_name
            FROM completed_deals
            WHERE deal_date >= ? AND deal_date < ?
            """,
            (day_from, day_to),
        ).fetchall()

        deals = {}
        lots = set()
        customers = set()
        suppliers = set()
        products = set()

        categories = {}
        regions = {}
        customer_counter = {}
        supplier_counter = {}
        product_group = {}

        for row in rows:
            deal_key = str(row["deal_id"] or "").strip()
            lot = str(row["lot_display_no"] or row["lot_id"] or "").strip()

            customer = str(row["customer_name"] or row["customer_inn"] or "????????").strip()
            supplier = str(row["provider_name"] or row["provider_inn"] or "????????").strip()
            category = str(row["category_name"] or "????????").strip()
            region = str(row["customer_region_name"] or "????????").strip()
            product = str(row["product_name"] or "????????").strip()

            amount = _digest_num(row["amount"])
            unit = str(row["unit_name"] or "").strip()

            if deal_key and deal_key not in deals:
                deals[deal_key] = {
                    "deal_id": deal_key,
                    "lot": lot or "?",
                    "contract": _digest_num(row["deal_cost"]),
                    "customer": customer,
                    "supplier": supplier,
                    "category": category,
                    "region": region,
                }

                categories[category] = categories.get(category, 0) + 1
                regions[region] = regions.get(region, 0) + 1
                customer_counter[customer] = customer_counter.get(customer, 0) + 1
                supplier_counter[supplier] = supplier_counter.get(supplier, 0) + 1

                if customer != "????????":
                    customers.add(customer)
                if supplier != "????????":
                    suppliers.add(supplier)

            if lot:
                lots.add(lot)

            if product and product != "????????":
                pkey = product.casefold()
                products.add(pkey)

                item = product_group.setdefault(
                    pkey,
                    {
                        "name": product,
                        "quantity": 0.0,
                        "units": {},
                        "deals": set(),
                    }
                )

                if amount > 0:
                    item["quantity"] += amount

                if unit:
                    item["units"][unit] = item["units"].get(unit, 0) + 1

                if deal_key:
                    item["deals"].add(deal_key)

        deal_rows = list(deals.values())
        total_contract = sum(x["contract"] for x in deal_rows)
        deals_count = len(deal_rows)

        top_products = []
        for item in product_group.values():
            unit_name = ""
            if item["units"]:
                unit_name = max(item["units"], key=item["units"].get)

            top_products.append({
                "name": item["name"],
                "quantity": round(item["quantity"], 4),
                "unit_name": unit_name,
                "deals_count": len(item["deals"]),
            })

        top_products.sort(
            key=lambda x: (x["quantity"], x["deals_count"]),
            reverse=True,
        )

        prev_rows = conn.execute(
            """
            SELECT deal_id, deal_cost
            FROM completed_deals
            WHERE deal_date >= ? AND deal_date < ?
            """,
            (prev_from, day_from),
        ).fetchall()

        prev_deals = {}
        for row in prev_rows:
            key = str(row["deal_id"] or "").strip()
            if key and key not in prev_deals:
                prev_deals[key] = _digest_num(row["deal_cost"])

        previous_deals = len(prev_deals)
        previous_contract = sum(prev_deals.values())

    def top_counter(d, limit=3):
        return [
            {"name": k, "deals_count": v}
            for k, v in sorted(
                d.items(),
                key=lambda x: x[1],
                reverse=True
            )[:limit]
        ]

    return {
        "market": market_key,
        "title": {
            "shop": "\u042d\u043b\u0435\u043a\u0442\u0440\u043e\u043d \u0434\u045e\u043a\u043e\u043d",
            "national": "\u041c\u0438\u043b\u043b\u0438\u0439 \u0434\u045e\u043a\u043e\u043d",
        }[market_key],
        "date": target_date.isoformat(),

        "deals_count": deals_count,
        "lots_count": len(lots),
        "customers_count": len(customers),
        "suppliers_count": len(suppliers),
        "products_count": len(products),

        "total_contract": round(total_contract, 2),
        "average_deal": round(
            total_contract / deals_count if deals_count else 0,
            2
        ),

        "previous_deals": previous_deals,
        "previous_contract": round(previous_contract, 2),

        "deals_change_pct": _digest_change(
            deals_count,
            previous_deals
        ),
        "contract_change_pct": _digest_change(
            total_contract,
            previous_contract
        ),

        "top_deals": sorted(
            deal_rows,
            key=lambda x: x["contract"],
            reverse=True
        )[:3],

        "top_products": top_products[:3],
        "top_categories": top_counter(categories),
        "top_regions": top_counter(regions),
        "top_customers": top_counter(customer_counter),
        "top_suppliers": top_counter(supplier_counter),
    }


@app.get("/api/digest/combined")
def web_daily_digest_combined(
    date_value: Optional[str] = Query(None, alias="date"),
):
    if date_value:
        try:
            target = datetime.strptime(date_value[:10], "%Y-%m-%d").date()
        except ValueError:
            raise HTTPException(400, "date format must be YYYY-MM-DD")
    else:
        target = date.today() - timedelta(days=1)

    shop = _build_web_digest("shop", target)
    national = _build_web_digest("national", target)

    deals_count = shop["deals_count"] + national["deals_count"]
    total_contract = shop["total_contract"] + national["total_contract"]
    previous_deals = shop["previous_deals"] + national["previous_deals"]
    previous_contract = shop["previous_contract"] + national["previous_contract"]

    return {
        "date": target.isoformat(),
        "title": "BIData \u043a\u0443\u043d\u043b\u0438\u043a \u0431\u043e\u0437\u043e\u0440 \u0434\u0430\u0439\u0436\u0435\u0441\u0442\u0438",

        "deals_count": deals_count,
        "lots_count": shop["lots_count"] + national["lots_count"],
        "customers_count": shop["customers_count"] + national["customers_count"],
        "suppliers_count": shop["suppliers_count"] + national["suppliers_count"],
        "products_count": shop["products_count"] + national["products_count"],

        "total_contract": round(total_contract, 2),
        "average_deal": round(
            total_contract / deals_count if deals_count else 0,
            2
        ),

        "previous_deals": previous_deals,
        "previous_contract": round(previous_contract, 2),

        "deals_change_pct": _digest_change(
            deals_count,
            previous_deals
        ),
        "contract_change_pct": _digest_change(
            total_contract,
            previous_contract
        ),

        "markets": {
            "shop": shop,
            "national": national,
        }
    }

@app.get("/api/digest/{market_key}")
def web_daily_digest(
    market_key: str,
    date_value: Optional[str] = Query(None, alias="date"),
):
    if market_key not in {"shop", "national"}:
        raise HTTPException(400, "market must be shop or national")

    if date_value:
        try:
            target = datetime.strptime(date_value[:10], "%Y-%m-%d").date()
        except ValueError:
            raise HTTPException(400, "date format must be YYYY-MM-DD")
    else:
        target = date.today() - timedelta(days=1)

    return _build_web_digest(market_key, target)


# === BIDATA TENDER ANALYTICS API V1 ===

TENDER_UNIFIED_DB = DATA / "uzbi_tender_unified_active_v1.sqlite3"

TENDER_MARKET_TITLES = {
    "etender": "UZEX eTender",
    "xt_xarid": "XT-Xarid Tender",
    "shaffof": "Shaffof Xarid",
    "mc": "MC.UZ",
}



# =============================================================================
# BIDATA TENDER REGION NORMALIZATION V1
# =============================================================================

TENDER_REGION_CANONICAL = [
    "Qoraqalpog'iston Respublikasi",
    "Andijon viloyati",
    "Buxoro viloyati",
    "Jizzax viloyati",
    "Qashqadaryo viloyati",
    "Navoiy viloyati",
    "Namangan viloyati",
    "Samarqand viloyati",
    "Surxondaryo viloyati",
    "Sirdaryo viloyati",
    "Toshkent viloyati",
    "Farg'ona viloyati",
    "Xorazm viloyati",
    "Toshkent shahri",
]

TENDER_REGION_ALIASES = {
    "Qoraqalpog'iston Respublikasi": [
        "%qoraqalpog%",
        "%Р В Р’В Р РЋРІР‚СњР В Р’В Р вЂ™Р’В°Р В Р Р‹Р В РІР‚С™Р В Р’В Р вЂ™Р’В°Р В Р’В Р РЋРІР‚СњР В Р’В Р вЂ™Р’В°Р В Р’В Р вЂ™Р’В»Р В Р’В Р РЋРІР‚вЂќР В Р’В Р вЂ™Р’В°Р В Р’В Р РЋРІР‚Сњ%",
    ],
    "Andijon viloyati": [
        "%andijon%",
        "%Р В Р’В Р вЂ™Р’В°Р В Р’В Р В РІР‚В¦Р В Р’В Р СћРІР‚ВР В Р’В Р РЋРІР‚ВР В Р’В Р вЂ™Р’В¶Р В Р’В Р вЂ™Р’В°Р В Р’В Р В РІР‚В¦%",
    ],
    "Buxoro viloyati": [
        "%buxoro%",
        "%Р В Р’В Р вЂ™Р’В±Р В Р Р‹Р РЋРІР‚СљР В Р Р‹Р Р†Р вЂљР’В¦Р В Р’В Р вЂ™Р’В°Р В Р Р‹Р В РІР‚С™%",
    ],
    "Jizzax viloyati": [
        "%jizzax%",
        "%Р В Р’В Р СћРІР‚ВР В Р’В Р вЂ™Р’В¶Р В Р’В Р РЋРІР‚ВР В Р’В Р вЂ™Р’В·Р В Р’В Р вЂ™Р’В°Р В Р’В Р РЋРІР‚Сњ%",
    ],
    "Qashqadaryo viloyati": [
        "%qashqadaryo%",
        "%Р В Р’В Р РЋРІР‚СњР В Р’В Р вЂ™Р’В°Р В Р Р‹Р Р†РІР‚С™Р’В¬Р В Р’В Р РЋРІР‚СњaР В Р’В Р СћРІР‚ВР В Р’В Р вЂ™Р’В°Р В Р Р‹Р В РІР‚С™%",
        "%Р В Р’В Р РЋРІР‚СњР В Р’В Р вЂ™Р’В°Р В Р Р‹Р Р†РІР‚С™Р’В¬Р В Р’В Р РЋРІР‚СњР В Р’В Р вЂ™Р’В°Р В Р’В Р СћРІР‚ВР В Р’В Р вЂ™Р’В°Р В Р Р‹Р В РІР‚С™%",
    ],
    "Navoiy viloyati": [
        "%navoiy%",
        "%Р В Р’В Р В РІР‚В¦Р В Р’В Р вЂ™Р’В°Р В Р’В Р В РІР‚В Р В Р’В Р РЋРІР‚СћР В Р’В Р РЋРІР‚В%",
    ],
    "Namangan viloyati": [
        "%namangan%",
        "%Р В Р’В Р В РІР‚В¦Р В Р’В Р вЂ™Р’В°Р В Р’В Р РЋР’ВР В Р’В Р вЂ™Р’В°Р В Р’В Р В РІР‚В¦Р В Р’В Р РЋРІР‚вЂњР В Р’В Р вЂ™Р’В°Р В Р’В Р В РІР‚В¦%",
    ],
    "Samarqand viloyati": [
        "%samarqand%",
        "%Р В Р Р‹Р В РЎвЂњР В Р’В Р вЂ™Р’В°Р В Р’В Р РЋР’ВР В Р’В Р вЂ™Р’В°Р В Р Р‹Р В РІР‚С™Р В Р’В Р РЋРІР‚СњР В Р’В Р вЂ™Р’В°Р В Р’В Р В РІР‚В¦Р В Р’В Р СћРІР‚В%",
    ],
    "Surxondaryo viloyati": [
        "%surxondaryo%",
        "%Р В Р Р‹Р В РЎвЂњР В Р Р‹Р РЋРІР‚СљР В Р Р‹Р В РІР‚С™Р В Р Р‹Р Р†Р вЂљР’В¦Р В Р’В Р вЂ™Р’В°Р В Р’В Р В РІР‚В¦Р В Р’В Р СћРІР‚ВР В Р’В Р вЂ™Р’В°Р В Р Р‹Р В РІР‚С™%",
    ],
    "Sirdaryo viloyati": [
        "%sirdaryo%",
        "%Р В Р Р‹Р В РЎвЂњР В Р Р‹Р Р†Р вЂљРІвЂћвЂ“Р В Р Р‹Р В РІР‚С™Р В Р’В Р СћРІР‚ВР В Р’В Р вЂ™Р’В°Р В Р Р‹Р В РІР‚С™%",
    ],
    "Toshkent shahri": [
        "%toshkent shah%",
        "%Р В Р Р‹Р Р†Р вЂљРЎв„ўР В Р’В Р вЂ™Р’В°Р В Р Р‹Р Р†РІР‚С™Р’В¬Р В Р’В Р РЋРІР‚СњР В Р’В Р вЂ™Р’ВµР В Р’В Р В РІР‚В¦Р В Р Р‹Р Р†Р вЂљРЎв„ў Р В Р’В Р РЋРІР‚вЂњР В Р’В Р РЋРІР‚СћР В Р Р‹Р В РІР‚С™Р В Р’В Р РЋРІР‚СћР В Р’В Р СћРІР‚В%",
        "%Р В Р’В Р РЋРІР‚вЂњР В Р’В Р РЋРІР‚СћР В Р Р‹Р В РІР‚С™Р В Р’В Р РЋРІР‚СћР В Р’В Р СћРІР‚В Р В Р Р‹Р Р†Р вЂљРЎв„ўР В Р’В Р вЂ™Р’В°Р В Р Р‹Р Р†РІР‚С™Р’В¬Р В Р’В Р РЋРІР‚СњР В Р’В Р вЂ™Р’ВµР В Р’В Р В РІР‚В¦Р В Р Р‹Р Р†Р вЂљРЎв„ў%",
        "%Р В Р’В Р РЋРІР‚вЂњ. Р В Р Р‹Р Р†Р вЂљРЎв„ўР В Р’В Р вЂ™Р’В°Р В Р Р‹Р Р†РІР‚С™Р’В¬Р В Р’В Р РЋРІР‚СњР В Р’В Р вЂ™Р’ВµР В Р’В Р В РІР‚В¦Р В Р Р‹Р Р†Р вЂљРЎв„ў%",
    ],
    "Toshkent viloyati": [
        "%toshkent vil%",
        "%Р В Р Р‹Р Р†Р вЂљРЎв„ўР В Р’В Р вЂ™Р’В°Р В Р Р‹Р Р†РІР‚С™Р’В¬Р В Р’В Р РЋРІР‚СњР В Р’В Р вЂ™Р’ВµР В Р’В Р В РІР‚В¦Р В Р Р‹Р Р†Р вЂљРЎв„ўР В Р Р‹Р В РЎвЂњР В Р’В Р РЋРІР‚Сњ%Р В Р’В Р РЋРІР‚СћР В Р’В Р вЂ™Р’В±Р В Р’В Р вЂ™Р’В»Р В Р’В Р вЂ™Р’В°Р В Р Р‹Р В РЎвЂњР В Р Р‹Р Р†Р вЂљРЎв„ў%",
    ],
    "Farg'ona viloyati": [
        "%farg'ona%",
        "%farg ona%",
        "%fargР В Р вЂ Р В РІР‚С™Р вЂ™Р’Вona%",
        "%fargona%",
        "%Р В Р Р‹Р Р†Р вЂљРЎвЂєР В Р’В Р вЂ™Р’ВµР В Р Р‹Р В РІР‚С™Р В Р’В Р РЋРІР‚вЂњР В Р’В Р вЂ™Р’В°Р В Р’В Р В РІР‚В¦%",
    ],
    "Xorazm viloyati": [
        "%xorazm%",
        "%Р В Р Р‹Р Р†Р вЂљР’В¦Р В Р’В Р РЋРІР‚СћР В Р Р‹Р В РІР‚С™Р В Р’В Р вЂ™Р’ВµР В Р’В Р вЂ™Р’В·Р В Р’В Р РЋР’В%",
    ],
}


def _tender_region_aliases(region):
    region = str(region or "").strip()
    return TENDER_REGION_ALIASES.get(region, [region])


def _tender_where(
    market=None,
    q=None,
    subject_group=None,
    sector_group=None,
    customer=None,
    region=None,
    status=None,
):
    where = ["1=1"]
    params = []

    if market:
        markets = [
            x.strip()
            for x in str(market).split(",")
            if x.strip()
        ]
        if markets:
            ph = ",".join("?" for _ in markets)
            where.append(f"market IN ({ph})")
            params.extend(markets)

    if q:
        qv = f"%{str(q).strip()}%"
        where.append("""
        (
            coalesce(product_name,'') LIKE ? OR
            coalesce(product_code,'') LIKE ? OR
            coalesce(tender_name,'') LIKE ? OR
            coalesce(source_tender_id,'') LIKE ? OR
            coalesce(source_lot_id,'') LIKE ?
        )
        """)
        params.extend([qv] * 5)

    if subject_group:
        where.append("subject_group = ?")
        params.append(subject_group)

    if sector_group:
        where.append("sector_group = ?")
        params.append(sector_group)

    if customer:
        where.append("coalesce(customer_name,'') LIKE ?")
        params.append(f"%{str(customer).strip()}%")

    if region:
        aliases = _tender_region_aliases(region)
        parts = []
        for alias in aliases:
            parts.append("lower(coalesce(region_name,'')) LIKE lower(?)")
            params.append(alias)
        where.append("(" + " OR ".join(parts) + ")")

    if status:
        where.append("status = ?")
        params.append(status)

    return " AND ".join(where), params


@app.get("/api/tenders/active/summary")
def tender_active_summary(
    market: Optional[str] = None,
    q: Optional[str] = None,
    subject_group: Optional[str] = None,
    sector_group: Optional[str] = None,
    customer: Optional[str] = None,
    region: Optional[str] = None,
    status: Optional[str] = None,
):
    w, p = _tender_where(
        market,
        q,
        subject_group,
        sector_group,
        customer,
        region,
        status,
    )

    with ro(TENDER_UNIFIED_DB) as c:

        r = c.execute(
            f"""
            SELECT
                COUNT(DISTINCT market || ':' || source_tender_id)
                    AS tenders,

                COUNT(
                    DISTINCT market || ':' ||
                    source_tender_id || ':' ||
                    coalesce(source_lot_id,'')
                ) AS lots,

                SUM(
                    CASE
                        WHEN source_level='item' THEN 1
                        ELSE 0
                    END
                ) AS items,

                COUNT(
                    DISTINCT CASE
                        WHEN coalesce(product_name,'') <> ''
                        THEN product_name
                    END
                ) AS products,

                COUNT(
                    DISTINCT CASE
                        WHEN coalesce(customer_name,'') <> ''
                        THEN customer_name
                    END
                ) AS customers,

                COUNT(
                    DISTINCT CASE
                        WHEN coalesce(subject_group,'') <> ''
                        THEN subject_group
                    END
                ) AS subject_groups,

                COUNT(
                    DISTINCT CASE
                        WHEN coalesce(sector_group,'') <> ''
                        THEN sector_group
                    END
                ) AS sector_groups

            FROM fact_tender_active
            WHERE {w}
            """,
            p,
        ).fetchone()

        # Amount must be counted once per tender/lot rather than once
        # for every product row.
        amount_row = c.execute(
            f"""
            SELECT coalesce(SUM(x.amount),0) AS amount
            FROM (
                SELECT
                    market,
                    source_tender_id,
                    coalesce(source_lot_id,'') AS source_lot_id,
                    MAX(coalesce(amount,0)) AS amount
                FROM fact_tender_active
                WHERE {w}
                GROUP BY
                    market,
                    source_tender_id,
                    coalesce(source_lot_id,'')
            ) x
            """,
            p,
        ).fetchone()

        out = dict(r)
        out["amount"] = amount_row["amount"] or 0
        out["rows"] = c.execute(
            f"SELECT COUNT(*) FROM fact_tender_active WHERE {w}",
            p,
        ).fetchone()[0]

        return out


@app.get("/api/tenders/active")
def tender_active_list(
    market: Optional[str] = None,
    q: Optional[str] = None,
    subject_group: Optional[str] = None,
    sector_group: Optional[str] = None,
    customer: Optional[str] = None,
    region: Optional[str] = None,
    status: Optional[str] = None,
    limit: int = Query(50, ge=1, le=500),
    offset: int = Query(0, ge=0),
):
    w, p = _tender_where(
        market,
        q,
        subject_group,
        sector_group,
        customer,
        region,
        status,
    )

    with ro(TENDER_UNIFIED_DB) as c:
        total = c.execute(
            f"""
            SELECT COUNT(*)
            FROM fact_tender_active
            WHERE {w}
            """,
            p,
        ).fetchone()[0]

        rows = [
            dict(x)
            for x in c.execute(
                f"""
                SELECT
                    row_id,
                    market,
                    source_tender_id,
                    source_lot_id,
                    source_item_id,

                    tender_name,
                    product_name,
                    product_code,
                    category_name,

                    quantity,
                    unit_name,
                    unit_price,
                    amount,

                    status,

                    customer_id,
                    customer_inn,
                    customer_name,

                    region_name,
                    district_name,
                    address,

                    subject_group,
                    sector_group,
                    product_type,

                    subject_method,
                    subject_confidence,
                    sector_method,
                    sector_confidence,

                    source_sector,
                    object_type,
                    service_type,

                    start_date,
                    end_date,
                    first_seen_at,
                    last_seen_at,

                    winner_id,
                    winner_name,
                    public_url,
                    source_level

                FROM fact_tender_active
                WHERE {w}

                ORDER BY
                    coalesce(end_date,'9999-12-31') ASC,
                    market,
                    source_tender_id,
                    source_lot_id,
                    row_id

                LIMIT ? OFFSET ?
                """,
                p + [limit, offset],
            ).fetchall()
        ]

        for row in rows:
            row["market_title"] = TENDER_MARKET_TITLES.get(
                row.get("market"),
                row.get("market"),
            )

        return {
            "total": total,
            "limit": limit,
            "offset": offset,
            "items": rows,
        }


@app.get("/api/tenders/active/filters")
def tender_active_filters():
    with ro(TENDER_UNIFIED_DB) as c:

        def values(col):
            return [
                r[0]
                for r in c.execute(
                    f"""
                    SELECT DISTINCT {col}
                    FROM fact_tender_active
                    WHERE coalesce({col},'') <> ''
                    ORDER BY {col}
                    """
                ).fetchall()
            ]

        markets = []
        for r in c.execute("""
            SELECT
                market,
                COUNT(DISTINCT source_tender_id) AS tenders
            FROM fact_tender_active
            GROUP BY market
            ORDER BY market
        """):
            markets.append({
                "key": r["market"],
                "title": TENDER_MARKET_TITLES.get(
                    r["market"],
                    r["market"],
                ),
                "tenders": r["tenders"],
            })

        return {
            "markets": markets,
            "subject_groups": values("subject_group"),
            "sector_groups": values("sector_group"),
            "regions": TENDER_REGION_CANONICAL,
            "statuses": values("status"),
        }


@app.get("/api/tenders/active/top/{dim}")
def tender_active_top(
    dim: str,
    market: Optional[str] = None,
    q: Optional[str] = None,
    subject_group: Optional[str] = None,
    sector_group: Optional[str] = None,
    customer: Optional[str] = None,
    region: Optional[str] = None,
    status: Optional[str] = None,
    limit: int = Query(10, ge=1, le=50),
):
    dimensions = {
        "subjects": "subject_group",
        "sectors": "sector_group",
        "customers": "customer_name",
        "products": "product_name",
    }

    where, params = _tender_where(
        market=market,
        q=q,
        subject_group=subject_group,
        sector_group=sector_group,
        customer=customer,
        region=region,
        status=status,
    )

    # ------------------------------------------------------
    # REGIONS: normalize raw source spellings into canonical
    # ------------------------------------------------------
    if dim == "regions":
        out = []

        with ro(TENDER_UNIFIED_DB) as c:
            for canonical in TENDER_REGION_CANONICAL:

                aliases = _tender_region_aliases(canonical)

                alias_sql = []
                alias_params = []

                for alias in aliases:
                    alias_sql.append(
                        "lower(coalesce(region_name,'')) LIKE lower(?)"
                    )
                    alias_params.append(alias)

                region_sql = "(" + " OR ".join(alias_sql) + ")"

                row = c.execute(
                    f"""
                    SELECT
                        COUNT(
                            DISTINCT market || ':' ||
                            source_tender_id || ':' ||
                            coalesce(source_lot_id,'')
                        ) AS tenders,
                        COUNT(*) AS rows
                    FROM fact_tender_active
                    WHERE ({where})
                      AND {region_sql}
                    """,
                    params + alias_params,
                ).fetchone()

                tenders = int(row["tenders"] or 0)
                rows_count = int(row["rows"] or 0)

                if tenders == 0:
                    continue

                amount_row = c.execute(
                    f"""
                    SELECT COALESCE(SUM(x.amount),0) AS amount
                    FROM (
                        SELECT
                            market,
                            source_tender_id,
                            coalesce(source_lot_id,'') AS source_lot_id,
                            MAX(amount) AS amount
                        FROM fact_tender_active
                        WHERE ({where})
                          AND {region_sql}
                        GROUP BY
                            market,
                            source_tender_id,
                            coalesce(source_lot_id,'')
                    ) x
                    """,
                    params + alias_params,
                ).fetchone()

                out.append({
                    "name": canonical,
                    "tenders": tenders,
                    "rows": rows_count,
                    "amount": float(amount_row["amount"] or 0),
                })

        out.sort(
            key=lambda x: (
                int(x["tenders"]),
                float(x["amount"]),
            ),
            reverse=True,
        )

        return out[:limit]

    # ------------------------------------------------------
    # Other dimensions
    # ------------------------------------------------------
    if dim not in dimensions:
        raise HTTPException(
            400,
            "dim must be subjects, sectors, customers, regions or products",
        )

    col = dimensions[dim]
    where2 = f"({where}) AND coalesce({col},'') <> ''"

    with ro(TENDER_UNIFIED_DB) as c:
        rows = c.execute(
            f"""
            SELECT
                {col} AS name,

                COUNT(
                    DISTINCT market || ':' ||
                    source_tender_id || ':' ||
                    coalesce(source_lot_id,'')
                ) AS tenders,

                COUNT(*) AS rows

            FROM fact_tender_active
            WHERE {where2}

            GROUP BY {col}
            ORDER BY tenders DESC
            LIMIT ?
            """,
            params + [limit],
        ).fetchall()

        out = []

        for r in rows:
            d = dict(r)

            amount_row = c.execute(
                f"""
                SELECT COALESCE(SUM(x.amount),0)
                FROM (
                    SELECT
                        market,
                        source_tender_id,
                        coalesce(source_lot_id,'') AS source_lot_id,
                        MAX(amount) AS amount
                    FROM fact_tender_active
                    WHERE {where2}
                      AND {col}=?
                    GROUP BY
                        market,
                        source_tender_id,
                        coalesce(source_lot_id,'')
                ) x
                """,
                params + [d["name"]],
            ).fetchone()

            d["amount"] = float(amount_row[0] or 0)
            out.append(d)

        return out

@app.get("/api/tenders/active/grouped")
def tender_active_grouped(
    market: Optional[str] = None,
    q: Optional[str] = None,
    subject_group: Optional[str] = None,
    sector_group: Optional[str] = None,
    customer: Optional[str] = None,
    region: Optional[str] = None,
    status: Optional[str] = None,
    limit: int = Query(50, ge=1, le=200),
    offset: int = Query(0, ge=0),
):
    where, params = _tender_where(
        market=market,
        q=q,
        subject_group=subject_group,
        sector_group=sector_group,
        customer=customer,
        region=region,
        status=status,
    )

    base_where = ("WHERE " + where) if where else ""

    with ro(TENDER_UNIFIED_DB) as c:

        total = c.execute(
            f"""
            SELECT COUNT(*)
            FROM (
                SELECT
                    market,
                    source_tender_id
                FROM fact_tender_active
                {base_where}
                GROUP BY
                    market,
                    source_tender_id
            )
            """,
            params,
        ).fetchone()[0]

        rows = c.execute(
            f"""
            SELECT
                market,
                source_tender_id,

                MIN(COALESCE(source_lot_id,'')) AS source_lot_id,

                MAX(tender_name) AS tender_name,
                MAX(product_name) AS sample_product_name,

                MAX(subject_group) AS subject_group,
                MAX(sector_group) AS sector_group,

                MAX(customer_name) AS customer_name,
                MAX(region_name) AS region_name,
                MAX(district_name) AS district_name,

                MAX(status) AS status,
                MAX(end_date) AS end_date,

                SUM(COALESCE(amount,0)) AS amount,

                COUNT(*) AS row_count,

                COUNT(
                    DISTINCT CASE
                        WHEN source_lot_id IS NOT NULL
                         AND TRIM(CAST(source_lot_id AS TEXT)) <> ''
                        THEN CAST(source_lot_id AS TEXT)
                    END
                ) AS lot_count,

                SUM(
                    CASE
                        WHEN source_level='item' THEN 1
                        ELSE 0
                    END
                ) AS item_count,

                MAX(public_url) AS public_url

            FROM fact_tender_active
            {base_where}

            GROUP BY
                market,
                source_tender_id

            ORDER BY
                CASE
                    WHEN ? IS NULL OR ? = ''
                    THEN (
                        SELECT COUNT(*)
                        FROM fact_tender_active z
                        WHERE z.market = fact_tender_active.market
                          AND (
                              COALESCE(z.end_date,'9999') <
                              COALESCE(MAX(fact_tender_active.end_date),'9999')
                              OR (
                                  COALESCE(z.end_date,'9999') =
                                  COALESCE(MAX(fact_tender_active.end_date),'9999')
                                  AND CAST(z.source_tender_id AS TEXT) <
                                      CAST(fact_tender_active.source_tender_id AS TEXT)
                              )
                          )
                    )
                    ELSE 0
                END,
                market,
                CASE WHEN MAX(end_date) IS NULL THEN 1 ELSE 0 END,
                MAX(end_date),
                source_tender_id

            LIMIT ? OFFSET ?
            """,
            params + [market, market, limit, offset],
        ).fetchall()

        items = []

        for r in rows:
            x = dict(r)
            x["market_title"] = TENDER_MARKET_TITLES.get(
                x["market"],
                x["market"],
            )
            items.append(x)

        return {
            "total": total,
            "limit": limit,
            "offset": offset,
            "items": items,
        }


ETENDER_PENDING_DB = DATA / "etender_pending_lots.sqlite3"

def _etender_rich_detail(tender_id=None, lot_id=None):
    if not ETENDER_PENDING_DB.exists():
        return None

    conn = ro(ETENDER_PENDING_DB)
    conn.row_factory = sqlite3.Row

    try:
        lot = None

        if tender_id:
            try:
                lot = conn.execute("""
                    SELECT *
                    FROM pending_lots
                    WHERE lot_id = ?
                    LIMIT 1
                """, (int(tender_id),)).fetchone()
            except Exception:
                pass

        if lot is None and lot_id:
            lot = conn.execute("""
                SELECT *
                FROM pending_lots
                WHERE CAST(display_no AS TEXT) = ?
                   OR CAST(lot_id AS TEXT) = ?
                LIMIT 1
            """, (str(lot_id), str(lot_id))).fetchone()

        if lot is None:
            return None

        lot = dict(lot)

        products = conn.execute("""
            SELECT
                id,
                lot_id,
                item_id,
                order_num,
                category_id,
                category_name,
                product_id,
                product_name,
                product_type_id,
                product_code,
                quantity,
                unit_name,
                price,
                cost,
                description,
                guarantee_term,
                delivery_term,
                term_period_name,
                source_json
            FROM pending_products
            WHERE lot_id = ?
            ORDER BY COALESCE(order_num, 999999), id
        """, (lot["lot_id"],)).fetchall()

        items = []
        for r in products:
            x = dict(r)

            props = []
            try:
                import json
                raw = json.loads(x.get("source_json") or "{}")
                props = raw.get("Js_Properties") or []
            except Exception:
                props = []

            x["properties"] = props
            x.pop("source_json", None)

            items.append(x)

        return {
            "market": "etender",
            "platform_name": "UZEX eTender",
            "lot": lot,
            "items": items,
            "items_count": len(items),
        }

    finally:
        conn.close()



SHAFFOF_DB = DATA / "shaffof_tenders.sqlite3"
MC_TENDERS_DB = DATA / "mc_tenders.sqlite3"
XT_XARID_TENDER_DB = DATA / "xt_xarid_tenders.sqlite3"



def _xt_xarid_rich_detail(tender_id=None, lot_id=None):
    if not XT_XARID_TENDER_DB.exists():
        return None

    conn = ro(XT_XARID_TENDER_DB)
    conn.row_factory = sqlite3.Row

    try:
        try:
            proc_id = int(tender_id)
        except Exception:
            return None

        tender_row = conn.execute(
            "SELECT * FROM tenders WHERE proc_id=? LIMIT 1",
            (proc_id,)
        ).fetchone()

        if not tender_row:
            return None

        tender = dict(tender_row)

        lots = [
            dict(r)
            for r in conn.execute(
                """
                SELECT *
                FROM lots
                WHERE proc_id=?
                ORDER BY lot_id
                """,
                (proc_id,)
            ).fetchall()
        ]

        try:
            selected_lot_id = int(lot_id)
        except Exception:
            selected_lot_id = lots[0]["lot_id"] if lots else None

        selected_lot = next(
            (x for x in lots if x.get("lot_id") == selected_lot_id),
            None
        )

        items = [
            dict(r)
            for r in conn.execute(
                """
                SELECT *
                FROM items
                WHERE proc_id=? AND lot_id=?
                ORDER BY item_id
                """,
                (proc_id, selected_lot_id)
            ).fetchall()
        ]

        prop_rows = [
            dict(r)
            for r in conn.execute(
                """
                SELECT *
                FROM item_properties
                WHERE proc_id=? AND lot_id=?
                ORDER BY item_id, rowid
                """,
                (proc_id, selected_lot_id)
            ).fetchall()
        ]

        props_by_item = {}

        for pr in prop_rows:
            props_by_item.setdefault(pr["item_id"], []).append(pr)

        for item in items:
            item["properties"] = props_by_item.get(item["item_id"], [])

        return {
            "tender": tender,
            "lots": lots,
            "selected_lot": selected_lot,
            "items": items,
        }

    finally:
        conn.close()


def _mc_rich_detail(tender_id=None, lot_id=None):
    if not MC_TENDERS_DB.exists():
        return None

    conn = ro(MC_TENDERS_DB)
    conn.row_factory = sqlite3.Row

    try:
        row = None

        # Internal MC.UZ tender ID
        for value in (tender_id, lot_id):
            if value in (None, ""):
                continue

            try:
                internal_id = int(str(value))
            except Exception:
                internal_id = None

            if internal_id is not None:
                row = conn.execute("""
                    SELECT *
                    FROM tenders
                    WHERE id = ?
                    LIMIT 1
                """, (internal_id,)).fetchone()

                if row:
                    break

        # Public lot number fallback
        if row is None:
            for value in (lot_id, tender_id):
                if value in (None, ""):
                    continue

                value = str(value)

                row = conn.execute("""
                    SELECT *
                    FROM tenders
                    WHERE unique_name = ?
                       OR lotid = ?
                    LIMIT 1
                """, (value, value)).fetchone()

                if row:
                    break

        if row is None:
            return None

        tender = dict(row)
        source_id = tender.get("id")

        items = [
            dict(x)
            for x in conn.execute("""
                SELECT *
                FROM tender_items
                WHERE tender_id = ?
                ORDER BY COALESCE(order_id, 999999), id
            """, (source_id,)).fetchall()
        ]

        props = [
            dict(x)
            for x in conn.execute("""
                SELECT *
                FROM tender_item_props
                WHERE tender_id = ?
                ORDER BY item_id, prop_numb, val_numb, id
            """, (source_id,)).fetchall()
        ]

        props_by_item = {}

        for x in props:
            props_by_item.setdefault(x.get("item_id"), []).append(x)

        for item in items:
            item["properties"] = props_by_item.get(item.get("id"), [])

        files = [
            dict(x)
            for x in conn.execute("""
                SELECT *
                FROM tender_files
                WHERE tender_id = ?
                ORDER BY id
            """, (source_id,)).fetchall()
        ]

        history = [
            dict(x)
            for x in conn.execute("""
                SELECT *
                FROM tender_status_history
                WHERE tender_id = ?
                ORDER BY seen_at, id
            """, (source_id,)).fetchall()
        ]

        return {
            "tender": tender,
            "items": items,
            "files": files,
            "history": history,
        }

    finally:
        conn.close()


def _shaffof_rich_detail(tender_id=None, lot_id=None):
    if not SHAFFOF_DB.exists():
        return None

    conn = ro(SHAFFOF_DB)
    conn.row_factory = sqlite3.Row

    try:
        row = None

        # Unified source_tender_id odatda tenders.id
        if tender_id:
            try:
                row = conn.execute("""
                    SELECT *
                    FROM tenders
                    WHERE id = ?
                    LIMIT 1
                """, (int(tender_id),)).fetchone()
            except Exception:
                pass

        # Fallback: public unique_name
        if row is None and lot_id:
            row = conn.execute("""
                SELECT *
                FROM tenders
                WHERE CAST(unique_name AS TEXT) = ?
                   OR CAST(id AS TEXT) = ?
                LIMIT 1
            """, (str(lot_id), str(lot_id))).fetchone()

        if row is None:
            return None

        tender = dict(row)
        tid = tender["id"]

        process_rows = conn.execute("""
            SELECT
                id,
                tender_id,
                action,
                created_at,
                end_date
            FROM tender_process
            WHERE tender_id = ?
            ORDER BY created_at, id
        """, (tid,)).fetchall()

        file_rows = conn.execute("""
            SELECT
                id,
                tender_id,
                source_kind,
                source_file_id,
                file_url,
                file_name,
                file_type,
                role,
                created_at,
                updated_at
            FROM tender_files
            WHERE tender_id = ?
            ORDER BY id
        """, (tid,)).fetchall()

        return {
            "market": "shaffof",
            "platform_name": "Shaffof Xarid",
            "tender": tender,
            "process": [dict(x) for x in process_rows],
            "files": [dict(x) for x in file_rows],
            "process_count": len(process_rows),
            "files_count": len(file_rows),
        }

    finally:
        conn.close()

@app.get("/api/tenders/active/detail")
def tender_active_detail(
    market: str,
    tender_id: str,
    lot_id: Optional[str] = None,
):
    where = [
        "market = ?",
        "CAST(source_tender_id AS TEXT) = ?",
    ]
    params = [market, tender_id]

    if lot_id is not None:
        where.append("COALESCE(CAST(source_lot_id AS TEXT),'') = ?")
        params.append(lot_id)

    sql_where = " AND ".join(where)

    with ro(TENDER_UNIFIED_DB) as c:

        rows = c.execute(
            f"""
            SELECT *
            FROM fact_tender_active
            WHERE {sql_where}
            ORDER BY
                CASE WHEN product_name IS NULL OR product_name='' THEN 1 ELSE 0 END,
                product_name,
                source_item_id
            """,
            params,
        ).fetchall()

        if not rows:
            raise HTTPException(404, "Tender not found")

        first = dict(rows[0])

        header = {
            "market": first.get("market"),
            "market_title": TENDER_MARKET_TITLES.get(
                first.get("market"),
                first.get("market"),
            ),
            "source_tender_id": first.get("source_tender_id"),
            "source_lot_id": first.get("source_lot_id"),
            "tender_name": first.get("tender_name"),
            "status": first.get("status"),
            "customer_id": first.get("customer_id"),
            "customer_inn": first.get("customer_inn"),
            "customer_name": first.get("customer_name"),
            "region_name": first.get("region_name"),
            "district_name": first.get("district_name"),
            "address": first.get("address"),
            "subject_group": first.get("subject_group"),
            "sector_group": first.get("sector_group"),
            "source_sector": first.get("source_sector"),
            "object_type": first.get("object_type"),
            "service_type": first.get("service_type"),
            "start_date": first.get("start_date"),
            "end_date": first.get("end_date"),
            "amount": first.get("amount"),
            "winner_id": first.get("winner_id"),
            "winner_name": first.get("winner_name"),
            "public_url": first.get("public_url"),
        }

        positions = []

        for r in rows:
            d = dict(r)

            if d.get("source_level") != "item":
                continue

            positions.append({
                "source_item_id": d.get("source_item_id"),
                "product_name": d.get("product_name"),
                "product_code": d.get("product_code"),
                "category_name": d.get("category_name"),
                "quantity": d.get("quantity"),
                "unit_name": d.get("unit_name"),
                "unit_price": d.get("unit_price"),
                "amount": d.get("amount"),
                "subject_group": d.get("subject_group"),
                "product_type": d.get("product_type"),
            })

        header["positions"] = positions
        header["positions_count"] = len(positions)
        header["rows_count"] = len(rows)

        # Lot/tender total:
        # item-level sources -> sum position amounts
        # tender-only sources -> keep source amount
        item_amounts = [
            float(x["amount"])
            for x in positions
            if x.get("amount") is not None
        ]

        if item_amounts:
            header["amount"] = sum(item_amounts)

        # eTender-specific rich detail
        if market == "etender":
            rich = _etender_rich_detail(
                tender_id=tender_id,
                lot_id=lot_id,
            )

            if rich:
                lot = rich.get("lot") or {}
                items = rich.get("items") or []

                header["detail_type"] = "etender_rich"

                header["etender"] = {
                    "lot_id": lot.get("lot_id"),
                    "display_no": lot.get("display_no"),
                    "type_name": lot.get("type_name"),
                    "status_id": lot.get("status_id"),
                    "status_name": lot.get("status_name"),

                    "start_date": lot.get("start_date"),
                    "end_date": lot.get("end_date"),
                    "clarification_date": lot.get("clarification_date"),

                    "start_cost": lot.get("start_cost"),
                    "currency_name": lot.get("currency_name"),
                    "currency_code": lot.get("currency_code"),

                    "valuation_name": lot.get("valuation_name"),
                    "tech_coef": lot.get("tech_coef"),
                    "cost_coef": lot.get("cost_coef"),
                    "min_point": lot.get("min_point"),

                    "customer_name": lot.get("customer_name"),
                    "customer_tin": lot.get("customer_tin"),
                    "customer_region_name": lot.get("customer_region_name"),
                    "customer_district_name": lot.get("customer_district_name"),
                    "customer_street": lot.get("customer_street"),
                    "customer_zip": lot.get("customer_zip"),

                    "delivering_region_name": lot.get("delivering_region_name"),
                    "delivering_district_name": lot.get("delivering_district_name"),
                    "delivering_address": lot.get("delivering_address"),
                    "delivering_phone": lot.get("delivering_phone"),

                    "payment_type_name": lot.get("payment_type_name"),
                    "advance_payment_perc": lot.get("advance_payment_perc"),
                    "term_payment_days": lot.get("term_payment_days"),

                    "pledge_name": lot.get("pledge_name"),
                    "pledge_value": lot.get("pledge_value"),
                    "block_pledge_type": lot.get("block_pledge_type"),
                    "is_bank_guarantee": lot.get("is_bank_guarantee"),

                    "financing_source": lot.get("financing_source"),
                    "consider_procedure": lot.get("consider_procedure"),

                    "addon_description": lot.get("addon_description"),
                    "technical_description": lot.get("technical_description"),

                    "has_benefits": lot.get("has_benefits"),
                    "is_local_manufacturs": lot.get("is_local_manufacturs"),

                    "is_prolonged": lot.get("is_prolonged"),
                    "prolong_date": lot.get("prolong_date"),
                    "prolonged_days": lot.get("prolonged_days"),

                    "authorized_count": lot.get("authorized_count"),
                    "nonauthorized_count": lot.get("nonauthorized_count"),

                    "detail_url": lot.get("detail_url"),
                }

                header["etender_positions"] = items
                header["etender_positions_count"] = len(items)


        # MC.UZ-specific rich detail

        if market == "xt_xarid":
            rich = _xt_xarid_rich_detail(
                tender_id=tender_id,
                lot_id=lot_id,
            )

            if rich:
                header["detail_type"] = "xt_xarid_rich"
                header["xt_xarid"] = rich.get("tender") or {}
                header["xt_lots"] = rich.get("lots") or []
                header["xt_lots_count"] = len(header["xt_lots"])
                header["xt_selected_lot"] = rich.get("selected_lot") or {}
                header["xt_items"] = rich.get("items") or []
                header["xt_items_count"] = len(header["xt_items"])

        if market == "mc":
            rich = _mc_rich_detail(
                tender_id=tender_id,
                lot_id=lot_id,
            )

            if rich:
                t = rich.get("tender") or {}
                items = rich.get("items") or []
                files = rich.get("files") or []
                history = rich.get("history") or []

                header["detail_type"] = "mc_rich"

                header["mc"] = t
                header["mc_items"] = items
                header["mc_items_count"] = len(items)

                header["mc_files"] = files
                header["mc_files_count"] = len(files)

                header["mc_history"] = history
                header["mc_history_count"] = len(history)


        # Shaffof Xarid-specific rich detail
        if market == "shaffof":
            rich = _shaffof_rich_detail(
                tender_id=tender_id,
                lot_id=lot_id,
            )

            if rich:
                t = rich.get("tender") or {}

                header["detail_type"] = "shaffof_rich"

                header["shaffof"] = {
                    "id": t.get("id"),
                    "unique_name": t.get("unique_name"),
                    "gnk_id": t.get("gnk_id"),

                    "name": t.get("name"),

                    "status_code": t.get("status_code"),
                    "status_name": t.get("status_name"),

                    "contract_status_code": t.get("contract_status_code"),
                    "contract_status_name": t.get("contract_status_name"),

                    "lot_type_name": t.get("lot_type_name"),
                    "start_price": t.get("start_price"),
                    "with_nds": t.get("with_nds"),

                    "created_at": t.get("created_at"),
                    "confirmed_date": t.get("confirmed_date"),
                    "placement_term": t.get("placement_term"),
                    "placement_term_days": t.get("placement_term_days"),

                    "discussion_start_date": t.get("discussion_start_date"),
                    "discussion_end_date": t.get("discussion_end_date"),

                    "end_term_work": t.get("end_term_work"),
                    "end_term_work_days": t.get("end_term_work_days"),

                    "region": t.get("region"),
                    "district": t.get("district"),
                    "address": t.get("address"),

                    "object_type": t.get("object_type"),
                    "service_type": t.get("service_type"),

                    "financial_source": t.get("financial_source"),
                    "category_score": t.get("category_score"),
                    "complexity_category": t.get("complexity_category"),

                    "dastur": t.get("dastur"),
                    "soxa": t.get("soxa"),
                    "asos": t.get("asos"),
                    "resolution": t.get("resolution"),

                    "required_percent": t.get("required_percent"),
                    "required_percent_amount": t.get("required_percent_amount"),

                    "customer_inn": t.get("customer_inn"),
                    "customer_name": t.get("customer_name"),
                    "customer_director": t.get("customer_director"),

                    "consulting_inn": t.get("consulting_inn"),
                    "consulting_name": t.get("consulting_name"),
                    "consulting_contact_person": t.get("consulting_contact_person"),
                    "consulting_email": t.get("consulting_email"),
                    "consulting_phone": t.get("consulting_phone"),

                    "winner_inn": t.get("winner_inn"),
                    "winner_name": t.get("winner_name"),
                    "winner_rating": t.get("winner_rating"),
                    "winner_rating_title": t.get("winner_rating_title"),
                    "winner_region": t.get("winner_region"),

                    "bidders_count": t.get("bidders_count"),

                    "contract_file_url": t.get("contract_file_url"),
                    "protocol_file_url": t.get("protocol_file_url"),
                    "dalolatnoma_file_url": t.get("dalolatnoma_file_url"),

                    "latitude": t.get("latitude"),
                    "longitude": t.get("longitude"),
                }

                header["shaffof_process"] = rich.get("process") or []
                header["shaffof_files"] = rich.get("files") or []
                header["shaffof_process_count"] = rich.get("process_count", 0)
                header["shaffof_files_count"] = rich.get("files_count", 0)

        return header












# ETENDER_VIEW_V1: isolated completed eTender routes.
from etender_backend import register_etender
register_etender(app, platform, ro)

# XT_TENDER_VIEW_V1: isolated read-only tender routes.
from xt_tender_backend import register_xt_tender
register_xt_tender(app, platform, ro)

# SHAFFOF_TENDER_VIEW_V1: isolated read-only tender routes.
from shaffof_tender_backend import register_shaffof_tender
register_shaffof_tender(app, platform, ro)

# MC_TENDER_VIEW_V1: isolated read-only analytics.
from mc_tender_backend import register_mc_tender
register_mc_tender(app, platform, ro)

# TENDER_FINANCE_V2
from tender_finance_backend import register as register_tender_finance
register_tender_finance(globals())


# === DAILY DIGEST V3 SNAPSHOTS ===
from digest_daily_backend import register_digest_routes
register_digest_routes(app, ROOT)
